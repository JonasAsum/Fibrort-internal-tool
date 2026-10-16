"""Fixtures for the Excel import tests. Synthetic workbooks are built here; the real
sample folder is used only when EXCEL_SAMPLE_DIR points at it (or the default exists)."""
import importlib
import os
import sys

import openpyxl
import pytest

ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))  # importer/
CRM_IMPORTER = os.environ.get("CRM_IMPORTER_DIR", ROOT)  # where the excel/ and visits/ packages live
SAMPLE_DIR = os.environ.get(
    "EXCEL_SAMPLE_DIR",
    "/Users/jonas/Mac Documents/Shanghai Company/CRM Excel Lists 2026-10-11",
)
GOLDEN = os.path.join(os.path.dirname(os.path.abspath(__file__)), "golden_real_folder.json")

if CRM_IMPORTER not in sys.path:
    sys.path.append(CRM_IMPORTER)  # appended: our tests/ package must win over the importer's tests/
if ROOT not in sys.path:
    sys.path.insert(0, ROOT)


def make_book(path, sheets, title="synthetic"):
    """sheets: {sheet_name: {"cells": {"A1": value}, "merges": ["F9:F12"], "hidden": bool}}.
    At least one sheet must be visible. Values are invented by the caller."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.properties.title = title  # changes the file bytes without changing any sheet
    for name, spec in sheets.items():
        ws = wb.create_sheet(name)
        for coord, value in spec.get("cells", {}).items():
            ws[coord] = value
        for rng in spec.get("merges", ()):
            ws.merge_cells(rng)
        if spec.get("hidden"):
            ws.sheet_state = "hidden"
    wb.save(path)
    return str(path)


@pytest.fixture
def book(tmp_path):
    """Writes a synthetic workbook into tmp_path and returns its path."""
    def _make(name, sheets, title="synthetic"):
        return make_book(tmp_path / name, sheets, title)
    return _make


@pytest.fixture(scope="session")
def sample_dir():
    if not os.path.isdir(SAMPLE_DIR):
        pytest.skip(f"sample folder not found; set EXCEL_SAMPLE_DIR (looked at {SAMPLE_DIR})")
    return SAMPLE_DIR


@pytest.fixture(scope="session")
def real_files(sample_dir):
    from tests.excel_scan import scan_folder
    return scan_folder(sample_dir)


@pytest.fixture(scope="session")
def golden_counts():
    import json
    if not os.path.exists(GOLDEN):
        pytest.skip("no golden file yet: python -m tests.excel_scan \"$EXCEL_SAMPLE_DIR\" --golden > tests/golden_real_folder.json")
    with open(GOLDEN, encoding="utf-8") as fh:
        return json.load(fh)


@pytest.fixture
def importer():
    """Imports a module of the existing importer (excel.run, visits.reader, ...) or skips with the reason."""
    def _load(name):
        try:
            return importlib.import_module(name)
        except ImportError as exc:
            pytest.skip(f"{name} not importable from {CRM_IMPORTER}: {exc}")
    return _load
