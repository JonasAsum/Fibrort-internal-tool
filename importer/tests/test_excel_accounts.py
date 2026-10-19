"""Customer-list importer (folders 04 and 05): every layout, merging across files, the finance fields that
must never come through, and the Margince adapter against a fake CRM. Every value here is invented; the
real-folder tests assert counts only."""
import glob
import itertools
import json
import os
import shutil

import openpyxl
import pytest

from common.client import ApiError
from excel import accounts, accounts_margince as am
from excel.companies import Master
from tests.conftest import SAMPLE_DIR

ERP_HEADER = ["序号", "客户编码", "客户名称", "客户简称", "纳税人登记号", "助记码", "客户分类编码", "客户分类名称", "地区名称",
              "所属行业", "地址", "开户银行", "银行账号", "法人", "Email地址", "联系人", "电话", "传真", "手机",
              "专营业务员名称", "销售默认收付款协议编码", "发货地址", "发运方式名称", "发货仓库名称", "分管部门名称",
              "所属银行名称", "币种", "潜在客户编码", "税率%", "省直辖市", "Purpose_Topics Visit"]
# Values that must never reach a record, a review item or a CRM payload.
FINANCE = {"纳税人登记号": "TAXSENTINEL0042X", "开户银行": "哨兵银行某支行", "银行账号": "BANKACCOUNT0042",
           "法人": "法人哨兵甲", "传真": "010-55550043", "所属银行名称": "哨兵银行",
           "税率%": 13.0, "币种": "人民币币哨兵", "发货地址": "发货哨兵路1号", "销售默认收付款协议编码": "PAYTERM42"}
# The contact person of an ERP row (owner's decision 2026-10-16): read into the account's contacts.
CONTACT_COLUMNS = {"Email地址": "sentinel.mail@example.invalid", "联系人": "联系哨兵乙", "电话": "010-55550042", "手机": "13900000042"}
MASTER_HEADER = ["Customer Code", "CN name", "EN name", "Province", "Region", "Area", "Previous / Current Supporting",
                 "Dealer", "First Order Time from 2020"]


@pytest.fixture(autouse=True)
def _decisions(tmp_path, monkeypatch):
    monkeypatch.setenv("COMPANY_DECISIONS_FILE", str(tmp_path / "company_decisions.json"))


@pytest.fixture
def master():
    return Master.from_records([])


def book(tmp_path, name, sheets, merges=None):
    """sheets: {title: [row, ...]} written from A1; merges: {title: ["A1:I1", ...]}."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for rng in (merges or {}).get(title, ()):
            ws.merge_cells(rng)
    p = str(tmp_path / name)
    wb.save(p)
    return p


def erp_row(code, name, **over):
    values = {"序号": 1, "客户编码": code, "客户名称": name, **FINANCE, **CONTACT_COLUMNS, **over}
    return [values.get(h) for h in ERP_HEADER]


def erp_book(tmp_path, rows, name="erp.xlsx", title_row=False):
    head = ([["List of customers in the North region"]] if title_row else []) + [ERP_HEADER]
    return book(tmp_path, name, {"North": head + rows})


def master_book(tmp_path, rows, name="master.xlsx"):
    return book(tmp_path, name, {"客户": [["Customer List"], MASTER_HEADER] + rows})


OLD_HEAD = ["序号\nno.", "客户名称\ncustomer name", "English name", "地址\nadress", "联系人\nContact Name", "电话\nphone",
            "主要应用型号\nMain application models", "竞争对手\ncompetitor", "预测全年采购数量\nExpected usage quantity per year"]
COMP_HEAD = ["序号\nno.", "厂家名称\nManufacturer's name", "English name", "产地\nadress", "销售代表\nsales representative",
             "电话\nphone", "1.已建立联系\nContact has been established.\n2.两家国产\nTwo domestic", None, None]
ATT_HEAD = ["序号\nno.", "厂家名称\nManufacturer's name", "English name", "产地\nadress", "销售代表\nsales representative",
            "电话\nphone", "可能需求转台型号\nPossible requirement for turntable model", "所属区域\nregion", None]


def effective_book(tmp_path, name="regional.xlsx"):
    rows = [
        ["意向老客户名单\nList of Intended Old Customers in the Western Region"],
        OLD_HEAD,
        [1, "测试甲机电有限公司", "ceshijia", "湖北省", "张三丰", 13800000001.0, "1.ER15\n2.ER14", "Brand X\nBrand Y", 10.0],
        [2, "测试乙机床有限公司", None, "陕西省", None, None, None, None, "many"],
        [None] * 9,
        ["新客户名单\nList of new customers"],
        OLD_HEAD,
        [1, "测试丙科技", "ceshibing", "西安xi'an", "李四", 13800000002.0, "ER series", "莱曼", 5.0],
        [None] * 9,
        ["展会转台竞争对手\nCompetitors of the rotating stage at the exhibition"],
        COMP_HEAD,
        [1, "测试丁精密机械有限公司", "ceshiding", "山东shandong", "王五", 13800000003.0, None, None, None],
        [2, "测试戊转台", "ceshiwu", "成都chengdu", None, None, None, None, None],
        [None] * 9,
        ["参会机床厂家\nMachine tool manufacturers participating in the event"],
        ATT_HEAD,
        [1, "测试己数控有限公司", "ceshiji", "银川yichuan", "赵六", 13800000004.0, "ER", "西北区域\nNorthwest Region", None],
        [2, "测试甲机电有限公司", "ceshijia", "武汉wuhan", "张三丰", 13800000001.0, "ORANGE NC", "西北区域\nNorthwest Region", None],
    ]
    return book(tmp_path, name, {"Effective": rows}, {"Effective": [
        "A1:I1", "A6:I6", "A10:I10", "G11:I13", "A15:I15", "H17:I18"]})


# ---- cell parsers ----

def test_parse_code_strips_zeros_and_ignores_placeholders():
    assert accounts.parse_code("00CR991038") == ("CR991038", False)
    assert accounts.parse_code("CR0002070") == ("CR0002070", False)
    assert accounts.parse_code("New") == ("", False)
    assert accounts.parse_code("广西某某有限公司CR0001234") == ("CR0001234", True)
    assert accounts.parse_code(None) == ("", False)


def test_parse_first_order_keeps_only_what_the_number_says():
    assert accounts.parse_first_order(2021.12) == "2021-12"
    assert accounts.parse_first_order(2022.9) == "2022-09"
    assert accounts.parse_first_order(2020.1) == "2020"  # January or October
    assert accounts.parse_first_order("2023.05") == "2023-05"
    assert accounts.parse_first_order(2020) == "2020"
    assert accounts.parse_first_order("soon") == ""
    assert accounts.parse_first_order(None) == ""


def test_parse_forecast_number_or_text():
    assert accounts.parse_forecast(10.0) == (10.0, "")
    assert accounts.parse_forecast("5台") == (5.0, "")
    assert accounts.parse_forecast("many") == (None, "many")
    assert accounts.parse_forecast(None) == (None, "")
    assert accounts.parse_forecast("-") == (None, "")


def test_latin_tail_of_bilingual_cells():
    assert accounts.latin_tail("机械行业 Machinery") == "Machinery"
    assert accounts.latin_tail("FIBRO经销商FIBRO Dealer") == "FIBRO Dealer"
    assert accounts.latin_tail("铸造业") == ""
    assert accounts.latin_tail("Machinery") == "Machinery"


# ---- ERP lists (folder 04) ----

def test_erp_reads_the_skeleton(tmp_path, master):
    p = erp_book(tmp_path, [erp_row("00CR100001", "测试甲机电有限公司", 客户简称="测试甲", 助记码="Ceshi Jia Electromechanical",
                                    所属行业="机械行业 Machinery", 地址="北京市东城区测试路1号", 地区名称="北京市Beijing",
                                    省直辖市="北京市", 专营业务员名称="Richard.Yu",
                                    **{"Purpose_Topics Visit": "Dealer of rotary tables."})], title_row=True)
    r = accounts.read_paths([p], master)
    [a] = r.accounts
    assert (a.code, a.cn, a.short, a.en) == ("CR100001", "测试甲机电有限公司", "测试甲", "Ceshi Jia Electromechanical")
    assert (a.industry, a.address, a.province, a.city, a.erp_rep) == \
        ("Machinery", "北京市东城区测试路1号", "北京市", "北京市 Beijing", "Richard Yu")
    assert a.lifecycle == "customer" and a.relationships == ["customer"] and a.source_id == "account:CR100001"
    assert any("Dealer of rotary tables" in n for n in a.notes)
    assert accounts.summarize(r)["accounts"] == 1


def test_erp_rows_without_code_or_name_are_skipped_with_a_reason(tmp_path, master):
    p = erp_book(tmp_path, [erp_row("CR100001", "测试甲机电有限公司"), erp_row("", "测试无码公司"), erp_row("CR100002", None),
                            [None] * len(ERP_HEADER), erp_row("测试脏格CR100003", "测试脏格有限公司")])
    r = accounts.read_paths([p], master)
    assert sorted(a.code for a in r.accounts) == ["CR100001", "CR100003"]
    reasons = {i.reason for i in r.review}
    assert {"no_code", "no_name", "code_cell_has_text"} <= reasons
    assert accounts.summarize(r)["blank_rows"] == 1


def test_erp_do_not_use_is_flagged_and_not_a_customer(tmp_path, master):
    p = erp_book(tmp_path, [erp_row("CR100001", "测试甲机电有限公司", 所属行业="Do Not use")])
    [a] = accounts.read_paths([p], master).accounts
    assert a.industry == "" and a.lifecycle == "" and a.relationships == []
    assert any("Do Not use" in n for n in a.notes)


def test_do_not_use_wins_over_another_list_that_says_customer(tmp_path, master):
    erp = erp_book(tmp_path, [erp_row("CR100001", "测试甲机电有限公司", 所属行业="Do Not use")])
    sales = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", None, None, None, None, None, None, 2021.05]])
    [a] = accounts.read_paths([erp, sales], master).accounts
    assert a.lifecycle == "" and a.relationships == []


def test_finance_columns_never_reach_a_record(tmp_path, master):
    p = erp_book(tmp_path, [erp_row("CR100001", "测试甲机电有限公司"), erp_row("CR100002", "测试乙机床有限公司")])
    r = accounts.read_paths([p], master)
    assert len(r.accounts) == 2
    dumped = repr(r) + json.dumps([vars(i) for i in r.review], ensure_ascii=False, default=str)
    for column, value in FINANCE.items():
        assert str(value) not in dumped, column
    names = {f for cls in (accounts.Row, accounts.Account, accounts.Lead) for f in cls.__dataclass_fields__}
    assert not names & {"tax_id", "bank", "bank_account", "legal_person", "fax", "tax_rate"}


def test_tax_number_typed_into_a_name_column_is_dropped(tmp_path, master):
    tax = "91320000MA1ABCD23X"
    p = erp_book(tmp_path, [erp_row("CR100001", "测试甲机电有限公司", 助记码=tax)])
    q = master_book(tmp_path, [["CR100002", "测试乙机床有限公司", tax, None, None, None, None, None, None]])
    r = accounts.read_paths([p, q], master)
    assert {a.code for a in r.accounts} == {"CR100001", "CR100002"}
    assert tax not in repr(r) and all(a.en == "" for a in r.accounts)
    assert [(i.field, i.value, i.reason) for i in r.review] == [("en", "", "id_number_in_text_field")]  # the ERP mnemonic is not even taken: capitals only
    crm = FakeCRM()
    am.commit(crm, r)
    assert tax not in json.dumps(crm.writes, ensure_ascii=False, default=str)


# ---- Customer_List master (folder 05) ----

def test_master_list_fields_and_new_customers(tmp_path, master):
    p = master_book(tmp_path, [
        ["CR0002070", "测试甲机电有限公司", "Ceshi Jia Co., Ltd.", "上海市Shanghai", "East", "Motion", "Monica", "Oubori existing", 2021.12],
        ["New", "测试新客户有限公司", None, "江苏省Jiangsu", "East", None, None, "new", None],
        [None] * 9,
    ])
    r = accounts.read_paths([p], master)
    a = next(x for x in r.accounts if x.code)
    assert (a.code, a.en, a.region, a.product_area, a.supporting, a.dealer, a.first_order, a.province) == \
        ("CR0002070", "Ceshi Jia Co., Ltd.", "East", "Motion", "Monica", "Oubori existing", "2021-12", "上海市 Shanghai")
    assert a.lifecycle == "customer"  # it has a first order
    new = next(x for x in r.accounts if not x.code)
    assert new.cn == "测试新客户有限公司" and new.lifecycle == "" and new.source_id.startswith("account:")
    assert new.key.startswith("name:")


def test_same_code_twice_in_one_list_is_one_account(tmp_path, master):
    rows = [["CR1000001", "测试甲机电有限公司", None, None, "East", None, None, None, None],
            ["CR1000001", "测试甲机电有限公司", "Ceshi Jia", "上海市Shanghai", "East", "Cutting", "Carrie", None, None]]
    [a] = accounts.read_paths([master_book(tmp_path, rows)], master).accounts
    assert (a.en, a.product_area, a.supporting) == ("Ceshi Jia", "Cutting", "Carrie")


# ---- handover list and registration form ----

def test_supporting_list_and_registration_form(tmp_path, master):
    sup = book(tmp_path, "handover.xlsx", {"Supporting": [
        ["序号\nItem", "客户\nCustomer ", "Company English Name", "City", "Province", "行业\nIndustry segment",
         "FIBRO Tables Category", "FIBRO Existing Customer (Y/N)", "Active Status/Order Placed in 2024-2025 (Y/N)",
         "have been supported by Wang Yao or Liu Hao ", "U8 Asigned to (Previously supported by)"],
        [1, "测试庚机械", "Ceshi Geng Machinery", "Langfang", "Hebei", "Machinery", "Rotocutting", "Y", "Y", "Wang Yao", "Richard Yu"],
        [2, "测试辛设备", None, "-", "-", "Dealer", "Rotomotion", "N", "Pls fill in", None, None],
    ]})
    reg = book(tmp_path, "registration.xlsx", {"customer registration": [
        ["Customer Name客户名称", "Location地点", "Industry所属行业", "Application应用", "Current Brand/Type in use 现使用品牌: ",
         "Division分度(工位数)", "Indexing speed旋转速度", "Accuracy精度要求", "MTM最大惯量"],
        ["Test Registered Technology Co., Ltd.", "Shenyang 沈阳", "Automation", "Machining parts", "no", "4 division", "no demand", None, None],
    ]})
    r = accounts.read_paths([sup, reg], master)
    a = next(x for x in r.accounts if x.cn == "测试庚机械")
    assert (a.erp_rep, a.supporting, a.product_area, a.city, a.province, a.lifecycle) == \
        ("Richard Yu", "Wang Yao", "Rotocutting", "Langfang", "Hebei", "customer")
    assert any("Order placed 2024-2025: Y" in n for n in a.notes)
    b = next(x for x in r.accounts if x.cn == "测试辛设备")
    assert b.lifecycle == "" and b.city == "" and not b.notes
    [lead] = r.leads
    assert (lead.company, lead.origin, lead.location, lead.application) == \
        ("Test Registered Technology Co., Ltd.", "registration", "Shenyang 沈阳", "Machining parts")
    assert "Division: 4 division" in lead.details and lead.source_id.startswith("lead:")
    assert not any(x.en == "Test Registered Technology Co., Ltd." for x in r.accounts)


# ---- the regional analysis workbook (several lists on one sheet) ----

def test_effective_workbook_sections(tmp_path, master):
    r = accounts.read_paths([effective_book(tmp_path)], master)
    by = {a.cn: a for a in r.accounts}
    a = by["测试甲机电有限公司"]
    assert (a.application, a.competitor, a.forecast, a.lifecycle) == ("1.ER15 / 2.ER14", "Brand X / Brand Y", 10.0, "customer")
    assert a.province == "湖北省" and a.en == ""  # "ceshijia" is a pinyin tag, not an English name
    assert a.contacts == [("张三丰", "13800000001")]
    assert any("Possible requirement: ORANGE NC" in n for n in a.notes)  # also attended: a note, not a second record
    b = by["测试乙机床有限公司"]
    assert b.forecast is None and any(i.reason == "forecast_not_a_number" for i in r.review)
    c = by["测试丙科技"]
    assert (c.lifecycle, c.forecast, c.city) == ("prospect", 5.0, "西安 xi'an")
    for name in ("测试丁精密机械有限公司", "测试戊转台"):
        assert by[name].relationships == ["competitor"] and by[name].lifecycle == ""
    assert any("Competitors of the rotating stage" in n for n in by["测试丁精密机械有限公司"].notes)
    assert [l.company for l in r.leads] == ["测试己数控有限公司"]
    lead = r.leads[0]
    assert (lead.origin, lead.contact, lead.application, lead.location) == ("exhibition", "赵六", "ER", "银川 yichuan")
    assert "Northwest Region" in lead.details
    s = accounts.summarize(r)
    assert (s["accounts"], s["competitors"], s["leads"], s["with_forecast"]) == (5, 2, 1, 2)
    stat = r.per_file_stats[0]
    assert (stat.layout, stat.account_rows, stat.competitor_rows, stat.lead_rows) == ("effective", 3, 2, 2)


def test_xls_layout_is_read_like_xlsx(tmp_path, master):
    xlwt = pytest.importorskip("xlwt")
    wb = xlwt.Workbook()
    ws = wb.add_sheet("List")
    for c, h in enumerate(["序号", "客户编码", "客户名称", "客户简称", "所属行业", "专营业务员名称"]):
        ws.write(0, c, h)
    for c, v in enumerate([1.0, "00CR100009", "测试壬有限公司", "壬", "汽车行业Automotive", "Mike.Gao"]):
        ws.write(1, c, v)
    p = str(tmp_path / "list.xls")
    wb.save(p)
    [a] = accounts.read_paths([p], master).accounts
    assert (a.code, a.industry, a.erp_rep) == ("CR100009", "Automotive", "Mike Gao")


# ---- merging across files ----

def test_same_customer_in_two_files_merges_on_code_and_flags_conflicts(tmp_path, master):
    erp = erp_book(tmp_path, [erp_row("00CR100001", "测试甲机电有限公司", 省直辖市="江苏省 Jiangsu", 地址="测试路1号",
                                      专营业务员名称="Mike.Gao")])
    lst = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", "Ceshi Jia Co., Ltd.", "浙江省Zhejiang", "East", "Motion",
                                  "Carrie", None, 2021.12]])
    r = accounts.read_paths([erp, lst], master)
    [a] = r.accounts
    assert (a.code, a.erp_rep, a.en, a.region, a.first_order, a.address) == \
        ("CR100001", "Mike Gao", "Ceshi Jia Co., Ltd.", "East", "2021-12", "测试路1号")
    assert a.province in ("江苏省 Jiangsu", "浙江省 Zhejiang")
    conflicts = [i for i in r.review if i.reason == "conflict"]
    assert [i.field for i in conflicts] == ["province"]
    assert conflicts[0].candidates  # names the value that won
    assert {i.reason for i in r.review} <= {"conflict"}
    assert len(a.sources) == 2


def test_richer_copy_wins_and_blank_copy_creates_nothing(tmp_path, master):
    full = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", "Ceshi Jia", "上海市Shanghai", "East", "Cutting", "Carrie", "new", 2022.4]],
                       "aaaaaaaaaaaa_Customer_List.xlsx")
    thin = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", None, None, None, None, None, None, None], [None] * 9],
                       "bbbbbbbbbbbb_Customer_List.xlsx")
    r = accounts.read_paths([thin, full], master)
    [a] = r.accounts
    assert (a.en, a.region, a.dealer) == ("Ceshi Jia", "East", "new") and not r.review
    assert r.accepted_accounts() == [a]


def test_byte_identical_copies_are_read_once(tmp_path, master):
    p = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", None, None, None, None, None, None, None]])
    copy = str(tmp_path / "copy_of_master.xlsx")
    shutil.copy(p, copy)
    r = accounts.read_paths([p, copy], master)
    assert len(r.per_file_stats) == 1 and len(r.accounts) == 1
    assert len(r.skipped_files) == 1 and "identical" in r.skipped_files[0]["error"]


def test_row_without_code_joins_the_account_the_master_knows(tmp_path):
    known = Master.from_records([{"source": "master", "code": "CR100001", "cn": "测试甲机电有限公司",
                                  "en": "Ceshi Jia Electromechanical Co., Ltd.", "short": ""}])
    sup = book(tmp_path, "handover.xlsx", {"S": [
        ["客户\nCustomer ", "Company English Name", "FIBRO Existing Customer (Y/N)"],
        ["测试甲机电", "Ceshi Jia Electromechanical", "Y"]]})
    [a] = accounts.read_paths([sup], known).accounts
    assert a.code == "CR100001" and a.key == "erp:CR100001" and a.cn == "测试甲机电有限公司"
    assert "测试甲机电" in a.names and a.lifecycle == "customer"


def test_crm_company_without_code_gets_no_code_from_its_id(tmp_path):
    """The CRM master gives a company without an ERP code the stand-in crm:<id>; the id's tail looks like a code
    (letters then digits) and must not be written as one."""
    known = Master.from_records([{"source": "crm", "code": "crm:01a142fd-6e8f-7331-a331-eeb928935537",
                                  "cn": "测试乙机电有限公司", "en": "", "short": ""}])
    sup = book(tmp_path, "handover.xlsx", {"S": [
        ["客户\nCustomer ", "Company English Name", "FIBRO Existing Customer (Y/N)"],
        ["测试乙机电", "", "Y"]]})
    accs = accounts.read_paths([sup], known).accounts
    assert [a.code for a in accs] == [""]


def test_doubtful_name_is_queued_not_merged(tmp_path):
    # same name, other province: doubtful. (The same province is accepted since the place check.)
    known = Master.from_records([{"source": "master", "code": "CR100001", "cn": "曼弗莱德智能制造（江苏）有限公司",
                                  "en": "", "short": ""}])
    lst = master_book(tmp_path, [["New", "曼弗莱德智能制造（山东）有限公司", None, None, None, None, None, "new", None]])
    r = accounts.read_paths([lst], known)
    [a] = r.accounts
    assert a.queued and r.accepted_accounts() == []
    assert [i.reason for i in r.review] == ["company"]
    assert accounts.summarize(r)["queued_by_reason"] == {"company": 1}


def test_unreadable_file_is_recorded_and_the_rest_is_read(tmp_path, master):
    bad = tmp_path / "broken.xlsx"
    bad.write_text("not a workbook")
    good = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", None, None, None, None, None, None, None]])
    r = accounts.read_paths([str(bad), good], master)
    assert len(r.accounts) == 1
    assert [s["file"] for s in r.skipped_files] == [str(bad)] and "could not be read" in r.skipped_files[0]["error"]


def test_file_kind_tells_customer_lists_from_other_workbooks(tmp_path):
    lst = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", None, None, None, None, None, None, None]])
    other = book(tmp_path, "other.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})
    (tmp_path / "junk.xlsx").write_text("no")
    assert accounts.file_kind(lst) == ("accounts", "master")
    assert accounts.file_kind(effective_book(tmp_path)) == ("accounts", "effective")
    assert accounts.file_kind(other)[0] == "unrecognised"
    assert accounts.file_kind(str(tmp_path / "junk.xlsx"))[0] == "unrecognised"


def test_a_sheet_that_is_no_customer_list_is_ignored(tmp_path, master):
    p = book(tmp_path, "other.xlsx", {"Sheet1": [["a", "b"], [1, 2]]})
    r = accounts.read_paths([p], master)
    assert r.accounts == [] and r.leads == [] and r.per_file_stats[0].layout == "none"


# ---- the CRM side, against a fake ----

class FakeCRM:
    """Just enough of common.client.MargenceClient for excel.accounts_margince. Keeps every write."""

    def __init__(self):
        self.ids = itertools.count(1)
        self.companies, self.contacts, self.leads, self.activities = {}, {}, {}, {}
        self.employment, self.fields, self.writes = [], {"company": [], "lead": []}, []
        self.fail_company = set()

    def _id(self, prefix):
        return f"{prefix}{next(self.ids)}"

    def add_company(self, name, **fields):
        cid = self._id("co")
        self.companies[cid] = {"id": cid, "display_name": name, "lifecycle": "unknown", "relationship_types": [], **fields}
        return cid

    def add_field(self, object_, label, type_="text"):
        f = {"label": label, "type": type_, "column_name": "cf_" + label.lower().replace(" ", "_")}
        self.fields[object_].append(f)
        return f

    # vocabulary
    def custom_fields(self, object_):
        return list(self.fields[object_])

    def create_custom_field(self, body):
        self.writes.append(("field", body))
        return self.add_field(body["object"], body["label"], body["type"])

    def seats_by_name(self):
        return {"Richard Yu": "seat-richard"}

    # reads
    def paginate(self, path, params=None):
        store = {"/v1/companies": self.companies, "/v1/leads": self.leads, "/v1/activities": self.activities}[path]
        yield from (dict(v) for v in list(store.values()))

    def contacts_for_company(self, company_id):
        return [{"full_name": self.contacts[c]["full_name"], "contact_id": c} for co, c in self.employment if co == company_id]

    # writes
    def post(self, path, json=None, **kw):
        assert path == "/v1/companies"
        self.writes.append(("company", json))
        if json["display_name"] in self.fail_company:
            raise _problem(500)
        cid = self._id("co")
        self.companies[cid] = {"lifecycle": "unknown", "relationship_types": [], **json, "id": cid}
        return dict(self.companies[cid])

    def patch(self, path, json, **kw):
        kind, ident = path.strip("/").split("/")[1:3]
        store = {"companies": self.companies, "leads": self.leads}[kind]
        self.writes.append((kind + "-patch", json))
        store[ident].update(json)
        return dict(store[ident])

    def post_status(self, path, json=None, **kw):
        assert path == "/v1/leads"
        self.writes.append(("lead", json))
        for l in self.leads.values():
            if (l["source_system"], l["source_id"]) == (json["source_system"], json["source_id"]):
                return 200, dict(l)
        lid = self._id("ld")
        self.leads[lid] = {**json, "id": lid}
        return 201, dict(self.leads[lid])

    def create_contact(self, body):
        self.writes.append(("contact", body))
        cid = self._id("ct")
        self.contacts[cid] = {**body, "id": cid}
        return dict(self.contacts[cid])

    def link_employment(self, contact_id, company_id):
        self.employment.append((company_id, contact_id))
        return {}

    def log_activity(self, body):
        for a in self.activities.values():
            if (a["source_system"], a["source_id"]) == (body["source_system"], body["source_id"]):
                return False, dict(a)
        self.writes.append(("note", body))
        aid = self._id("ac")
        self.activities[aid] = {**body, "id": aid}
        return True, dict(self.activities[aid])

    def n(self, kind):
        return sum(1 for k, _ in self.writes if k == kind)


def _problem(status):
    import requests
    resp = requests.Response()
    resp.status_code = status
    resp._content = b'{"title": "boom"}'
    resp.request = requests.Request("POST", "http://fake/v1/companies").prepare()
    return ApiError(resp)


def crm_read(tmp_path, master):
    """A read with every kind of record: ERP account, list account, opportunity data, a competitor, a lead, notes."""
    erp = erp_book(tmp_path, [erp_row("00CR100001", "测试甲机电有限公司", 所属行业="机械行业 Machinery", 地址="测试路1号",
                                      省直辖市="江苏省", 专营业务员名称="Richard.Yu", 地区名称="苏州市Suzhou",
                                      **{"Purpose_Topics Visit": "Dealer of rotary tables."})])
    lst = master_book(tmp_path, [["CR100001", "测试甲机电有限公司", "Ceshi Jia Co., Ltd.", None, "East", "Motion", "Carrie", None, 2021.12]])
    return accounts.read_paths([erp, lst, effective_book(tmp_path)], master)


def test_commit_writes_accounts_contacts_notes_and_leads(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    report = am.commit(crm, r)
    assert report["failed"] == [] and report["conflicts"] == []
    s = accounts.summarize(r)
    assert report["companies_created"] == s["accounts"] == 5 and len(crm.companies) == 5
    cols = {f["label"]: f["column_name"] for f in crm.fields["company"]}
    jia = next(c for c in crm.companies.values() if c["display_name"] == "测试甲机电有限公司")
    assert jia[cols["ERP customer code"]] == "CR100001" and jia[cols["Sales region"]] == "East"
    assert jia[cols["Competitor"]] == "Brand X / Brand Y" and jia[cols["Expected annual volume"]] == 10
    assert jia[cols["Application"]] == "1.ER15 / 2.ER14" and jia[cols["First order"]] == "2021-12"
    assert (jia["lifecycle"], jia["relationship_types"], jia["industry"], jia["legal_name"]) == \
        ("customer", ["customer"], "Machinery", "Ceshi Jia Co., Ltd.")
    assert jia["address"] == {"line1": "测试路1号", "city": "苏州市 Suzhou", "region": "江苏省"}
    assert jia["owner_id"] == "seat-richard"
    comp = next(c for c in crm.companies.values() if c["display_name"] == "测试丁精密机械有限公司")
    assert comp["relationship_types"] == ["competitor"] and comp["lifecycle"] == "unknown"
    assert next(c for c in crm.companies.values() if c["display_name"] == "测试丙科技")["lifecycle"] == "prospect"
    assert [c["full_name"] for c in crm.contacts.values()].count("张三丰") == 1
    phones = next(c for c in crm.contacts.values() if c["full_name"] == "张三丰")["phones"]
    assert phones == [{"phone": "+8613800000001", "phone_type": "mobile", "is_primary": True, "position": 0}]
    assert report["notes_created"] == s["notes"] == len(crm.activities)
    assert all(a["links"][0]["entity_id"] in crm.companies for a in crm.activities.values())
    [lead] = crm.leads.values()
    assert (lead["company_name"], lead["full_name"], lead["source"]) == ("测试己数控有限公司", "赵六", "import")
    lead_cols = {f["label"]: f["column_name"] for f in crm.fields["lead"]}
    assert lead[lead_cols["Application"]] == "ER" and lead[lead_cols["Phone"]] == "13800000004"
    assert report["leads_created"] == 1 and len(report["fields_created"]) == len(am.COMPANY_FIELDS) + len(am.LEAD_FIELDS)


def test_second_commit_writes_nothing(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    am.commit(crm, r)
    writes = len(crm.writes)
    again = am.commit(crm, accounts.read_paths(
        sorted(glob.glob(str(tmp_path / "*.xlsx"))), master))
    assert len(crm.writes) == writes
    assert (again["companies_created"], again["companies_updated"], again["contacts_created"], again["notes_created"],
            again["leads_created"], again["failed"]) == (0, 0, 0, 0, 0, [])
    assert again["companies_unchanged"] == 5 and again["leads_present"] == 1 and again["contacts_present"] == 4  # + the ERP row's contact
    assert len(crm.companies) == 5 and len(crm.leads) == 1


def test_preview_reads_only_and_counts(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    out = am.preview(crm, r)
    assert crm.writes == [] and not crm.companies
    assert (out["companies_new"], out["companies_present"], out["contacts_new"], out["leads_new"]) == (5, 0, 4, 1)  # 3 of the analysis list + the ERP row's contact
    assert out["notes_new"] == out["notes"] == 4 and len(out["fields_missing"]) == len(am.COMPANY_FIELDS)
    am.commit(crm, r)
    out = am.preview(crm, r)
    assert (out["companies_new"], out["companies_present"], out["companies_to_fill"], out["contacts_new"],
            out["notes_new"], out["leads_new"]) == (0, 5, 0, 0, 0, 0)
    assert out["fields_missing"] == []


def test_account_the_crm_already_holds_is_enriched_not_duplicated(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    held = crm.add_company("测试甲机电有限公司", industry="Foundry", legal_name="Ceshi Jia Co., Ltd.", lifecycle="prospect",
                           address={"line1": "旧路2号"})
    report = am.commit(crm, r)
    assert len(crm.companies) == 5 and report["companies_created"] == 4 and report["companies_updated"] == 1
    cur = crm.companies[held]
    assert cur["industry"] == "Foundry" and cur["lifecycle"] == "prospect"  # a person's value stays
    assert cur["address"]["line1"] == "旧路2号" and cur["address"]["region"] == "江苏省"  # blanks filled
    cols = {f["label"]: f["column_name"] for f in crm.fields["company"]}
    assert cur[cols["ERP customer code"]] == "CR100001"
    fields = {c["field"] for c in report["conflicts"]}
    assert fields == {"Industry", "Address line1"} and all(c["account"] == "account:CR100001" for c in report["conflicts"])


def test_account_is_found_by_erp_code_even_when_the_name_differs(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    code_field = crm.add_field("company", "ERP customer code")
    held = crm.add_company("Totally Other Name", **{code_field["column_name"]: "00CR100001"})
    am.commit(crm, r)
    assert len(crm.companies) == 5 and crm.companies[held]["display_name"] == "Totally Other Name"
    assert crm.companies[held]["industry"] == "Machinery"


def test_company_held_twice_is_left_alone(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    a, b = crm.add_company("测试甲机电有限公司"), crm.add_company("测试甲机电")
    crm.companies[b]["legal_name"] = "测试甲机电有限公司"
    report = am.commit(crm, r)
    assert report["companies_ambiguous"] == ["account:CR100001"]
    assert len(crm.companies) == 6 and crm.companies[a].get("industry") is None


def test_company_with_another_erp_code_is_another_customer(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    code_field = crm.add_field("company", "ERP customer code")
    crm.add_company("测试甲机电有限公司", **{code_field["column_name"]: "CR999999"})
    report = am.commit(crm, r)
    assert report["companies_created"] == 5 and len(crm.companies) == 6


def test_lead_for_a_company_the_crm_holds_is_not_duplicated(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    crm.add_company("测试己数控有限公司")
    report = am.commit(crm, r)
    assert report["leads_created"] == 0 and report["leads_already_accounts"] == 1 and not crm.leads


def test_one_failing_account_does_not_stop_the_run(tmp_path, master):
    r = crm_read(tmp_path, master)
    crm = FakeCRM()
    crm.fail_company = {"测试丙科技"}
    report = am.commit(crm, r)
    assert [f["id"] for f in report["failed"]][0].startswith("account:")
    assert report["companies_created"] == 4 and len(crm.leads) == 1


def test_queued_accounts_and_leads_are_not_written(tmp_path):
    known = Master.from_records([{"source": "master", "code": "CR100001", "cn": "曼弗莱德智能制造（江苏）有限公司", "en": "", "short": ""}])
    lst = master_book(tmp_path, [["New", "曼弗莱德智能制造（山东）有限公司", None, None, None, None, None, "new", None]])
    r = accounts.read_paths([lst], known)
    crm = FakeCRM()
    report = am.commit(crm, r)
    assert crm.writes == [] or all(k == "field" for k, _ in crm.writes)
    assert not crm.companies and report["companies_created"] == 0
    assert am.plan_counts(r)["queued"] == 1


def test_finance_fields_never_reach_the_crm(tmp_path, master):
    erp = erp_book(tmp_path, [erp_row("CR100001", "测试甲机电有限公司", 所属行业="机械行业 Machinery", 地址="测试路1号"),
                              erp_row("CR100002", "测试乙机床有限公司")])
    crm = FakeCRM()
    report = am.commit(crm, accounts.read_paths([erp], master))
    assert report["companies_created"] == 2 and crm.n("company") == 2
    sent = json.dumps(crm.writes + list(crm.companies.values()) + list(crm.contacts.values()), ensure_ascii=False, default=str)
    for column, value in FINANCE.items():
        assert str(value) not in sent, column
    # the contact person of each ERP row is a contact of its account (owner's decision), with all his numbers and his address
    assert len(crm.contacts) == 2 and crm.activities == {}
    for c in crm.contacts.values():
        assert c["full_name"] == "联系哨兵乙"
        assert [p["phone"] for p in c["phones"]] == ["+8613900000042", "+861055550042"]
        assert [e["email"] for e in c["emails"]] == ["sentinel.mail@example.invalid"]


# ---- the real sample folder: counts only ----

FOLDER = os.path.join(SAMPLE_DIR, "04 ERP Customer Lists by Region"), os.path.join(SAMPLE_DIR, "05 Customer Lists")
REAL = [p for d in FOLDER for p in sorted(glob.glob(os.path.join(d, "*")))
        if not os.path.basename(p).startswith((".", "~$"))]
needs_folder = pytest.mark.skipif(len(REAL) < 7, reason="sample folders 04 / 05 not present")


@needs_folder
def test_real_folder_every_file_is_read_and_none_is_empty():
    r = accounts.read_paths(REAL, Master.from_records([]))
    report = {x["file"]: x for x in accounts.file_report(r)}
    assert r.skipped_files == [] and all(x["error"] == "" for x in report.values())
    assert len(report) == 7 and all(x["rows"] > 0 for x in report.values())  # the three that used to give 0 records too
    assert {x["layout"] for x in report.values()} == {"erp", "master", "effective", "supporting", "registration"}
    regional = next(x for n, x in report.items() if n.startswith("Regional"))
    assert (regional["account_rows"], regional["competitor_rows"], regional["lead_rows"]) == (16, 6, 16)
    assert (regional["with_competitor"], regional["with_forecast"]) == (10, 10)
    s = accounts.summarize(r)
    assert s["accounts"] > 1300 and s["leads"] >= 10 and s["with_forecast"] >= 5 and s["queued_by_reason"].get("company", 0) < 60


@needs_folder
def test_real_folder_bank_and_tax_values_do_not_come_through():
    import re
    import xlrd
    r = accounts.read_paths(REAL, Master.from_records([]))
    dumped = repr(r)
    idlike = re.compile(r"[0-9A-Z]{8,}")
    checked, leaked = 0, []
    for p in REAL:
        if not p.endswith(".xls") or "Regional" in p:
            continue
        sh = xlrd.open_workbook(p).sheet_by_index(0)
        header = {str(sh.cell_value(h, c)).strip(): c for h in range(min(3, sh.nrows)) for c in range(sh.ncols)}
        for label in ("银行账号", "纳税人登记号", "开户银行"):
            if label not in header:
                continue
            for row in range(sh.nrows):
                v = str(sh.cell_value(row, header[label])).strip()
                if v == label or (label != "开户银行" and not idlike.fullmatch(v)) or len(v) < 6:
                    continue  # a cell that is not an account or tax number (a name typed into the column) is not checked
                checked += 1
                if v in dumped:
                    leaked.append(label)
    assert checked > 100
    assert leaked == []  # never assert `in` on the dump itself: a failure would try to print a megabyte


@needs_folder
def test_real_folder_commit_to_a_fake_crm_sends_no_id_numbers_or_addresses_of_people():
    import re
    r = accounts.read_paths(REAL, Master.from_records([]))
    crm = FakeCRM()
    report = am.commit(crm, r)
    assert report["failed"] == [] and report["companies_created"] == len(r.accepted_accounts())
    # a note that carries a contact's own details (a number or address with no person's name, an extension) is the one place
    # such text is allowed: contact details are no bank or tax data
    sent = [json.dumps(b, ensure_ascii=False, default=str) for k, b in crm.writes if k in ("company", "company-patch", "note", "lead")
            and not (k == "note" and "Customer list: Contact" in b.get("subject", ""))]
    assert len(sent) > 1300
    odd = [k for k in ("@", "银行", "账号", "纳税") if any(k in s for s in sent)]
    ids = sum(bool(re.search(r"(?<![0-9A-Za-z])(?:\d{10,}|(?=[0-9A-Z]*\d)[0-9A-Z]{15,20})(?![0-9A-Za-z])", s)) for s in sent)
    assert odd == [] and ids == 0
    again = am.commit(crm, r)
    assert (again["companies_created"], again["notes_created"], again["leads_created"], again["contacts_created"]) == (0, 0, 0, 0)
