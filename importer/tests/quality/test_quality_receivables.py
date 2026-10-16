"""Quality checks for the receivables reader (folder 09), run on every row of every AR list and every sheet.

The oracle (oracle_receivables.py) is built from the cell ledger alone. The numbers are the point: every
invoice amount, the open and overdue total of every customer, the grand totals, and the sheets' own control
totals are compared in Decimal to the cent.

THE OVERDUE RULE (documented, and checked against the sheet): the sheet's `Today` cell (A5, `=TODAY()` as last
saved by Excel) is the day the figures are counted from. `Due days` (column H) = Due Date - Today, so an invoice
is overdue when Due days < 0; due today (0) is not overdue. The reader applies "due date < as-of day" with the as-of
day taken from that Today cell; the oracle applies the file's own `Due days` column. Both must agree on every row,
and both must equal the totals the sheet states itself (`Total Amount Overdue`, `SUMIF(H<0)`).

Declared drop rules (a row/sheet that is not in the book must be named here):
  rows:   placeholder rows (`#N/A` customer and `No Invoice`, no amount), zero amounts
  cells:  header row, the rows above it (totals, Today), the side pivot tables beside the list (columns after M)
  sheets: hidden, bank cash flows, accounts payable, fixed costs, finance summaries, customer down payments,
          aging-bucket ledgers (the files named Laepple are kept out by design)
"""
import os
import re
from collections import Counter
from decimal import Decimal

import pytest

from excel import receivables
from tests.quality import fam3_common as oc
from tests.quality import oracle_receivables as oo
from tests.quality import snapshot, snapshot_receivables

KNOWN_BUGS = {}
ALL_RULES = sorted(set(KNOWN_BUGS) | {
    "sheet-kind", "no-hidden-invoice-list", "sheet-facts", "invoice-rows", "invoice-fields", "amount-exact",
    "overdue-rule", "sheet-totals", "control-totals", "book-snapshot", "book-customers", "book-grand-totals",
    "book-audit-detail", "field-shape", "idempotent", "data-consistency",
})
REASON_CATEGORY = (  # reader reason text -> oracle category
    ("hidden", "hidden"), ("bank cash", "bank"), ("accounts payable", "payable"), ("fixed costs", "fixed"),
    ("finance summary", "summary"), ("down payments", "down_payment"), ("aging", "aging"))
ANNOTATION = re.compile(r"\s*[（(][^）)]*\btransfer\b[^）)]*[）)]?\s*$", re.I)  # declared: a "(transfer from ...)" note is not part of the name


def _v(rule, file, sheet, row, detail):
    return {"rule": rule, "file": file, "sheet": sheet, "row": row, "detail": str(detail)}


def _category(reason):
    return next((c for needle, c in REASON_CATEGORY if needle in reason), "other")


def _same(a, b):
    return a is not None and b is not None and oc.cent(Decimal(a)) == oc.cent(Decimal(b))


@pytest.fixture(scope="session")
def facts(sample_dir):
    if not os.path.isdir(sample_dir):
        pytest.skip("sample folder not found")
    return oo.all_facts(sample_dir)


@pytest.fixture(scope="session")
def workbooks(facts):
    out = {}
    for f in facts:
        info = receivables.FileInfo(path=f.path, name=receivables.file_display_name(f.path))
        out[f.path] = (receivables.read_workbook(f.path, info), info)
    return out


@pytest.fixture(scope="session")
def book(facts):
    return receivables.build_book([f.path for f in facts])


def _sheet_checks(facts, workbooks, out):
    for f in facts:
        read, info = workbooks[f.path]
        if set(info.sheets_read) != set(f.ar):
            out.append(_v("sheet-kind", f.file, "", 0, f"reader reads {info.sheets_read}, oracle {list(f.ar)}"))
        got = {n: _category(why) for n, why in info.excluded_sheets}
        if got != f.other:
            diff = {n: (got.get(n), f.other.get(n)) for n in set(got) | set(f.other) if got.get(n) != f.other.get(n)}
            out.append(_v("sheet-kind", f.file, "", 0, f"reader/oracle category differs: {dict(list(diff.items())[:4])}"))
        for hidden in oo.hidden_sheet_has_invoice_list(f.path):
            out.append(_v("no-hidden-invoice-list", f.file, hidden, 0, "a hidden sheet holds an invoice list"))
        by_name = {s.sheet: s for s in read}
        for name, o in f.ar.items():
            s = by_name.get(name)
            if s is None:
                continue
            if s.as_of != o.today:
                out.append(_v("sheet-facts", f.file, name, 0, f"as-of {s.as_of!r}, the Today cell says {o.today!r}"))
            if "today" not in s.as_of_source:
                out.append(_v("sheet-facts", f.file, name, 0, f"as-of taken from {s.as_of_source!r}, not the sheet's Today cell"))
            if s.currency != "CNY":
                out.append(_v("sheet-facts", f.file, name, 0, f"currency {s.currency!r}"))
            if not (_same(s.claimed_open, o.claimed_open) and _same(s.claimed_overdue, o.claimed_overdue)):
                out.append(_v("sheet-facts", f.file, name, 0, f"claimed totals {s.claimed_open}/{s.claimed_overdue} vs {o.claimed_open}/{o.claimed_overdue}"))
            if s.skipped_rows != len(o.dropped):
                out.append(_v("sheet-facts", f.file, name, 0, f"reader skipped {s.skipped_rows} rows, oracle {len(o.dropped)}"))
            _invoice_checks(f, name, o, s, out)


def _invoice_checks(f, name, o, s, out):
    mine = {i.row: i for i in s.invoices}
    want = {r.row: r for r in o.rows}
    for row in sorted(set(want) - set(mine)):
        out.append(_v("invoice-rows", f.file, name, row, f"invoice {want[row].invoice_no or '(no number)'} {want[row].amount} not read"))
    for row in sorted(set(mine) - set(want)):
        out.append(_v("invoice-rows", f.file, name, row, "reader has an invoice the oracle does not"))
    for row in sorted(set(mine) & set(want)):
        i, r = mine[row], want[row]
        if i.amount != r.amount:
            out.append(_v("amount-exact", f.file, name, row, f"{r.invoice_no}: cell {r.amount}, reader {i.amount}"))
        for label, a, b in (("customer_cn", i.customer_cn, r.cn), ("customer_en", i.customer_en, ANNOTATION.sub("", r.en).strip()),
                            ("invoice_no", i.invoice_no or "", r.invoice_no), ("invoice_date", i.invoice_date or "", r.invoice_date),
                            ("due_date", i.due_date or "", r.due_date)):
            if a != b:
                out.append(_v("invoice-fields", f.file, name, row, f"{label}: cell {b!r}, reader {a!r}"))
        reader_overdue = bool(i.due_date and s.as_of and i.due_date < s.as_of)
        if reader_overdue != o.overdue(r):
            out.append(_v("overdue-rule", f.file, name, row, f"{r.invoice_no}: due {r.due_date}, Due days {r.due_days}: "
                          f"file says overdue={o.overdue(r)}, reader says {reader_overdue}"))
    if s.open_total != o.open_total or s.overdue_total != o.overdue_total:
        out.append(_v("sheet-totals", f.file, name, 0, f"reader open/overdue {s.open_total}/{s.overdue_total}, rows give {o.open_total}/{o.overdue_total}"))


def _data_checks(facts, out):
    """Properties of the DATA that make the oracle trustworthy: the file's own control totals agree with its rows."""
    import datetime as dt
    for f in facts:
        for name, o in f.ar.items():
            if o.claimed_open is not None and oc.cent(o.open_total) != oc.cent(o.claimed_open):
                out.append(_v("control-totals", f.file, name, 0, f"Total Account Receivables {o.claimed_open}, rows sum to {o.open_total}"))
            if o.claimed_overdue is not None and oc.cent(o.overdue_total) != oc.cent(o.claimed_overdue):
                out.append(_v("control-totals", f.file, name, 0, f"Total Amount Overdue {o.claimed_overdue}, rows with Due days < 0 sum to {o.overdue_total}"))
            today = dt.date.fromisoformat(o.today)
            for r in o.rows:
                if r.due_date and r.due_days is not None and (dt.date.fromisoformat(r.due_date) - today).days != r.due_days:
                    out.append(_v("data-consistency", f.file, name, r.row, f"Due days {r.due_days} != {r.due_date} - {o.today}"))
                if r.amount <= 0:
                    out.append(_v("data-consistency", f.file, name, r.row, f"non-positive amount {r.amount}"))
            cust = o.customers()
            for pivot_name, total in o.pivot.items():
                if pivot_name not in cust or oc.cent(cust[pivot_name]["open"]) != oc.cent(total):
                    out.append(_v("control-totals", f.file, name, 0, f"the file's own pivot says {pivot_name}: {total}, rows give {cust.get(pivot_name, {}).get('open')}"))


def _book_checks(facts, book, out):
    candidates = [(o.today, f, o) for f in facts for o in f.ar.values()]
    newest = max(c[0] for c in candidates)
    top = [c for c in candidates if c[0] == newest]
    f, o = top[0][1], top[0][2]
    if book.as_of != newest or book.source != f.display:
        out.append(_v("book-snapshot", book.source, "", 0, f"book is {book.source!r} as of {book.as_of}, the newest list is {f.display!r} as of {newest}"))
    for info in book.files:
        want_old = any(c[0] < newest and oc.display_name(c[1].path) == info.name for c in candidates)
        if want_old != bool(info.superseded_by):
            out.append(_v("book-snapshot", info.name, "", 0, f"superseded_by={info.superseded_by!r}, should be {want_old}"))
    expected = o.customers()
    by_written = {}
    for c in book.customers.values():
        for n in c.written_names:
            by_written.setdefault(n, c)
    grouped = {}
    for name, e in expected.items():
        cust = by_written.get(name)
        if cust is None:
            out.append(_v("book-customers", "", "", 0, f"customer {name!r} ({e['open']}) is not in the book"))
            continue
        g = grouped.setdefault(cust.key, {"cust": cust, "open": Decimal(0), "overdue": Decimal(0), "invoices": 0,
                                          "overdue_invoices": 0, "oldest_due": ""})
        g["open"] += e["open"]
        g["overdue"] += e["overdue"]
        g["invoices"] += e["invoices"]
        g["overdue_invoices"] += e["overdue_invoices"]
        if e["oldest_due"] and (not g["oldest_due"] or e["oldest_due"] < g["oldest_due"]):
            g["oldest_due"] = e["oldest_due"]
    for key in set(book.customers) - set(grouped):
        out.append(_v("book-customers", "", "", 0, f"book customer {key!r} has no row in the list"))
    for key, g in grouped.items():
        c = g["cust"]
        if c.open != g["open"] or c.overdue != g["overdue"]:
            out.append(_v("book-customers", "", "", 0, f"{c.name}: open {c.open} / overdue {c.overdue}, rows give {g['open']} / {g['overdue']}"))
        if (c.invoices, c.overdue_invoices) != (g["invoices"], g["overdue_invoices"]):
            out.append(_v("book-audit-detail", "", "", 0, f"{c.name}: invoices {c.invoices}/{c.overdue_invoices}, rows give {g['invoices']}/{g['overdue_invoices']}"))
        if (c.oldest_due or "") != g["oldest_due"]:
            out.append(_v("book-audit-detail", "", "", 0, f"{c.name}: oldest due {c.oldest_due!r}, rows give {g['oldest_due']!r}"))
    total_open = sum((c.open for c in book.customers.values()), Decimal(0))
    total_overdue = sum((c.overdue for c in book.customers.values()), Decimal(0))
    if total_open != o.open_total or total_open != o.claimed_open:
        out.append(_v("book-grand-totals", "", "", 0, f"customers' open {total_open}; rows {o.open_total}; the sheet states {o.claimed_open}"))
    if total_overdue != o.overdue_total or total_overdue != o.claimed_overdue:
        out.append(_v("book-grand-totals", "", "", 0, f"customers' overdue {total_overdue}; rows {o.overdue_total}; the sheet states {o.claimed_overdue}"))
    summary = receivables.summarize(book)
    if Decimal(summary["open_total"]) != total_open or Decimal(summary["overdue_total"]) != total_overdue:
        out.append(_v("book-grand-totals", "", "", 0, f"summarize() says {summary['open_total']} / {summary['overdue_total']}"))


def _shape_checks(book, out):
    if not re.fullmatch(r"\d{4}-\d{2}-\d{2}", book.as_of or ""):
        out.append(_v("field-shape", "", "", 0, f"as_of {book.as_of!r}"))
    if book.currency != "CNY":
        out.append(_v("field-shape", "", "", 0, f"currency {book.currency!r}"))
    for c in book.customers.values():
        for k in ("open", "overdue"):
            v = getattr(c, k)
            if not isinstance(v, Decimal) or v != v.quantize(Decimal("0.01")) or v < 0:
                out.append(_v("field-shape", "", "", 0, f"{c.name}: {k} {v!r} is not a non-negative amount in whole fen"))
        if c.overdue > c.open:
            out.append(_v("field-shape", "", "", 0, f"{c.name}: overdue {c.overdue} exceeds open {c.open}"))
        if not c.name or not c.written_names:
            out.append(_v("field-shape", "", "", 0, f"customer without a name: {c.key!r}"))
        if c.oldest_due and not re.fullmatch(r"\d{4}-\d{2}-\d{2}", c.oldest_due):
            out.append(_v("field-shape", "", "", 0, f"{c.name}: oldest due {c.oldest_due!r}"))


def _idempotent(facts, out):
    paths = [f.path for f in facts]
    a, b = snapshot_receivables.records_of(paths), snapshot_receivables.records_of(paths)
    if a != b:
        out.append(_v("idempotent", "", "", 0, snapshot.report(snapshot.diff(a, b))))


@pytest.fixture(scope="session")
def violations(facts, workbooks, book):
    out = []
    _sheet_checks(facts, workbooks, out)
    _data_checks(facts, out)
    _book_checks(facts, book, out)
    _shape_checks(book, out)
    _idempotent(facts, out)
    return out


def _param(rule):
    if rule in KNOWN_BUGS:
        return pytest.param(rule, marks=pytest.mark.xfail(strict=True, reason=KNOWN_BUGS[rule]))
    return rule


@pytest.mark.parametrize("rule", [_param(r) for r in ALL_RULES])
def test_rule_has_no_violations(rule, violations):
    bad = [v for v in violations if v["rule"] == rule]
    assert not bad, f"{len(bad)} violation(s) of '{rule}':\n" + "\n".join(
        f"  {v['file'][:40]} {v['sheet'][:12]} row {v['row']}: {v['detail'][:200]}" for v in bad[:6])


def test_population(facts, book):
    """Everything the rules ran over: 4 workbooks, 2 AR lists with 48 invoices each, 14 customers in the book."""
    assert len(facts) == 4
    assert sorted(len(o.rows) for f in facts for o in f.ar.values()) == [48, 48]
    assert sum(len(f.other) for f in facts) == 85 + 16 + 15 + 4  # every sheet that is not an AR list has a kind (85 + 16 + 15 + 4 sheets)
    assert len(book.customers) == 14


def test_duplicate_invoice_numbers_are_counted_once_per_row(facts, book):
    """SIN2602030 is typed on two rows of the 20260410 list (27 Feb and 27 Mar, 34,675 each). The sheet's own total
    counts both, so the book does too (decisions/receivables.md, RECV-2)."""
    o = next(o for f in facts for o in f.ar.values() if o.today == book.as_of)
    typed = [r for r in o.rows if r.invoice_no and "invoice_no" not in r.merged_from]
    dups = [n for n, k in Counter(r.invoice_no for r in typed).items() if k > 1]
    assert dups == ["SIN2602030"]
    assert sum(r.amount for r in o.rows if r.invoice_no == "SIN2602030") == Decimal("69350")


def test_one_invoice_over_two_amount_rows(facts, book):
    """SIN2603025 (Changzhou Miyou) is one merged invoice cell over two amount rows, 141,500 + 169,800; both count and the
    sheet's own total agrees (decisions/receivables.md, RECV-3)."""
    o = next(o for f in facts for o in f.ar.values() if o.today == book.as_of)
    rows = [r for r in o.rows if r.invoice_no == "SIN2603025"]
    assert [(r.row, str(r.amount), bool(r.merged_from)) for r in rows] == [(47, "141500", False), (48, "169800", True)]
    assert {r.due_date for r in rows} == {"2026-06-25"}


def test_due_on_the_as_of_day_is_not_overdue(book):
    """The real lists have no invoice due exactly on the Today cell, so the boundary is pinned on a synthetic list."""
    import datetime as dt
    import tempfile

    import openpyxl
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "AR"
    ws["A5"], ws["B5"] = dt.datetime(2026, 4, 13), "Today !"
    heads = ["Customer name(EN)", "Customer name", "Sales invoice no.", "Invoice Date", "Due Date", "Payment Term in Days",
             "Invoice Value Local Currency（RMB）", "Due days"]
    ws.append([]), ws.append([]), ws.append([]), ws.append([]), ws.append([]), ws.append([])
    ws.append(heads)
    for no, due, amount in (("S1", dt.datetime(2026, 4, 12), 100), ("S2", dt.datetime(2026, 4, 13), 200), ("S3", dt.datetime(2026, 4, 14), 400)):
        ws.append(["Acme", "测试有限公司", no, dt.datetime(2026, 3, 1), due, 30, amount, (due - dt.datetime(2026, 4, 13)).days])
    with tempfile.TemporaryDirectory() as d:
        path = os.path.join(d, "Financial_Status_20260413.xlsx")
        wb.save(path)
        sheet = receivables.read_workbook(path, receivables.FileInfo(path=path, name="x"))[0]
    assert (sheet.open_total, sheet.overdue_total) == (Decimal("700"), Decimal("100"))
