"""The cell ledger: every non-empty cell of a workbook, read with openpyxl / xlrd only (no importer code).

This is the independent record of what is written in the files. Text keeps its line breaks. Each cell
carries its address, whether its row or column is hidden, and whether it is the top-left cell of a
merged range, so a check can say exactly where a value came from.
"""
import datetime as dt
import os
import re
from dataclasses import dataclass


@dataclass(frozen=True)
class Cell:
    file: str  # base name of the workbook
    sheet: str  # exact sheet name, trailing spaces included
    row: int
    col: int
    coord: str  # "F9"
    text: str  # raw text, line breaks kept
    hidden: bool = False
    merged_origin: bool = False
    formula: str = ""  # the formula, when the cell holds one (text is then the cached result)


def col_letters(col: int) -> str:
    out = ""
    while col:
        col, rem = divmod(col - 1, 26)
        out = chr(65 + rem) + out
    return out


def _text(value) -> str:
    if value is None or isinstance(value, bool):
        return "" if value is None else str(value)
    if isinstance(value, dt.datetime):
        return value.isoformat(sep=" ") if value.time() != dt.time(0) else value.date().isoformat()
    if isinstance(value, dt.date):
        return value.isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value)


_COORD = re.compile(r"^([A-Z]+)(\d+)$")


def col_number(letters: str) -> int:
    n = 0
    for ch in letters:
        n = n * 26 + ord(ch) - 64
    return n


def _xlsx(path, merged_out=None):
    import openpyxl
    base = os.path.basename(path)
    values = openpyxl.load_workbook(path, data_only=True)
    try:
        formulas = openpyxl.load_workbook(path, data_only=False)
    except Exception:  # a workbook that only opens with cached values is still readable
        formulas = None
    for ws in values.worksheets:
        merged_origin = {(r.min_row, r.min_col) for r in ws.merged_cells.ranges}
        if merged_out is not None:
            merged_out[ws.title] = [(r.min_row, r.min_col, r.max_row, r.max_col) for r in ws.merged_cells.ranges]
        fws = formulas[ws.title] if formulas is not None and ws.title in formulas.sheetnames else None
        for row in ws.iter_rows():
            for c in row:
                if c.value is None or str(c.value) == "":
                    continue
                f = fws[c.coordinate].value if fws is not None else None
                hidden = bool(ws.row_dimensions[c.row].hidden or ws.column_dimensions[col_letters(c.column)].hidden)
                yield Cell(base, ws.title, c.row, c.column, c.coordinate, _text(c.value), hidden,
                           (c.row, c.column) in merged_origin, f if isinstance(f, str) and f.startswith("=") else "")


def _xls(path, merged_out=None):
    import xlrd
    base = os.path.basename(path)
    book = xlrd.open_workbook(path, formatting_info=True)
    for sh in book.sheets():
        merged_origin = {(r0 + 1, c0 + 1) for r0, r1, c0, c1 in sh.merged_cells}
        if merged_out is not None:
            merged_out[sh.name] = [(r0 + 1, c0 + 1, r1, c1) for r0, r1, c0, c1 in sh.merged_cells]
        for r in range(sh.nrows):
            info = sh.rowinfo_map.get(r)
            for c in range(sh.ncols):
                cell = sh.cell(r, c)
                if cell.ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK) or cell.value == "":
                    continue
                v = cell.value
                if cell.ctype == xlrd.XL_CELL_DATE:
                    v = xlrd.xldate_as_datetime(v, book.datemode)
                colinfo = sh.colinfo_map.get(c)
                hidden = bool((info and info.hidden) or (colinfo and colinfo.hidden))
                yield Cell(base, sh.name, r + 1, c + 1, f"{col_letters(c + 1)}{r + 1}", _text(v), hidden,
                           (r + 1, c + 1) in merged_origin)


_NAMELESS_STYLE = re.compile(r"<cellStyle(?![^>]*\bname=)[^>]*/>")


def _repaired_copy(path):
    """A temp copy of the workbook without nameless <cellStyle> entries (ERP exports have them; openpyxl refuses
    them). The file on disk is never touched."""
    import io
    import tempfile
    import zipfile
    buf = io.BytesIO()
    with zipfile.ZipFile(path) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                data = _NAMELESS_STYLE.sub("", data.decode("utf8")).encode("utf8")
            zout.writestr(item, data)
    fd, tmp = tempfile.mkstemp(suffix=".xlsx")
    with os.fdopen(fd, "wb") as f:
        f.write(buf.getvalue())
    return tmp


def cells(path, merged_out=None):
    """Every non-empty cell of the workbook at `path`. `merged_out`, when given, is filled with
    {sheet: [(first_row, first_col, last_row, last_col)]} for its merged ranges."""
    if path.lower().endswith(".xls"):
        return _xls(path, merged_out)
    try:
        return list(_xlsx(path, merged_out))
    except TypeError:  # a nameless <cellStyle>: read a repaired copy, but keep the real file name
        tmp = _repaired_copy(path)
        try:
            return [Cell(**{**c.__dict__, "file": os.path.basename(path)}) for c in _xlsx(tmp, merged_out)]
        finally:
            os.unlink(tmp)


class Ledger:
    """Lookup of one workbook's cells by (sheet, coordinate)."""

    def __init__(self, path):
        self.path = path
        self.file = os.path.basename(path)
        merged = {}
        self.by_coord = {(c.sheet, c.coord): c for c in cells(path, merged)}
        self._origin = {}  # (sheet, coord) of any cell in a merged range -> coord of its top-left cell
        for sheet, ranges in merged.items():
            for r0, c0, r1, c1 in ranges:
                origin = f"{col_letters(c0)}{r0}"
                for r in range(r0, r1 + 1):
                    for c in range(c0, c1 + 1):
                        self._origin[(sheet, f"{col_letters(c)}{r}")] = origin

    def text(self, sheet, coord) -> str:
        """The cell's text; a cell inside a merged range reads as the range's top-left cell."""
        c = self.by_coord.get((sheet, coord))
        if c is None and (sheet, coord) in self._origin:
            c = self.by_coord.get((sheet, self._origin[(sheet, coord)]))
        return c.text if c else ""

    def sheets(self):
        return sorted({s for s, _ in self.by_coord})

    def __len__(self):
        return len(self.by_coord)
