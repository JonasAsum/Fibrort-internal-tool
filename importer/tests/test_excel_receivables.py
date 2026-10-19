"""Receivables (AR) from the finance workbooks: which sheets are CRM data, the open/overdue figures per
customer, newest snapshot wins, and the Margince adapter against a fake CRM. Every value is invented; the
real-folder test asserts counts only."""
import copy
import datetime as dt
import glob
import json
import os
import shutil
from decimal import Decimal

import openpyxl
import pytest

from common.client import ApiError
from excel import receivables_margince as rm
from excel.companies import Master
from excel.receivables import (AGING, BANK, DOWN, FIXED, HIDDEN, PAYABLE, SUMMARY, build_book, read_workbook,
                               sheet_reason, summarize, FileInfo)
from excel.sheet import load
from tests.conftest import SAMPLE_DIR
from tests.fake_crm import FakeCRM

MASTER = [{"source": "master", "code": "KD00001", "cn": "宇宙星辰科技有限公司", "en": "Cosmos Stellar Co., Ltd.", "short": ""}]
BANK_NO, TAX_ID, AP_MARK, FIX_MARK, PAID_MARK = "6222000011112222", "TAXID-91310000TEST", "AP-SUPPLIER-777", "FIXED-COST-888", "DOWNPAY-999"
HEADER = ["Customer name(EN)", "Customer name", "Sales invoice no.", "Invoice Date", "Due Date", "Payment Term in Days",
          "Invoice Value Local Currency（RMB）", "Due days", "Days since Invoice", "Remark", "币种", "Original amt", "Currency"]
AS_OF = dt.datetime(2026, 4, 13)


def master():
    return Master.from_records(MASTER)


def inv(en="Acme Test Co.", cn="艾测试有限公司", no="SIN001", due=dt.datetime(2026, 1, 1), amount=1000, issued=dt.datetime(2025, 12, 1),
        remark=None, cur="CNY"):
    return [en, cn, no, issued, due, 30, amount, None, None, remark, "人民币", amount, cur]


def fin_book(tmp_path, name="2026_Financial_Status_20260410.xlsx", rows=(), today=AS_OF, report=dt.datetime(2026, 4, 10),
             claimed=None, header=HEADER, extra_sheets=True):
    """A finance workbook like the real one: Overview, AR (with the sheet's own totals) and the sheets that are not CRM data."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ov = wb.create_sheet("Overview")
    ov["G2"], ov["H2"] = "Date:", report
    ov["L6"], ov["M6"] = "BOC-RMB", BANK_NO
    ws = wb.create_sheet("AR")
    ws["G2"], ws["G4"] = "Total Account Receivables", "Total Amount Overdue"
    shown = [r for r in rows if isinstance(r[6], (int, float))]
    day = today.date() if today else None
    open_ = sum(r[6] for r in shown)
    overdue = sum(r[6] for r in shown if r[4] and day and r[4].date() < day)
    ws["A3"], ws["G3"] = "Account Receivables", claimed[0] if claimed else open_
    if today:
        ws["A5"], ws["B5"] = today, "Today !"
    ws["G5"] = claimed[1] if claimed else overdue
    for c, label in enumerate(header, 1):
        ws.cell(7, c, label)
    for i, row in enumerate(rows, 8):
        for c, v in enumerate(row, 1):
            ws.cell(i, c, v)
    n = 8 + len(rows)
    for k in range(3):  # the formula-error rows below the list in the real sheet
        ws.cell(n + k, 1, "#N/A")
        ws.cell(n + k, 5, "No Invoice")
    if extra_sheets:
        ap = wb.create_sheet("AP")
        ap.append(["Supplier", "Invoice", "Amount"])
        ap.append([AP_MARK, "P-1", 424242])
        for nm in ("Cash In-BOC", "Cash Out-BOC", "EBA In-ABC"):
            cash = wb.create_sheet(nm)
            cash.append(["Date", "Counterparty", "Account", "Amount"])
            cash.append([dt.datetime(2026, 4, 1), "Someone", BANK_NO, 5150])
        fx = wb.create_sheet("Fix Cost")
        fx.append(["Item", "Amount"])
        fx.append([FIX_MARK, 31337])
        hid = wb.create_sheet("EN mapping")
        hid.append(["Account", "Tax ID"])
        hid.append([BANK_NO, TAX_ID])
        hid.sheet_state = "hidden"
    p = str(tmp_path / name)
    wb.save(p)
    return p


# ---- the list ----

def test_reads_the_invoice_list_and_its_totals_agree_with_the_sheet(tmp_path):
    p = fin_book(tmp_path, rows=[inv(), inv(no="SIN002", due=dt.datetime(2026, 6, 1), amount=500.5), inv("Beta Test Co.", "贝塔测试有限公司", "SIN003", amount=200)])
    b = build_book([p], master())
    assert b.as_of == "2026-04-13" and b.currency == "CNY" and b.source == "2026_Financial_Status_20260410.xlsx"
    assert not b.sheet.warnings
    c = {x.name: x for x in b.customers.values()}
    assert (c["艾测试有限公司"].open, c["艾测试有限公司"].overdue, c["艾测试有限公司"].invoices) == (Decimal("1500.50"), Decimal(1000), 2)
    assert c["贝塔测试有限公司"].overdue == Decimal(200)
    s = summarize(b)
    assert (s["open_total"], s["overdue_total"], s["invoices"], s["customers"]) == ("1700.50", "1200.00", 3, 2)
    assert b.sheet.skipped_rows == 3  # the "#N/A / No Invoice" placeholder rows


def test_due_today_is_not_overdue_and_the_boundary_is_the_as_of_day(tmp_path):
    p = fin_book(tmp_path, rows=[inv(due=dt.datetime(2026, 4, 13)), inv(no="B", due=dt.datetime(2026, 4, 12))])
    assert next(iter(build_book([p], master()).customers.values())).overdue == Decimal(1000)


def test_row_without_invoice_number_counts_as_open_never_as_overdue(tmp_path):
    p = fin_book(tmp_path, rows=[inv(no=None, due=None, amount=169800)])
    b = build_book([p], master())
    c = next(iter(b.customers.values()))
    assert (c.open, c.overdue) == (Decimal(169800), Decimal(0))
    assert [r.reason for r in b.review] == ["invoice_doubtful", "no_due_date"]
    assert b.accepted() == []  # a total that may hold an invoice nobody can name is not written


def test_the_same_invoice_number_twice_holds_that_customer_out(tmp_path):
    p = fin_book(tmp_path, rows=[inv(no="SIN9"), inv(no="SIN9", amount=500), inv(en="Other Co.", cn="他公司有限公司", no="SIN1")])
    b = build_book([p], master())
    held = [c.name for c in b.customers.values() if c.doubts]
    assert len(held) == 1 and len(b.accepted()) == len(b.customers) - 1


def test_zero_and_placeholder_rows_are_not_invoices(tmp_path):
    p = fin_book(tmp_path, rows=[inv(amount=0), inv(no="B", amount=10)])
    b = build_book([p], master())
    assert len(b.sheet.invoices) == 1


def test_english_name_annotation_is_not_part_of_the_name(tmp_path):
    p = fin_book(tmp_path, rows=[inv(en="Acme Test (transfer from the old ledger)", cn="")])
    assert next(iter(build_book([p], master()).customers.values())).name == "Acme Test"


def test_sheet_totals_that_disagree_are_warned_about(tmp_path):
    p = fin_book(tmp_path, rows=[inv()], claimed=(5000, 1000))
    b = build_book([p], master())
    assert any("open total" in w for w in b.sheet.warnings) and not any("overdue total" in w for w in b.sheet.warnings)


def test_as_of_date_comes_from_the_today_cell_then_the_overview_then_the_file_name(tmp_path):
    rows = [inv()]
    assert build_book([fin_book(tmp_path, "a_20260101.xlsx", rows)], master()).as_of == "2026-04-13"
    b = build_book([fin_book(tmp_path, "b_20260101.xlsx", rows, today=None)], master())
    assert b.as_of == "2026-04-10" and b.sheet.as_of_source == "the workbook's Overview date"
    wb = openpyxl.load_workbook(fin_book(tmp_path, "c_20260201.xlsx", rows, today=None))
    wb["Overview"]["H2"] = None
    wb.save(tmp_path / "c_20260201.xlsx")
    b = build_book([str(tmp_path / "c_20260201.xlsx")], master())
    assert b.as_of == "2026-02-01" and b.sheet.as_of_source == "the date in the file name"


def test_a_list_that_does_not_name_its_currency_or_date_is_not_usable(tmp_path):
    header = list(HEADER)
    header[6] = "Invoice Value Local Currency"
    b = build_book([fin_book(tmp_path, rows=[inv()], header=header)], master())
    assert b.as_of is None and not b.customers and b.review[0].reason == "undated_or_unlabelled_snapshot"
    with pytest.raises(rm.ReceivablesError):
        rm.commit(ArCRM(), b)


# ---- which sheets are not CRM data ----

def test_everything_but_the_invoice_list_is_excluded_with_a_reason(tmp_path):
    p = fin_book(tmp_path, rows=[inv()])
    info = FileInfo(path=p, name="x")
    assert len(read_workbook(p, info)) == 1 and info.sheets_read == ["AR"]
    why = dict(info.excluded_sheets)
    assert why["Overview"] == SUMMARY and why["AP"] == PAYABLE and why["Fix Cost"] == FIXED
    assert why["Cash In-BOC"] == why["Cash Out-BOC"] == why["EBA In-ABC"] == BANK
    assert why["EN mapping"] == HIDDEN


def test_down_payments_and_aging_ledgers_are_not_receivables(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    a = wb.create_sheet("Received in advance 202601")
    a.append(["Customer Number", "Customer Name", "downpayment (RMB)"])
    a.append(["CR1", "Acme Test", 72000])
    g = wb.create_sheet("pure RT AR-20231231")
    g.append(["Customer No", "Customer", "AR Amount", "within payment terms", "1-30days overdue"])
    g.append(["CR1", "Acme Test", 5, 5, 0])
    p = str(tmp_path / "ledger.xlsx")
    wb.save(p)
    b = build_book([p], master())
    f = b.files[0]
    assert f.kind == "excluded" and DOWN in f.reason and AGING in f.reason and not b.customers
    assert dict(f.excluded_sheets) == {"Received in advance 202601": DOWN, "pure RT AR-20231231": AGING}


def test_hidden_sheets_are_never_read(tmp_path):
    wb = openpyxl.load_workbook(fin_book(tmp_path, rows=[inv()]))
    wb["AR"].sheet_state = "hidden"
    wb.save(tmp_path / "h.xlsx")
    b = build_book([str(tmp_path / "h.xlsx")], master())
    assert not b.customers and b.files[0].kind == "excluded"


def test_sheet_name_beats_what_the_first_rows_say(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Fix Cost"
    ws.append(["Bank charges", "Cash"])
    p = str(tmp_path / "f.xlsx")
    wb.save(p)
    assert sheet_reason(load(p)[0], False) == FIXED


# ---- newest snapshot wins ----

def test_newest_snapshot_replaces_older_ones_and_identical_copies_are_read_once(tmp_path):
    old = fin_book(tmp_path, "old_20260228.xlsx", [inv(amount=100), inv("Gone Test Co.", "消失测试有限公司", "X1", amount=50)],
                   today=dt.datetime(2026, 3, 4))
    new = fin_book(tmp_path, "new_20260410.xlsx", [inv(amount=300)])
    shutil.copy(new, tmp_path / "0123456789ab_new_20260410.xlsx")
    b = build_book([old, new, str(tmp_path / "0123456789ab_new_20260410.xlsx")], master())
    assert b.source == "new_20260410.xlsx" and b.as_of == "2026-04-13"
    assert [c.open for c in b.customers.values()] == [Decimal(300)]
    assert [(f.name, f.superseded_by) for f in b.files if f.superseded_by] == [("old_20260228.xlsx", "new_20260410.xlsx")]
    assert any("identical" in s["error"] for s in b.skipped_files)


def test_unreadable_file_is_reported_and_the_run_goes_on(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"nope")
    b = build_book([str(bad), fin_book(tmp_path, rows=[inv()])], master())
    assert len(b.customers) == 1 and b.skipped_files
    assert next(f for f in b.files if f.name == "bad.xlsx").kind == "excluded"


def test_group_company_and_unanswered_customers_are_flagged(tmp_path):
    p = fin_book(tmp_path, rows=[inv("FIBRO Test GmbH", "", "G1"), inv("Cosmic Stellar Company Ltd", "", "Q1")])
    b = build_book([p], master())
    by = {c.name: c for c in b.customers.values()}
    assert by["FIBRO Test GmbH"].group_company and by["Cosmic Stellar Company Ltd"].queued
    assert {r.reason for r in b.review} == {"group_company", "company"}


# ---- the Margince adapter ----

class ArCRM(FakeCRM):
    def __init__(self, refuse_fields=None, drop=()):
        super().__init__()
        self.refuse_fields, self.drop = refuse_fields, set(drop)

    def paginate(self, path, params=None):
        if path == "/v1/companies":
            yield from [{**c, "id": i} for i, c in self.companies.items()]
            return
        yield from super().paginate(path, params)

    def create_custom_field(self, body):
        if self.refuse_fields:
            raise self.refuse_fields
        return super().create_custom_field(body)

    def patch(self, path, body, **kw):
        return super().patch(path, {k: v for k, v in body.items() if k not in self.drop}, **kw)


def refused(status=501):
    import requests
    resp = requests.Response()
    resp.status_code = status
    resp._content = b'{"code": "not_implemented"}'
    resp.request = requests.Request("POST", "http://fake/v1/custom-fields").prepare()
    return ApiError(resp)


def snapshot(crm):
    return copy.deepcopy((crm.companies, crm.activities, crm.fields))


def two_customers(tmp_path, name="a_20260410.xlsx", today=AS_OF, a=1000, b=200, report=dt.datetime(2026, 4, 10)):
    return fin_book(tmp_path, name, [inv(amount=a), inv("Beta Test Co.", "贝塔测试有限公司", "B1", amount=b, due=dt.datetime(2026, 9, 1))],
                    today=today, report=report)


def cols(crm):
    return {f["label"]: f["column_name"] for f in crm.fields}


def test_commit_creates_the_fields_and_writes_whole_fen_with_the_date(tmp_path):
    crm = ArCRM()
    r = rm.commit(crm, build_book([two_customers(tmp_path)], master()))
    assert r["mode"] == "fields" and r["accounts_written"] == 2 and not r["errors"]
    assert {f["label"]: (f["type"], f.get("currency")) for f in crm.fields} == {
        "AR open (CNY)": ("currency", "CNY"), "AR overdue (CNY)": ("currency", "CNY"), "AR as-of date": ("date", None)}
    c, acme = cols(crm), next(x for x in crm.companies.values() if x["display_name"] == "艾测试有限公司")
    assert (acme[c["AR open (CNY)"]], acme[c["AR overdue (CNY)"]], acme[c["AR as-of date"]]) == (100000, 100000, "2026-04-13")
    beta = next(x for x in crm.companies.values() if x["display_name"] == "贝塔测试有限公司")
    assert (beta[c["AR open (CNY)"]], beta[c["AR overdue (CNY)"]]) == (20000, 0)  # not overdue is a real zero
    assert all(x["lifecycle"] == "customer" for x in crm.companies.values())


def test_second_run_writes_nothing_and_preview_writes_nothing(tmp_path):
    crm = ArCRM()
    b = build_book([two_customers(tmp_path)], master())
    before = snapshot(crm)
    p = rm.plan(crm, b)
    assert p["accounts"]["write"] == 2 and p["companies_new"] == 2 and snapshot(crm) == before
    rm.commit(crm, b)
    done = snapshot(crm)
    r = rm.commit(crm, build_book([two_customers(tmp_path)], master()))
    assert (r["accounts_written"], r["accounts_unchanged"], r["accounts_zeroed"]) == (0, 2, 0)
    assert snapshot(crm) == done
    assert rm.plan(crm, b)["accounts"] == {"write": 0, "same": 2, "newer": 0, "zero": 0}


def test_a_newer_file_replaces_the_figures_an_older_one_never_does(tmp_path):
    crm = ArCRM()
    old = two_customers(tmp_path, "old_20260228.xlsx", today=dt.datetime(2026, 3, 4), a=500, b=200)
    new = two_customers(tmp_path, "new_20260410.xlsx", a=1000, b=200)
    rm.commit(crm, build_book([old], master()))
    r = rm.commit(crm, build_book([new], master()))
    c = cols(crm)
    acme = next(x for x in crm.companies.values() if x["display_name"] == "艾测试有限公司")
    beta = next(x for x in crm.companies.values() if x["display_name"] == "贝塔测试有限公司")
    assert r["accounts_written"] == 2  # acme's figure moved, beta's figures stand but now hold the newer date
    assert (acme[c["AR open (CNY)"]], acme[c["AR as-of date"]]) == (100000, "2026-04-13")
    assert (beta[c["AR open (CNY)"]], beta[c["AR as-of date"]]) == (20000, "2026-04-13")
    r = rm.commit(crm, build_book([old], master()))  # the old list again, after the new one
    assert r["accounts_kept_newer"] == 2 and r["accounts_written"] == 0
    assert acme[c["AR open (CNY)"]] == 100000 and acme[c["AR as-of date"]] == "2026-04-13"


def test_same_figures_on_a_newer_date_only_move_the_date(tmp_path):
    crm = ArCRM()
    rm.commit(crm, build_book([two_customers(tmp_path, "a.xlsx")], master()))
    r = rm.commit(crm, build_book([two_customers(tmp_path, "b.xlsx", today=dt.datetime(2026, 4, 20))], master()))
    acme = next(x for x in crm.companies.values() if x["display_name"] == "艾测试有限公司")
    assert r["accounts_written"] == 2 and acme[cols(crm)["AR as-of date"]] == "2026-04-20"


def test_a_customer_missing_from_a_newer_snapshot_has_paid_and_goes_to_zero(tmp_path):
    crm = ArCRM()
    rm.commit(crm, build_book([two_customers(tmp_path, "old_20260228.xlsx", today=dt.datetime(2026, 3, 4))], master()))
    newer = fin_book(tmp_path, "new_20260410.xlsx", [inv(amount=1000)])
    r = rm.commit(crm, build_book([newer], master()))
    beta = next(x for x in crm.companies.values() if x["display_name"] == "贝塔测试有限公司")
    c = cols(crm)
    assert r["accounts_zeroed"] == 1 and (beta[c["AR open (CNY)"]], beta[c["AR overdue (CNY)"]], beta[c["AR as-of date"]]) == (0, 0, "2026-04-13")
    assert rm.commit(crm, build_book([newer], master()))["accounts_zeroed"] == 0


def test_nobody_is_zeroed_while_a_customer_of_the_snapshot_waits_for_an_answer(tmp_path):
    crm = ArCRM()
    rm.commit(crm, build_book([two_customers(tmp_path, "old_20260228.xlsx", today=dt.datetime(2026, 3, 4))], master()))
    newer = fin_book(tmp_path, "new_20260410.xlsx", [inv(amount=1000), inv("Cosmic Stellar Company Ltd", "", "Q1")])
    b = build_book([newer], master())
    r = rm.commit(crm, b)
    beta = next(x for x in crm.companies.values() if x["display_name"] == "贝塔测试有限公司")
    assert r["accounts_zeroed"] == 0 and beta[cols(crm)["AR open (CNY)"]] == 20000
    assert r["customers_waiting_for_a_company_answer"] == 1 and len(crm.companies) == 2


def test_a_value_the_crm_does_not_keep_is_reported(tmp_path):
    crm = ArCRM(drop=["cf_ar_overdue_(cny)"])
    r = rm.commit(crm, build_book([two_customers(tmp_path)], master()))
    assert len(r["errors"]) == 2 and "did not keep" in r["errors"][0]


def test_existing_fields_are_reused_and_a_field_of_another_type_is_not_touched(tmp_path):
    crm = ArCRM()
    crm.fields.append({"label": "AR open (CNY)", "type": "text", "column_name": "cf_someone_elses"})
    rm.commit(crm, build_book([two_customers(tmp_path)], master()))
    assert sorted(f["label"] for f in crm.fields) == ["AR as-of date", "AR open (CNY)", "AR open (CNY) (receivables)",
                                                      "AR overdue (CNY)"]
    n = len(crm.fields)
    rm.commit(crm, build_book([two_customers(tmp_path)], master()))
    assert len(crm.fields) == n


def test_when_the_crm_refuses_custom_fields_an_updated_note_per_account_holds_the_figures(tmp_path):
    crm = ArCRM(refuse_fields=refused(501))
    b = build_book([two_customers(tmp_path)], master())
    r = rm.commit(crm, b)
    assert r["mode"] == "notes" and r["notes_created"] == 2 and not crm.fields
    note = next(a for a in crm.activities.values() if a["source_id"].startswith("ar:") and "艾测试" in a["source_id"])
    assert "Open receivables: CNY 1,000.00" in note["body"] and "As of: 2026-04-13" in note["body"]
    assert note["source_system"] == rm.SOURCE_SYSTEM and note["links"][0]["entity_type"] == "company"
    before = snapshot(crm)
    assert rm.commit(crm, b)["notes_unchanged"] == 2 and snapshot(crm) == before
    newer = two_customers(tmp_path, "n_20260501.xlsx", today=dt.datetime(2026, 5, 2), a=1500)
    assert rm.commit(crm, build_book([newer], master()))["notes_updated"] == 2
    r = rm.commit(crm, b)  # the old list after the new one
    assert r["notes_kept_newer"] == 2 and r["notes_updated"] == 0
    assert "CNY 1,500.00" in note["body"]


def test_other_api_errors_on_field_creation_are_not_swallowed(tmp_path):
    with pytest.raises(ApiError):
        rm.commit(ArCRM(refuse_fields=refused(500)), build_book([two_customers(tmp_path)], master()))


# ---- what must never reach a record or a payload ----

def test_no_bank_account_tax_id_ap_fixed_cost_or_down_payment_value_reaches_any_record_or_payload(tmp_path):
    header = list(HEADER) + ["Bank Account", "Tax ID"]
    rows = [inv(remark=f"{BANK_NO} {TAX_ID}") + [BANK_NO, TAX_ID]]
    p = fin_book(tmp_path, rows=rows, header=header)
    down = openpyxl.Workbook()
    down.active.title = "Received in advance 202601"
    down.active.append(["Customer Number", "Customer Name", "downpayment (RMB)"])
    down.active.append([PAID_MARK, "Acme Test", 98765])
    down.save(tmp_path / "down.xlsx")
    for mode in (ArCRM(), ArCRM(refuse_fields=refused(501))):
        b = build_book([p, str(tmp_path / "down.xlsx")], master())
        rm.commit(mode, b)
        everything = json.dumps([b.customers, b.review, b.files, b.sheet, summarize(b), mode.companies, mode.activities,
                                 mode.fields, rm.plan(mode, b)], default=lambda o: getattr(o, "__dict__", str(o)),
                                ensure_ascii=False)
        for secret in (BANK_NO, TAX_ID, AP_MARK, FIX_MARK, PAID_MARK, "424242", "31337", "98765", "5150"):
            assert secret not in everything, secret


# ---- the real folder: counts only ----

def test_real_folder_receivables():
    folder = os.path.join(SAMPLE_DIR, "09 AR and Financial Status")
    if not os.path.isdir(folder):
        pytest.skip("sample folder not present")
    paths = sorted(glob.glob(os.path.join(folder, "*.xlsx")))
    assert len(paths) == 4
    b = build_book(paths, master())
    s = summarize(b)
    assert (s["files"], s["files_with_receivables"], s["files_excluded"]) == (4, 2, 2)
    assert s["snapshot"] == "2026_Financial_Status_Fibro_CN_20260410.xlsx" and s["as_of"] == "2026-04-13"
    assert (s["invoices"], s["customers"], s["currency"]) == (48, 14, "CNY")
    assert not b.sheet.warnings  # the sheet's own two totals equal the sums of the rows
    assert len(s["superseded"]) == 1 and not s["skipped_files"]
    assert {e["file"] for e in s["excluded"]} == {"Laepple_RT_downpayment_202602.xlsx", "Laepple_RT_APAR_20231231.xlsx"}


def test_file_kind_for_the_wiring(tmp_path):
    from excel.receivables import file_kind, looks_like_receivables
    assert file_kind(fin_book(tmp_path, rows=[inv()])) == ("receivables", "")
    wb = openpyxl.Workbook()
    wb.active.title = "Cash In-BOC"
    wb.save(tmp_path / "cash.xlsx")
    assert file_kind(str(tmp_path / "cash.xlsx")) == ("excluded", BANK)
    assert [looks_like_receivables(s) for s in load(fin_book(tmp_path, "z.xlsx", rows=[inv()]))][:2] == [False, True]
