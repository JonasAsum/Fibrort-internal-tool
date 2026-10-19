"""Workbooks Excel opens but openpyxl refuses (cell styles without a name) must still load."""
import io
import re
import zipfile

import openpyxl
import pytest

from excel import sheet


def nameless_style_book(path):
    """A small made-up workbook whose styles.xml holds a <cellStyle> without a name, as some ERP exports do."""
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sales Order List"
    ws["A1"], ws["B2"] = "Document No.", "SO0001"
    buf = io.BytesIO()
    wb.save(buf)
    zin = zipfile.ZipFile(io.BytesIO(buf.getvalue()))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "xl/styles.xml":
                text = data.decode("utf8")
                assert "<cellStyles" in text
                text = re.sub(r"(<cellStyles[^>]*>)", r'\1<cellStyle xfId="0" builtinId="0"/>', text, count=1)
                data = text.encode("utf8")
            zout.writestr(item, data)
    return str(path)


def test_the_made_up_book_really_breaks_openpyxl(tmp_path):
    path = nameless_style_book(tmp_path / "broken.xlsx")
    with pytest.raises(TypeError):
        openpyxl.load_workbook(path)


def test_load_reads_a_workbook_with_a_nameless_cell_style(tmp_path):
    path = nameless_style_book(tmp_path / "broken.xlsx")
    (s,) = sheet.load(path)
    assert (s.name, s.value(1, 1), s.value(2, 2)) == ("Sales Order List", "Document No.", "SO0001")


def test_a_normal_workbook_is_unchanged(tmp_path):
    wb = openpyxl.Workbook()
    wb.active["A1"] = "x"
    wb.save(tmp_path / "ok.xlsx")
    (s,) = sheet.load(str(tmp_path / "ok.xlsx"))
    assert s.value(1, 1) == "x"


def test_a_file_that_is_not_a_workbook_still_raises(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"this is not a zip file")
    with pytest.raises(Exception):
        sheet.load(str(bad))


def test_a_truncated_zip_still_raises(tmp_path):
    good = nameless_style_book(tmp_path / "broken.xlsx")
    data = open(good, "rb").read()
    cut = tmp_path / "cut.xlsx"
    cut.write_bytes(data[: len(data) // 2])
    with pytest.raises(Exception):
        sheet.load(str(cut))
