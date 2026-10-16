"""Every-file matrix: each file of the real sample folder is pushed through the importer of its family
and compared with tests/golden_files_matrix.json (counts and reasons only, no cell values). Regenerate with
REGEN_MATRIX=1. Synthetic tests at the bottom prove the importer paths reject bad files with a reason and
never raise.

Two layers, both from the same folder walk (nothing is hard-coded, so a new file fails until it has a row):
  1. the reader of the file's family, called directly (the top-level keys of each golden row)
  2. the SITE: the file is uploaded through /api/upload exactly as the page does, then dry-run, committed
     into an in-memory CRM and committed again (the "site" key of each golden row). Each file is its own
     test case; a file the front door cannot place, or that is not kept out on purpose, fails its case.

Families: 01-03 visit importer (excel.run.read_paths), 04-05 accounts (excel.accounts), 06 contacts
(excel.contacts), 07 to-dos (excel.todos), 08 orders (excel.orders), 09 receivables (excel.receivables).
"""
import functools
import hashlib
import json
import os
import re
import threading
import unicodedata
from collections import Counter
from unittest import mock

import pytest

from tests.conftest import SAMPLE_DIR
from tests.site_fixture import site  # noqa: F401  (the fixture)

GOLDEN_MATRIX = os.path.join(os.path.dirname(os.path.abspath(__file__)), "golden_files_matrix.json")
VISIT_FAMILIES = ("01", "02", "03")
CONTACT_FAMILIES = ("06",)
SKIP_TOP = {"import-tests"}


def _nfc(s):
    return unicodedata.normalize("NFC", s)


def _is_junk(name):
    return name.startswith(".") or name.startswith("~$") or name.endswith("#")


def _family_of(rel):
    top = rel.split(os.sep)[0]
    code = top[:2]
    return code if os.sep in rel and code.isdigit() and top[2:3] == " " and code != "00" else None


def walk_files(root):
    """[(relative path (NFC), family or None, absolute path)] for every file below a family
    folder; dot-directories, import-tests and top-level files are not part of the matrix."""
    out = []
    for dirpath, dirs, files in os.walk(root):
        rel_dir = os.path.relpath(dirpath, root)
        dirs[:] = sorted(d for d in dirs if not d.startswith(".") and not (rel_dir == "." and d in SKIP_TOP))
        for name in sorted(files):
            absolute = os.path.join(dirpath, name)
            rel = _nfc(os.path.relpath(absolute, root))
            fam = _family_of(os.path.relpath(absolute, root))
            if fam is None:
                continue
            out.append((rel, fam, absolute))
    return out


def _err(exc):
    return type(exc).__name__  # type only: messages can echo sheet names or cell text


def process_visit_file(path):
    from excel.run import read_paths
    try:
        res = read_paths([path])
    except Exception as exc:
        return "FAILED", {"error": "read_paths raised " + _err(exc)}
    stat = res.per_file_stats[0] if res.per_file_stats else None
    if stat is None or (stat.error and not stat.sheets):
        reason = "unreadable" if stat is None else stat.error.split(":")[0]
        return "FAILED", {"error": reason}
    queued = Counter()
    for item in res.review:
        queued[item.reason] += 1
    visits = len(res.parsed.visits)
    acc = len(res.auto_accepted)
    return "imported", {
        "layout": stat.family, "sheets": stat.sheets, "sheets_read": stat.read_sheets,
        "sheets_unrecognised": len(stat.unrecognised), "sheet_errors": len(stat.sheet_errors),
        "visits": visits, "auto_accepted": acc, "queued": visits - acc,
        "queued_by_reason": dict(sorted(queued.items())),
        "skipped": [s["error"].split(":")[0] for s in res.skipped_files],
    }


def process_contact_file(path):
    from excel import contacts
    try:
        res = contacts.read_paths([path])
    except Exception as exc:
        return "FAILED", {"error": "contacts.read_paths raised " + _err(exc)}
    stat = res.per_file_stats[0] if res.per_file_stats else None
    if stat is None or (stat.error and not stat.sheets):
        return "FAILED", {"error": "unreadable" if stat is None else stat.error.split(":")[0]}
    s = contacts.summarize(res)
    return "imported", {
        "layout": stat.layout, "sheets": s["sheets"], "sheets_read": s["read_sheets"],
        "sheets_unrecognised": len(stat.unrecognised), "sheet_errors": len(stat.sheet_errors),
        "rows": s["rows"], "blank_rows": s["blank_rows"], "contacts": s["contacts"],
        "contacts_accepted": s["contacts_accepted"], "companies": s["companies"],
        "companies_queued": s["companies_queued"], "queued_by_reason": s["queued_by_reason"],
        "skipped_by_reason": s["skipped_by_reason"], "skipped": [x["error"].split(":")[0] for x in res.skipped_files],
    }


def process_importer_file(path):
    """Families 04, 05, 07, 08, 09: the front door places the file, then the reader of its family counts it.
    A file the review says does not belong in the CRM is "kept_out" with the review's reason."""
    from webui import entry, families
    try:
        e = entry.classify(path)
        if e["kind"] == "kept_out":
            return "kept_out", {"importer": None, "what": e["what"], "reason": e["reason"], "folder": e["folder"]}
        if e["kind"] != "import":
            return "FAILED", {"error": f"front door: {e['kind']}"}
        adapter = families.ADAPTERS[e["family"]]
        res = adapter.read([path])
        records, unit = adapter.file_stats(res).get(path, (0, ""))
        reasons = Counter(i.reason for i in res.review)
        return ("imported" if records else "empty"), {
            "importer": e["family"], "layout": e["layout"], "sheets": e["sheets"], "sheets_used": e["sheets_used"],
            "records": records, "unit": unit, "ready": adapter.writable(res), "waiting": adapter.waiting(res),
            "counts": {r["label"]: r["value"] for r in adapter.counts(res) if isinstance(r["value"], int)},
            "review_by_reason": dict(sorted(reasons.items())), "files_skipped": len(adapter.skipped(res))}
    except Exception as exc:
        return "FAILED", {"error": "importer raised " + _err(exc)}


def process_file(name, fam, path):
    if _is_junk(name):
        return "ignored", {}
    if not name.lower().endswith((".xlsx", ".xls")):
        return "ignored", {"reason": "not a workbook"}
    if fam in VISIT_FAMILIES:
        return process_visit_file(path)
    if fam in CONTACT_FAMILIES:
        return process_contact_file(path)
    return process_importer_file(path)


def build_matrix(root):
    matrix = {}
    for rel, fam, absolute in walk_files(root):
        status, counts = process_file(os.path.basename(rel), fam, absolute)
        matrix[rel] = {"family": fam, "status": status, **counts}
    return matrix


def comparable(matrix):
    """Lock files and other junk are tolerated: they are ignored and never asserted."""
    return {k: v for k, v in matrix.items() if v["status"] != "ignored"}


GOLDEN_LOCK = threading.Lock()


def load_golden():
    if not os.path.exists(GOLDEN_MATRIX):
        return {}
    with open(GOLDEN_MATRIX, encoding="utf-8") as fh:
        return json.load(fh)


def save_golden(golden):
    with open(GOLDEN_MATRIX, "w", encoding="utf-8") as fh:
        json.dump(golden, fh, indent=1, sort_keys=True, ensure_ascii=False)
        fh.write("\n")


def readers_only(golden):
    """The reader-level rows of the golden: no "site" results, no whole-folder entry."""
    return {k: {f: x for f, x in v.items() if f != "site"} for k, v in golden.items() if not k.startswith("_")}


def diff_text(old, new):
    lines = []
    for k in sorted(set(old) | set(new)):
        if k not in new:
            lines.append(f"- missing now: {k}")
        elif k not in old:
            lines.append(f"+ new file: {k} -> {new[k]['status']}")
        elif old[k] != new[k]:
            for f in sorted(set(old[k]) | set(new[k])):
                if old[k].get(f) != new[k].get(f):
                    lines.append(f"~ {k}: {f}: {old[k].get(f)!r} -> {new[k].get(f)!r}")
    return "\n".join(lines)


@pytest.fixture(scope="module")
def matrix():
    if not os.path.isdir(SAMPLE_DIR):
        pytest.skip(f"sample folder not found; set EXCEL_SAMPLE_DIR (looked at {SAMPLE_DIR})")
    return build_matrix(SAMPLE_DIR)


def test_matrix_matches_golden(matrix):
    live = comparable(matrix)
    if os.environ.get("REGEN_MATRIX") == "1":
        with GOLDEN_LOCK:
            old = load_golden()  # the site results and the whole-folder entry survive a reader-level regeneration
            save_golden({**{k: v for k, v in old.items() if k.startswith("_")},
                         **{k: ({**v, "site": old[k]["site"]} if "site" in old.get(k, {}) else v) for k, v in live.items()}})
        return
    if not os.path.exists(GOLDEN_MATRIX):
        pytest.skip("no golden yet: REGEN_MATRIX=1 pytest tests/test_all_files_matrix.py")
    golden = readers_only(load_golden())
    assert live == golden, "matrix differs from golden:\n" + diff_text(golden, live)


def test_every_family_covered(matrix):
    fams = {v["family"] for v in matrix.values()}
    assert fams == {f"0{i}" for i in range(1, 10)}


def test_junk_is_ignored_not_failed(matrix):
    for rel, v in matrix.items():
        if _is_junk(os.path.basename(rel)):
            assert v["status"] == "ignored", rel


def test_no_unexpected_failures(matrix):
    failed = {k: v["error"] for k, v in matrix.items() if v["status"] == "FAILED"}
    assert not failed, f"FAILED files: {failed}"


def test_every_file_has_an_importer_or_a_review_reason_to_stay_out(matrix):
    for rel, v in matrix.items():
        if v["status"] == "ignored":
            continue
        assert v["status"] in ("imported", "kept_out", "empty"), (rel, v["status"])
        if v["status"] == "kept_out":  # only what the folder review says does not belong in the CRM
            assert v["family"] in ("08", "09") and v["reason"] and v["importer"] is None, rel
        elif v["family"] in ("04", "05", "07", "08", "09"):
            assert v["importer"], rel


def test_dot_dirs_and_locks_do_not_change_matrix(tmp_path):
    (tmp_path / "06 Contact Lists").mkdir()
    (tmp_path / "06 Contact Lists" / ".~lock.x.xlsx#").write_text("x")
    (tmp_path / "06 Contact Lists" / ".DS_Store").write_bytes(b"\0")
    (tmp_path / ".hidden").mkdir()
    (tmp_path / ".hidden" / "a.xlsx").write_bytes(b"")
    m = build_matrix(str(tmp_path))
    assert comparable(m) == {}
    assert all(v["status"] == "ignored" for v in m.values())


# ---------- synthetic files: invented data, every path rejects with a reason, never raises ----------

def _make_bad_files(tmp_path):
    import openpyxl
    files = {}
    p = tmp_path / "zero.xlsx"; p.write_bytes(b""); files["zero_xlsx"] = p
    p = tmp_path / "zero.xls"; p.write_bytes(b""); files["zero_xls"] = p
    p = tmp_path / "text.xlsx"; p.write_text("hello, this is not a workbook\n" * 5); files["text_xlsx"] = p
    p = tmp_path / "text.xls"; p.write_text("a,b,c\n1,2,3\n"); files["text_xls"] = p
    good = tmp_path / "real_bytes.xlsx"
    wb = openpyxl.Workbook(); wb.active["A1"] = "invented"; wb.save(good)
    p = tmp_path / "xlsx_as.xls"; p.write_bytes(good.read_bytes()); files["xls_with_xlsx_bytes"] = p
    p = tmp_path / "locked.xlsx"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + os.urandom(4096)); files["ole_random_xlsx"] = p
    p = tmp_path / "locked.xls"
    p.write_bytes(b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1" + os.urandom(4096)); files["ole_random_xls"] = p
    return files


@pytest.mark.parametrize("key", ["zero_xlsx", "zero_xls", "text_xlsx", "text_xls",
                                 "xls_with_xlsx_bytes", "ole_random_xlsx", "ole_random_xls"])
def test_bad_files_rejected_with_reason_never_raise(tmp_path, key):
    path = str(_make_bad_files(tmp_path)[key])
    for fam in ("01", "04", "06"):
        status, counts = process_file(os.path.basename(path), fam, path)  # must not raise
        if key == "xls_with_xlsx_bytes" and fam != "04":
            continue  # xlrd may accept or reject; only "no exception" is required
        assert status in ("FAILED", "imported", "empty", "kept_out")
        if status == "FAILED":
            assert counts.get("error"), (key, fam)
    # the visit importer itself records the skip instead of raising
    from excel.run import read_paths
    res = read_paths([path])
    if key != "xls_with_xlsx_bytes":
        assert res.skipped_files and res.skipped_files[0]["error"].startswith("could not be read")
        assert res.parsed.visits == []


def test_empty_sheets_workbook(book):
    path = book("empty.xlsx", {"One": {"cells": {}}, "Two": {"cells": {}}})
    assert process_file("empty.xlsx", "01", path)[0] == "imported"
    _, c = process_file("empty.xlsx", "01", path)
    assert c["visits"] == 0 and c["layout"] == "none" and c["sheets"] == 2
    status, c = process_file("empty.xlsx", "04", path)  # no list in it: the front door says so, it does not raise
    assert status == "FAILED" and c["error"] == "front door: unknown"
    _, c = process_file("empty.xlsx", "06", path)
    assert c["contacts"] == 0 and c["layout"] == "none" and c["sheets"] == 2
    assert process_file("empty.xlsx", "07", path) == ("FAILED", {"error": "front door: unknown"})


def test_forty_sheet_workbook(book):
    sheets = {f"S{i}": {"cells": {"A1": f"x{i}"}} for i in range(40)}
    path = book("forty.xlsx", sheets)
    _, c = process_file("forty.xlsx", "02", path)
    assert c["sheets"] == 40 and c["visits"] == 0
    assert process_file("forty.xlsx", "08", path) == ("FAILED", {"error": "front door: unknown"})
    assert process_file("forty.xlsx", "05", path) == ("FAILED", {"error": "front door: unknown"})


def test_very_wide_sheet(book):
    from openpyxl.utils import get_column_letter
    cells = {f"{get_column_letter(i)}1": f"h{i}" for i in range(1, 301)}
    path = book("wide.xlsx", {"Wide": {"cells": cells}})
    assert process_file("wide.xlsx", "09", path) == ("FAILED", {"error": "front door: unknown"})
    from excel.sheet import load
    sheet = load(path)[0]
    assert sheet.max_col == 40 and "columns beyond" in sheet.cut  # capped at excel.sheet.MAX_COL
    assert process_file("wide.xlsx", "01", path)[0] == "imported"


# ---------- the site layer: every file through /api/upload, a dry run, a commit and a second commit ----------

def _workbooks_on_disk(root):
    """Relative names (NFC) of every .xlsx/.xls below a family folder (01-09), found with a plain walk of its own:
    dot-folders (.arena), import-tests, lock files and .DS_Store are not part of the folder's lists."""
    found = set()
    for dirpath, dirs, names in os.walk(root):
        rel_dir = os.path.relpath(dirpath, root)
        dirs[:] = [d for d in dirs if not d.startswith(".") and not (rel_dir == "." and d == "import-tests")]
        for name in names:
            rel = _nfc(os.path.relpath(os.path.join(dirpath, name), root))
            if re.match(r"^0[1-9] ", rel) and name.lower().endswith((".xlsx", ".xls")) and not _is_junk(name):
                found.add(rel)
    return found


def _case_id(rel):
    """A readable id that names the file without its customers: folder number, the file's short name, a hash."""
    name = re.sub(r"^[0-9a-f]{12}_", "", os.path.basename(rel))
    slug = re.sub(r"[^A-Za-z0-9]+", "-", os.path.splitext(name)[0]).strip("-")[:40] or "file"
    return f"{rel[:2]}-{slug}-{hashlib.sha1(rel.encode()).hexdigest()[:6]}"


SITE_FILES = ([(rel, fam, absolute) for rel, fam, absolute in walk_files(SAMPLE_DIR)
               if rel.lower().endswith((".xlsx", ".xls")) and not _is_junk(os.path.basename(rel))]
              if os.path.isdir(SAMPLE_DIR) else [])
SITE_PARAMS = [pytest.param(rel, fam, absolute, id=_case_id(rel)) for rel, fam, absolute in SITE_FILES]


@pytest.fixture(scope="module")
def real_master(tmp_path_factory):
    """The real company master (customer lists + ERP lists), loaded once, with no earlier answers."""
    from excel import run
    answers = tmp_path_factory.mktemp("answers") / "company_decisions.json"
    answers.write_text("{}", encoding="utf-8")
    patch = pytest.MonkeyPatch()
    patch.setenv("COMPANY_DECISIONS_FILE", str(answers))
    run._master.cache_clear()
    master = run._master(run.IMPORTER_DIR)
    patch.setattr(run, "_master", functools.lru_cache(None)(lambda importer_dir: master))
    yield master
    patch.undo()
    run._master.cache_clear()


# KNOWN GAP, pinned so that fixing it is noticed: visits/pipeline.py add_reachability() patches a contact again on every
# run when the contact was created without an e-mail address because another person already holds it. Nothing else
# may be written by a second import (the golden pins how many of these each file repeats).
KNOWN_REPEATED_WRITES = {"contacts-patch"}


def _numbers(rows):
    """{label: value} of a family's lines, counts only (a money total or a name is not a count)."""
    return {r["label"]: r["value"] for r in rows if isinstance(r["value"], int)}


def through_the_site(site, rel, path):
    """What a person gets by putting this one file on the page: the row of the front door, the dry run, the first
    import and the second import. Counts and statuses only."""
    d = site.upload((path, rel))
    assert "no visit list recognised" not in json.dumps(d).lower(), rel
    (row,) = d["batch"]["entry"]
    batch = d["batch"]
    fam = next((f for f in batch["families"]), None)
    out = {"upload": {k: row[k] for k in ("kind", "status", "family", "folder", "what", "layout", "target", "records", "unit")}}
    if row["kind"] == "kept_out":
        out["upload"]["reason"] = row["reason"]
    if fam:
        out["upload"].update(ready=fam["ready"], waiting=fam["waiting"])
    if not batch["ready_families"]:  # kept out on purpose, a blank template, or every record waits for a company answer
        assert site.client.post("/api/commit").status_code == 400, rel
        assert site.client.post("/api/dry_run").status_code == 400, rel
        assert not site.crm.writes, rel
        return out

    writes, snapshot = len(site.crm.writes), site.crm.snapshot()
    plan = site.dry_run()
    assert len(site.crm.writes) == writes and site.crm.snapshot() == snapshot, f"the dry run wrote: {rel}"
    assert all(f["error"] == "" for f in plan["families"]), (rel, [f["error"] for f in plan["families"]])
    out["dry_run"] = {f["key"]: _numbers(f["rows"]) for f in plan["families"]}

    first = site.families(site.commit())
    out["commit"] = {k: {x: f[x] for x in ("status", "written", "present", "failed", "skipped")} for k, f in first.items()}
    assert all(f["failed"] == 0 and not f["errors"] for f in first.values()), (rel, {k: f["errors"] for k, f in first.items()})
    assert len(site.crm.writes) > writes, f"the import wrote nothing: {rel}"

    writes = len(site.crm.writes)
    second = site.families(site.commit())
    out["second_commit"] = {k: f["status"] for k, f in second.items()}
    repeated = Counter(kind for kind, _ in site.crm.writes[writes:])
    out["second_commit_writes"] = dict(sorted(repeated.items()))
    assert set(repeated) <= KNOWN_REPEATED_WRITES, f"the second import wrote something new: {rel}: {dict(repeated)}"
    return out


def test_every_file_of_the_folder_has_its_own_site_case(sample_dir):
    on_disk = _workbooks_on_disk(sample_dir)
    assert on_disk, "the folder holds no workbooks"
    cases = {rel for rel, _, _ in SITE_FILES}
    assert cases == on_disk, f"files without a case: {sorted(on_disk - cases)}; cases without a file: {sorted(cases - on_disk)}"
    assert len(SITE_PARAMS) == len(on_disk)
    assert len({p.id for p in SITE_PARAMS}) == len(SITE_PARAMS), "two files share a case id"
    golden = load_golden()
    if os.environ.get("REGEN_MATRIX") != "1":
        assert on_disk <= set(golden), f"files with no row in the golden: {sorted(on_disk - set(golden))}"
        assert {rel for rel in on_disk if "site" not in golden[rel]} == set(), "a golden row has no site result"


@pytest.mark.parametrize("rel,fam,path", SITE_PARAMS)
def test_file_goes_through_the_site(site, real_master, rel, fam, path):
    got = through_the_site(site, rel, path)
    up = got["upload"]
    # whatever the golden says, no real file may end up unplaced or unreadable
    assert up["status"] in ("ready", "kept_out", "empty"), f"{rel}: the front door answered {up['status']}"
    if up["status"] == "kept_out":  # only what the folder review keeps out, and with its reason
        assert up["folder"] in ("08", "09") and up["reason"], rel
    elif up["status"] == "ready":
        assert up["records"] > 0 and up["target"], rel
        assert up["folder"] == fam, f"{rel}: sent to folder {up['folder']} lists"
    if os.environ.get("REGEN_MATRIX") == "1":
        with GOLDEN_LOCK:
            golden = load_golden()
            golden.setdefault(rel, {"family": fam})["site"] = got
            save_golden(golden)
        return
    expected = load_golden().get(rel, {}).get("site")
    assert expected is not None, f"{rel} has no site result in the golden: REGEN_MATRIX=1"
    diff = [f"{k}: {expected.get(k)!r} -> {got.get(k)!r}" for k in sorted(set(expected) | set(got)) if expected.get(k) != got.get(k)]
    assert not diff, f"{rel} differs from the golden:\n" + "\n".join(diff)


def test_the_whole_folder_in_one_upload(site, real_master):
    """What a person does first: drops the whole folder. Every file is placed, the lists that belong in the CRM are
    imported in the safe order, one family failing would not block the others, and a second import adds nothing."""
    if not SITE_FILES:
        pytest.skip("sample folder not found")
    d = site.upload(*[(absolute, rel) for rel, _, absolute in SITE_FILES])
    assert "no visit list recognised" not in json.dumps(d).lower()
    entry = d["batch"]["entry"]
    assert len(entry) == len(SITE_FILES), "every file of the folder has a row in the front door's table"
    tally = Counter(r["status"] for r in entry)
    assert set(tally) <= {"ready", "kept_out", "empty", "duplicate", "superseded"}, tally
    assert all(r["reason"] for r in entry if r["status"] == "kept_out")
    got = {"files": len(entry), "status": dict(sorted(tally.items())), "ready_families": d["batch"]["ready_families"],
           "families": {f["key"]: {"files": len(f["files"]), "ready": f["ready"], "waiting": f["waiting"],
                                   "questions": len(f["questions"])} for f in d["batch"]["families"]}}

    before, writes = site.crm.snapshot(), len(site.crm.writes)
    plan = site.dry_run()
    assert site.crm.snapshot() == before and len(site.crm.writes) == writes, "the dry run wrote"
    assert [f["key"] for f in plan["families"]] == d["batch"]["ready_families"] and all(f["error"] == "" for f in plan["families"])

    job = site.commit()
    assert job["status"] == "done" and job["order"] == ["accounts", "contacts", "visits", "todos", "orders", "receivables"]
    first = site.families(job)
    got["commit"] = {k: {x: f[x] for x in ("status", "written", "present", "failed", "skipped")} for k, f in first.items()}
    assert all(f["status"] == "done" and f["failed"] == 0 for f in first.values()), {k: f["errors"] for k, f in first.items()}

    writes = len(site.crm.writes)
    second = site.families(site.commit())
    got["second_commit"] = {k: f["status"] for k, f in second.items()}
    repeated = Counter(kind for kind, _ in site.crm.writes[writes:])
    got["second_commit_writes"] = dict(sorted(repeated.items()))
    assert set(repeated) <= KNOWN_REPEATED_WRITES, dict(repeated)

    if os.environ.get("REGEN_MATRIX") == "1":
        with GOLDEN_LOCK:
            golden = load_golden()
            golden["_whole_folder"] = got
            save_golden(golden)
        return
    expected = load_golden().get("_whole_folder")
    assert expected is not None, "no whole-folder result in the golden: REGEN_MATRIX=1"
    diff = [f"{k}: {expected.get(k)!r} -> {got.get(k)!r}" for k in sorted(set(expected) | set(got)) if expected.get(k) != got.get(k)]
    assert not diff, "the whole-folder upload differs from the golden:\n" + "\n".join(diff)
