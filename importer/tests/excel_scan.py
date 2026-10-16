"""Structure-only scan of a folder of Excel workbooks, plus the small parsers the
test suite checks. Nothing here prints cell values. report_text() shows counts and
ids (#07); file names appear only with show_names=True.

Every per-file read catches its own errors and records them on FileInfo, so one bad
workbook never stops a scan.

Run from importer/:
  python -m tests.excel_scan "<folder>"                 report (no names)
  python -m tests.excel_scan "<folder>" --show-names    report with file names (local only)
  python -m tests.excel_scan "<folder>" --golden        counts-only JSON for tests/golden_real_folder.json
"""
from __future__ import annotations

import argparse
import hashlib
import json
import os
import re
import unicodedata
from collections import Counter
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

import openpyxl
import xlrd

from excel.sheet import open_workbook

ZIP_MAGIC = b"PK\x03\x04"
OLE_MAGIC = b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1"
IMPORTABLE = ("xlsx", "xls")
KINDS = ("xlsx", "xls", "lock", "hidden", "not_excel", "empty", "corrupt",
         "mislabelled", "encrypted", "unreadable")
CW_RE = re.compile(r"\bCW\s*0*(\d{1,2})\b", re.IGNORECASE)
HASH_PREFIX_RE = re.compile(r"^([0-9a-f]{12})_")
CURRENCY_RE = re.compile(r"[¥￥$€£元]|RMB|CNY|USD|EUR", re.IGNORECASE)
MONTH_DAY_RE = re.compile(r"(\d{1,2})月(\d{1,2})日?")
EXCEL_EPOCH = datetime(1899, 12, 30)


@dataclass
class SheetInfo:
    name: str
    hidden: bool
    rows: int
    cols: int
    nonempty: int
    merged: int
    cw: int | None


@dataclass
class FileInfo:
    path: str
    family: str
    kind: str
    size: int
    sha256: str = ""
    sheets: list[SheetInfo] = field(default_factory=list)
    error: str = ""

    @property
    def name(self) -> str:
        # macOS stores names in NFD; compare and print in NFC.
        return unicodedata.normalize("NFC", os.path.basename(self.path))

    @property
    def importable(self) -> bool:
        return self.kind in IMPORTABLE and not self.error

    @property
    def hash_prefix(self) -> str | None:
        m = HASH_PREFIX_RE.match(self.name)
        return m.group(1) if m else None


# ---------- discovery and classification ----------

def classify_name(name: str) -> str | None:
    """Kind decided by the name alone, or None when the content must be checked."""
    if name.startswith(".~lock.") or name.startswith("~$") or name.endswith("#"):
        return "lock"
    if name.startswith("."):
        return "hidden"
    if not name.lower().endswith((".xlsx", ".xls")):
        return "not_excel"
    return None


def sha256_of(path: str) -> str:
    h = hashlib.sha256()
    with open(path, "rb") as fh:
        for block in iter(lambda: fh.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def _stored_cells(ws):
    """Cells that exist in the file. iter_rows() walks the whole sheet dimension, which
    can be millions of empty styled cells (one 6.8 MB sample took 150 s); _cells holds
    only the stored ones. This is an openpyxl internal, stable across 3.x."""
    return list(ws._cells.values())


def _xlsx_sheets(path: str) -> list[SheetInfo]:
    out = []
    wb = open_workbook(path, data_only=True)
    try:
        for ws in wb.worksheets:
            nonempty = sum(1 for c in _stored_cells(ws) if c.value not in (None, ""))
            out.append(SheetInfo(ws.title, ws.sheet_state != "visible", ws.max_row, ws.max_column,
                                 nonempty, len(ws.merged_cells.ranges), cw_from_sheet_name(ws.title)))
    finally:
        wb.close()
    return out


def _xls_sheets(path: str) -> list[SheetInfo]:
    out = []
    book = xlrd.open_workbook(path, formatting_info=True)
    try:
        for sh in book.sheets():
            nonempty = sum(1 for r in range(sh.nrows) for c in range(sh.ncols)
                           if sh.cell_type(r, c) not in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK))
            out.append(SheetInfo(sh.name, sh.visibility != 0, sh.nrows, sh.ncols, nonempty,
                                 len(sh.merged_cells), cw_from_sheet_name(sh.name)))
    finally:
        book.release_resources()
    return out


def _fill(info: FileInfo, kind: str, reader) -> None:
    try:
        info.sheets = reader(info.path)
        info.kind = kind
    except Exception as exc:  # classify instead of stopping the scan
        text = str(exc).lower()
        info.kind = "encrypted" if ("encrypt" in text or "password" in text) else "corrupt"
        info.error = type(exc).__name__  # type name only: messages can echo sheet names


def scan_file(path, root=None) -> FileInfo:
    path = os.fspath(path)
    rel = os.path.relpath(path, root) if root else os.path.basename(path)
    parts = rel.split(os.sep)
    family = parts[0] if len(parts) > 1 else ""
    info = FileInfo(path=path, family=family, kind="", size=os.path.getsize(path))
    try:
        info.sha256 = sha256_of(path)
    except OSError:
        info.kind, info.error = "unreadable", "OSError"
        return info
    by_name = classify_name(os.path.basename(path))
    if by_name:
        info.kind = by_name
        return info
    if info.size == 0:
        info.kind = "empty"
        return info
    with open(path, "rb") as fh:
        head = fh.read(8)
    if path.lower().endswith(".xlsx"):
        if head.startswith(OLE_MAGIC):
            info.kind, info.error = "encrypted", "OLE container under .xlsx name"
        elif not head.startswith(ZIP_MAGIC):
            info.kind, info.error = "corrupt", "no zip signature"
        else:
            _fill(info, "xlsx", _xlsx_sheets)
    elif head.startswith(ZIP_MAGIC):
        info.kind, info.error = "mislabelled", "zip container under .xls name"
    else:
        _fill(info, "xls", _xls_sheets)
    return info


def scan_folder(root) -> list[FileInfo]:
    """Every file under root, sorted. Dot-folders (e.g. .arena, .venv) and a top-level import-tests folder are skipped;
    dot-files at any level are still reported as kind 'hidden' or 'lock'."""
    root = os.fspath(root)
    found = []
    for dirpath, dirnames, filenames in os.walk(root):
        dirnames[:] = sorted(d for d in dirnames if not d.startswith("."))
        if dirpath == root:  # our own output folder is not sample data
            dirnames[:] = [d for d in dirnames if d != "import-tests"]
        for fn in sorted(filenames):
            found.append(scan_file(os.path.join(dirpath, fn), root))
    return found


def duplicate_groups(files) -> dict[str, list[FileInfo]]:
    """Byte-identical files (same sha256), groups of two or more."""
    groups: dict[str, list[FileInfo]] = {}
    for f in files:
        if f.size > 0 and f.sha256:
            groups.setdefault(f.sha256, []).append(f)
    return {k: v for k, v in groups.items() if len(v) > 1}


# ---------- content comparison (not structure) ----------

def sheet_fingerprints(path: str) -> dict[str, str]:
    """sha256 of each sheet's cell values. The same sheet in two saved workbooks gives
    the same value even though the files differ (a re-sent weekly workbook)."""
    wb = open_workbook(path, read_only=True, data_only=True)
    try:
        out = {}
        for ws in wb.worksheets:
            rows = [tuple(r) for r in ws.iter_rows(values_only=True)]
            while rows and not any(v not in (None, "") for v in rows[-1]):
                rows.pop()  # trailing blank rows differ between saves
            h = hashlib.sha256()
            for row in rows:
                h.update(repr(row).encode("utf-8") + b"\n")
            out[ws.title] = h.hexdigest()
        return out
    finally:
        wb.close()


def formula_gaps(path: str) -> tuple[int, int]:
    """(formula cells, formula cells with no cached value) over all sheets."""
    formulas = open_workbook(path, data_only=False)
    values = open_workbook(path, data_only=True)
    try:
        total = missing = 0
        for ws in formulas.worksheets:
            vs = values[ws.title]
            for c in _stored_cells(ws):
                if isinstance(c.value, str) and c.value.startswith("="):
                    total += 1
                    if vs[c.coordinate].value is None:
                        missing += 1
        return total, missing
    finally:
        formulas.close()
        values.close()


# ---------- sheet names, text, parsers ----------

def cw_from_sheet_name(name: str) -> int | None:
    m = CW_RE.search(name)
    if not m:
        return None
    n = int(m.group(1))
    return n if 1 <= n <= 53 else None


def fold_text(text) -> str:
    """Comparable form of a label: full-width to half-width (NFKC), casefold, collapsed
    whitespace, accents dropped. '（已联络）' -> '(已联络)', 'Müller' -> 'muller'."""
    s = unicodedata.normalize("NFKC", str(text)).casefold()
    s = unicodedata.normalize("NFKD", s)
    s = "".join(ch for ch in s if not unicodedata.combining(ch))
    return " ".join(s.split())


def find_header_row(rows, required, max_scan: int = 20) -> int | None:
    """0-based index of the first row (within max_scan) that holds every required label."""
    want = {fold_text(r) for r in required}
    for i, row in enumerate(rows[:max_scan]):
        cells = {fold_text(v) for v in row if isinstance(v, str)}
        if want <= cells:
            return i
    return None


def _safe_date(y: int, m: int, d: int) -> date | None:
    try:
        return date(y, m, d)
    except ValueError:
        return None


def parse_cell_date(value, year: int | None = None) -> date | None:
    """Excel cell to date. Accepts datetime/date, 1900-system serials, ISO text,
    day-first dotted text (05.03.2026) and Chinese 'M月D日' (needs year).
    Slash text such as 5/3/2026 is ambiguous and returns None, never a guess."""
    if value is None or isinstance(value, bool) or value == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if isinstance(value, (int, float)):
        if not 61 <= value < 2958466:  # 60 is Excel's fake 1900-02-29
            return None
        return (EXCEL_EPOCH + timedelta(days=float(value))).date()
    text = unicodedata.normalize("NFKC", str(value)).strip()
    m = re.fullmatch(r"(\d{4})[-./](\d{1,2})[-./](\d{1,2})", text)
    if m:
        return _safe_date(int(m[1]), int(m[2]), int(m[3]))
    m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", text)
    if m:
        return _safe_date(int(m[3]), int(m[2]), int(m[1]))
    m = MONTH_DAY_RE.fullmatch(text)
    if m and year:
        return _safe_date(year, int(m[1]), int(m[2]))
    return None


def parse_number(value) -> float | None:
    """Number from a cell. Currency signs and % are handled, and so are both decimal
    conventions. When both separators appear, the last one is the decimal point:
    '1,234.5' and '1.234,5' both give 1234.5. A lone comma before exactly three digits
    is a thousands separator: '1,234' gives 1234. '12,5' gives 12.5. '12.5%' gives 0.125.
    '(120)' gives -120. Anything else gives None."""
    if value is None or isinstance(value, bool):
        return None
    if isinstance(value, (int, float)):
        return float(value)
    text = unicodedata.normalize("NFKC", str(value)).strip()
    text = CURRENCY_RE.sub("", text).replace(" ", "").replace(" ", "")
    if not text:
        return None
    negative = text.startswith("(") and text.endswith(")")
    text = text.strip("()")
    percent = text.endswith("%")
    text = text.rstrip("%")
    if "," in text and "." in text:
        if text.rfind(",") > text.rfind("."):
            text = text.replace(".", "").replace(",", ".")
        else:
            text = text.replace(",", "")
    elif "," in text:
        text = text.replace(",", "") if re.fullmatch(r"-?\d{1,3}(,\d{3})+", text) else text.replace(",", ".")
    if not re.fullmatch(r"-?\d+(\.\d+)?", text):
        return None
    num = float(text) / (100 if percent else 1)
    return -num if negative else num


# ---------- reporting (structure and counts only) ----------

def golden(files) -> dict:
    """Counts only. Written to tests/golden_real_folder.json and compared by the real-folder tests."""
    files = [f for f in files if f.kind != "lock"]  # a lock file only means someone has the workbook open
    ok = [f for f in files if f.importable]
    sheets = [s for f in ok for s in f.sheets]
    gaps = [formula_gaps(f.path) for f in ok if f.kind == "xlsx"]
    groups = duplicate_groups(files)
    return {
        "files": len(files),
        "kinds": dict(sorted(Counter(f.kind for f in files).items())),
        "importable": len(ok),
        "sheets": len(sheets),
        "hidden_sheets": sum(s.hidden for s in sheets),
        "empty_sheets": sum(s.nonempty == 0 for s in sheets),
        "cw_sheets": sum(s.cw is not None for s in sheets),
        "formula_cells": sum(t for t, _ in gaps),
        "formula_cells_without_cache": sum(m for _, m in gaps),
        "duplicate_groups": len(groups),
        "redundant_files": sum(len(g) - 1 for g in groups.values()),
    }


def report_text(files, show_names: bool = False) -> str:
    out = [f"files on disk: {len(files)}"]
    out += [f"  {kind:<12}{n}" for kind, n in sorted(Counter(f.kind for f in files).items())]
    ok = [f for f in files if f.importable]
    sheets = [s for f in ok for s in f.sheets]
    out.append(f"importable: {len(ok)} workbooks, {len(sheets)} sheets "
               f"(hidden {sum(s.hidden for s in sheets)}, empty {sum(s.nonempty == 0 for s in sheets)}, "
               f"CW-named {sum(s.cw is not None for s in sheets)}, "
               f"merged ranges {sum(s.merged for s in sheets)})")
    groups = duplicate_groups(files)
    idx = {id(f): i for i, f in enumerate(files)}
    out.append(f"duplicate groups: {len(groups)} (redundant copies: {sum(len(g) - 1 for g in groups.values())})")
    for sha, group in groups.items():
        out.append(f"  sha {sha[:8]}: " + " ".join(f"#{idx[id(f)]:02d}" for f in group))
    out.append("folders:")
    out += [f"  {fam or '(root)'}: {n}" for fam, n in sorted(Counter(f.family for f in files).items())]
    out.append("skipped or unreadable:")
    for i, f in enumerate(files):
        if f.importable:
            continue
        label = f.name if show_names else f"#{i:02d}"
        reason = f" ({f.error})" if f.error else ""
        out.append(f"  {label} {f.kind}{reason}")
    return "\n".join(out)


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description="Structure-only scan of a folder of Excel workbooks.")
    ap.add_argument("folder")
    ap.add_argument("--show-names", action="store_true", help="include file names (never share the output)")
    ap.add_argument("--golden", action="store_true", help="print counts-only JSON instead of the report")
    args = ap.parse_args(argv)
    files = scan_folder(args.folder)
    if args.golden:
        print(json.dumps(golden(files), indent=2, ensure_ascii=False, sort_keys=True))
    else:
        print(report_text(files, args.show_names))
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
