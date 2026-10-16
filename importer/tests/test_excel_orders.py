"""Sales order list reader, on synthetic workbooks (all names and numbers invented), plus a
structure-only run over the real export folder when it is present."""
import datetime as dt
import glob
import os

import openpyxl
import pytest

from excel.orders import OrdersError, looks_like_orders, parse_date, parse_number, read_orders
from excel.sheet import load

EN = ["Document No.", "Document Date", "Customer", "Salesperson", "Order Status", "Product Code",
      "Product Name", "Sales Quantity", "Sales Unit", "Amount Excl. Tax", "Tax Amount", "Amount Incl. Tax"]
D = dt.datetime(2026, 3, 5)


def line(doc="SO001", date=D, cust="Acme Test GmbH", sp="Pat", st="Open", code="P-1", name="Widget",
         qty=2, unit="pc", net=100, tax=13, gross=113):
    return [doc, date, cust, sp, st, code, name, qty, unit, net, tax, gross]


def book(tmp_path, sheets, name="orders.xlsx", hidden=(), merges=None):
    """sheets: {title: [row, row, ...]} written from A1."""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        for r in rows:
            ws.append(r)
        for rng in (merges or {}).get(title, ()):
            ws.merge_cells(rng)
        if title in hidden:
            ws.sheet_state = "hidden"
    p = str(tmp_path / name)
    wb.save(p)
    return p


def std(rows, title_rows=1):
    return [["Sales Order List"]] * title_rows + [EN] + rows


def test_english_basic(tmp_path):
    r = read_orders(book(tmp_path, {"S": std([line(), line("SO002", qty=1)])}))
    assert r.layout == "en" and len(r.lines) == 2 and not r.skipped
    ln = r.lines[0]
    assert (ln.document_no, ln.date_iso, ln.customer, ln.quantity, ln.net_amount, ln.gross_amount) == \
        ("SO001", "2026-03-05", "Acme Test GmbH", 2.0, 100.0, 113.0)
    assert ln.row == 3 and ln.unit == "pc" and ln.tax_amount == 13.0
    assert r.headers_found["S"]["document_no"] == "Document No."


def test_chinese_bilingual_headers_and_order_change(tmp_path):
    hdr = ["金额 无税金额 Net Value", "含税金额 Gross Amount", "销售数量 Sales Qty", "客户 Customer Name",
           "单据编号   Sales Order No.", "单据日期 Order Date", "商品编码 Artikle Code", "商品名称 Name",
           "销售业务员 Sales Rep", "订单状态 Order Status", "销售单位", "表体税额 Tax"]
    hdr[0] = "无税金额 Net Value"
    rows = [[100, 113, 2, "测试客户", "SO9", D, "X-1", "转台", "Li", "开立", "件", 13]]
    r = read_orders(book(tmp_path, {"S": [["销售订单列表"], hdr] + rows}))
    assert r.layout == "zh"
    ln = r.lines[0]
    assert (ln.document_no, ln.customer, ln.quantity, ln.net_amount, ln.gross_amount, ln.unit) == \
        ("SO9", "测试客户", 2.0, 100.0, 113.0, "件")


def test_german_headers_and_comma_text_numbers(tmp_path):
    hdr = ["Belegnummer", "Belegdatum", "Kunde", "Artikelnummer", "Menge", "Einheit", "Nettobetrag", "Bruttobetrag"]
    rows = [["A1", "05.03.2026", "Muster AG", "Z1", "3", "Stk", "1.234,50", "1.469,06"]]
    r = read_orders(book(tmp_path, {"S": [hdr] + rows}))
    ln = r.lines[0]
    assert r.layout == "de" and ln.date_iso == "2026-03-05"
    assert (ln.quantity, ln.net_amount, ln.gross_amount) == (3.0, 1234.5, 1469.06)


def test_fullwidth_case_and_punctuation_in_headers(tmp_path):
    hdr = ["ＤＯＣＵＭＥＮＴ  ＮＯ．", "document   DATE:", "CUSTOMER", "Sales  Quantity", "amount incl. tax"]
    r = read_orders(book(tmp_path, {"S": [hdr, ["SO1", D, "Test Co", 1, 5]]}))
    assert len(r.lines) == 1 and r.lines[0].gross_amount == 5.0


def test_dates_datetime_serial_text():
    assert parse_date(dt.datetime(2026, 1, 2, 5)) == "2026-01-02"
    assert parse_date(46086) == "2026-03-05"
    assert parse_date("2026/03/05") == "2026-03-05"
    assert parse_date("5.3.2026") == "2026-03-05"
    assert parse_date("soon") is None and parse_date(None) is None


def test_numbers():
    assert parse_number("1,234.50") == 1234.5
    assert parse_number("1.234,50") == 1234.5
    assert parse_number("1,234") == 1234.0
    assert parse_number("12,5") == 12.5
    assert parse_number("RMB 1 000") == 1000.0
    assert parse_number("(5)") == -5.0 and parse_number("5-") == -5.0 and parse_number("-3") == -3.0
    assert parse_number("abc") is None and parse_number("") is None and parse_number(True) is None


def test_text_and_bad_values_reported_not_crashing(tmp_path):
    rows = [line(qty="2", net="1,000.00", gross="n/a", date="not a date")]
    ln = read_orders(book(tmp_path, {"S": std(rows)})).lines[0]
    assert ln.quantity == 2.0 and ln.net_amount == 1000.0 and ln.gross_amount is None and ln.date_iso is None
    assert "gross_amount: not a number" in ln.issues and "date: not a date" in ln.issues


def test_duplicate_documents_grouped_with_totals(tmp_path):
    r = read_orders(book(tmp_path, {"S": std([line(), line(qty=3, net=50, tax=6, gross=56), line("SO2")])}))
    d = r.documents["SO001"]
    assert (d.line_count, d.quantity, d.net_amount, d.gross_amount) == (2, 5.0, 150.0, 169.0)
    assert len(r.documents) == 2 and len(r.lines) == 3


def test_negative_quantity_and_zero_amounts_kept(tmp_path):
    r = read_orders(book(tmp_path, {"S": std([line(qty=-2, net=-100, tax=-13, gross=-113), line("SO2", net=0, tax=0, gross=0)])}))
    assert r.lines[0].quantity == -2.0 and r.lines[0].gross_amount == -113.0
    assert r.lines[1].net_amount == 0.0 and r.lines[1].gross_amount == 0.0 and not r.lines[1].issues
    assert r.documents["SO001"].quantity == -2.0


def test_skips_blank_total_repeated_header_banner(tmp_path):
    rows = [line(), [], ["Total", None, None, None, None, None, None, 2, None, 100, 13, 113],
            ["合计"], EN, ["Note: all prices in test currency"], line("SO2")]
    r = read_orders(book(tmp_path, {"S": std(rows)}, merges={"S": ["A8:L8"]}))
    assert [ln.document_no for ln in r.lines] == ["SO001", "SO2"]
    why = {s.row: s.reason for s in r.skipped}
    assert why[4] == "blank row" and "total" in why[5] and "total" in why[6]
    assert why[7] == "repeated header row" and why[8] == "merged banner row"


def test_merged_document_cell_repeats_for_each_line(tmp_path):
    rows = [line("SO5"), line(None, code="P-2"), line(None, code="P-3")]
    r = read_orders(book(tmp_path, {"S": std(rows)}, merges={"S": ["A3:A5"]}))
    assert [ln.document_no for ln in r.lines] == ["SO5"] * 3 and r.documents["SO5"].line_count == 3


def test_row_without_document_number_skipped(tmp_path):
    r = read_orders(book(tmp_path, {"S": std([line(), line(None)])}))
    assert len(r.lines) == 1 and r.skipped[0].reason == "no document number"


def test_extra_columns_kept_by_header(tmp_path):
    hdr = EN + ["Remarks", "Remarks", "Weird Col"]
    r = read_orders(book(tmp_path, {"S": [hdr, line() + ["a", "b", 7]]}))
    assert r.lines[0].extra == {"Remarks": "a", "Remarks (2)": "b", "Weird Col": 7}


def test_currency_column_when_present(tmp_path):
    r = read_orders(book(tmp_path, {"S": [EN + ["Currency"], line() + ["EUR"]]}))
    assert r.lines[0].currency == "EUR"
    assert read_orders(book(tmp_path, {"S": std([line()])}, name="b.xlsx")).lines[0].currency is None


def test_missing_columns_warn(tmp_path):
    hdr = ["Document No.", "Customer", "Product Name"]
    r = read_orders(book(tmp_path, {"S": [hdr, ["SO1", "Test Co", "Thing"]]}))
    assert len(r.lines) == 1 and r.lines[0].quantity is None
    assert any("no column found for" in w and "quantity" in w and "gross_amount" in w for w in r.warnings)


def test_no_document_column_means_not_orders(tmp_path):
    p = book(tmp_path, {"S": [["Customer", "Product Name", "Sales Quantity"], ["Test Co", "Thing", 1]]})
    r = read_orders(p)
    assert not r.lines and r.layout is None
    assert any("no sheet looks like" in w for w in r.warnings)
    assert not looks_like_orders(load(p)[0])


def test_two_sheets_and_hidden_sheet(tmp_path):
    p = book(tmp_path, {"A": std([line("SO1")]), "B": std([line("SO1"), line("SO2")]), "H": std([line("SO3")]),
                        "Notes": [["just", "notes"]]}, hidden=("H",))
    r = read_orders(p)
    assert sorted(ln.sheet for ln in r.lines) == ["A", "B", "B"]
    assert ("H", "sheet is hidden") in r.skipped_sheets and "SO3" not in r.documents
    assert r.documents["SO1"].sheets == ["A", "B"]
    assert any("more than one sheet" in w for w in r.warnings)
    assert [looks_like_orders(s) for s in load(p)] == [True, True, True, False]


def test_header_below_preamble(tmp_path):
    pre = [["Query"], ["Status", "Open"], [], []]
    r = read_orders(book(tmp_path, {"S": pre + [EN, line()]}))
    assert len(r.lines) == 1 and r.lines[0].row == 6


def test_formula_without_cached_value_reported(tmp_path):
    rows = [line(), line("SO2")]
    rows[1][11] = "=J4+K4"  # openpyxl writes no cached value
    rows[0][0] = "=\"SO\"&\"1\""
    r = read_orders(book(tmp_path, {"S": std(rows)}))
    assert [s.reason for s in r.skipped] == ["document_no: formula without cached value"]
    assert r.lines[0].gross_amount is None and r.lines[0].issues == ["gross_amount: formula without cached value"]
    assert any("formula without a cached value" in w for w in r.warnings)


def test_wrong_type_empty_and_corrupt(tmp_path):
    txt = tmp_path / "a.txt"
    txt.write_text("hello")
    with pytest.raises(OrdersError, match="not an Excel"):
        read_orders(str(txt))
    empty = tmp_path / "e.xlsx"
    empty.write_bytes(b"")
    with pytest.raises(OrdersError, match="empty"):
        read_orders(str(empty))
    bad = tmp_path / "b.xlsx"
    bad.write_bytes(b"this is not a zip file")
    with pytest.raises(OrdersError, match="cannot open"):
        read_orders(str(bad))
    with pytest.raises(OrdersError, match="not found"):
        read_orders(str(tmp_path / "nope.xlsx"))


def test_workbook_without_orders_returns_empty_result(tmp_path):
    r = read_orders(book(tmp_path, {"S": [[]]}))
    assert not r.lines and r.warnings


# ---- real export folder: structure only, no values asserted ----
# EXCEL_SAMPLE_DIR is the folder that holds the nine family folders (as in the other real-folder tests).
SAMPLE_ROOT = os.environ.get("EXCEL_SAMPLE_DIR", "/Users/jonas/Mac Documents/Shanghai Company/CRM Excel Lists 2026-10-11")
SAMPLE_DIR = os.path.join(SAMPLE_ROOT, "08 Sales Order Exports")
GOLDEN_FILES = 9
GOLDEN_LINES = 1034


def test_real_sales_order_folder_structure():
    if not os.path.isdir(SAMPLE_DIR):
        pytest.skip("sample folder not present")
    files = sorted(glob.glob(os.path.join(SAMPLE_DIR, "*.xlsx")))
    assert len(files) == GOLDEN_FILES
    total = 0
    for p in files:
        r = read_orders(p)  # must not raise
        assert r.lines or r.warnings, os.path.basename(p)
        assert r.layout in (None, "en", "zh", "de")
        total += len(r.lines)
    assert total == GOLDEN_LINES


# ======================================================================================
# What a file is, the book of orders (newest export wins), and the Margince adapter
# ======================================================================================
import copy
import json
import shutil

from excel import orders_margince as om
from excel.companies import Master
from excel.orders import FORMULA_REASON, INVOICE_REASON, STOCK_REASON, build_book, status_key, summarize
from tests.fake_crm import FakeCRM

# a company list with one entry that looks like none of the invented customers below
MASTER = [{"source": "master", "code": "KD00001", "cn": "宇宙星辰科技有限公司", "en": "Cosmos Stellar Co., Ltd.", "short": ""}]
FULL = EN + ["Sales Organization", "Amount Incl. Tax in Functional Currency", "Amount Excl. Tax in Functional Currency"]
CN_ORG, US_ORG = "测试精密部件（上海）有限公司", "Test Rotary US Inc."


def master():
    return Master.from_records(MASTER)


def fline(doc="SO001", date=D, cust="Acme Test GmbH", status="Open", code="P-1", name="Widget", qty=2,
          net=100, tax=13, gross=113, org=CN_ORG, loc_gross=None, loc_net=None, sp="Pat"):
    return line(doc, date, cust, sp, status, code, name, qty, "pc", net, tax, gross) + [
        org, gross if loc_gross is None else loc_gross, net if loc_net is None else loc_net]


def export(tmp_path, name, rows, stamp=None, sheet="Sales Order List", header=FULL):
    top = [["Sales Order List"]] if stamp is None else [["Sales Order List"], ["Export time", stamp]]
    return book(tmp_path, {sheet: top + [header] + rows}, name=name)


def test_status_words():
    assert status_key("Open") == "open" and status_key("开立") == "open"
    assert status_key("待发货 Pending Delivery") == "to_ship" and status_key("To Ship") == "to_ship"
    assert status_key("已完成 Order Completed") == "completed" and status_key("Completed") == "completed"
    assert status_key("Partially Received") == "in_progress" and status_key("To Receive") == "in_progress"
    assert status_key("Cancelled") == "cancelled" and status_key("whatever") == "other" and status_key(None) == "other"


# ---- what a file is: only the ERP sales-order lists are CRM data ----

def invoice_book(tmp_path, name="2026-03-09_-_317_Fibro_Order_List.xlsx"):
    sheets = {"Invoice": [["Test Rotary US Inc."], [None, None, None, None, "Invoice #:", "INV-TEST-1"],
                          ["Item", "Qty", "Unit price", "Amount"], ["Gadget", 2, 50, 100]],
              "Packing List Crate 1": [["Test Rotary US Inc."], [None, None, None, None, "Invoice #:", " "],
                                       ["Crate", "Weight"], ["C-1", 55]],
              "Original Stock": [[None, "TEST ROTARY US INC."], [None, "Inventory - December 2025"], ["Part", "Qty"]]}
    return book(tmp_path, sheets, name=name)


def test_invoice_and_packing_list_named_order_list_is_excluded_by_content(tmp_path):
    r = read_orders(invoice_book(tmp_path))
    assert r.kind == "invoice_packing" and r.reason == INVOICE_REASON and not r.lines
    assert {w for _, w in r.excluded_sheets} == {INVOICE_REASON, STOCK_REASON}


def test_stock_list_alone_is_excluded(tmp_path):
    p = book(tmp_path, {"Original Stock": [["TEST INC"], ["Inventory - December 2025"], ["Part", "Qty"], ["X", 1]],
                        "Sheet1": [["P-1", 25.26]]})
    r = read_orders(p)
    assert r.kind == "stock_list" and r.reason == STOCK_REASON


def test_workbook_with_formulas_over_a_copy_of_the_list_is_excluded(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    calc = wb.create_sheet("Actual")
    calc["A1"] = "Actual 2026 from the Sales Order List"
    for i in range(2, 8):
        calc[f"B{i}"] = f"=SUMIF('Sales Order List'!B:B,A{i},'Sales Order List'!L:L)"
    ws = wb.create_sheet("Sales Order List")
    for row in [["Sales Order List"], FULL, fline("SO1")]:
        ws.append(row)
    p = str(tmp_path / "Actual_calc.xlsx")
    wb.save(p)
    r = read_orders(p)
    assert r.kind == "formula_workbook" and r.reason == FORMULA_REASON
    assert len(r.lines) == 1  # the lines are still readable; callers import only kind == "orders"
    assert build_book([p], master()).orders == []
    assert build_book([p], master()).files[0].reason == FORMULA_REASON


def test_order_file_with_a_side_sheet_stays_orders_and_names_the_side_sheet(tmp_path):
    p = book(tmp_path, {"S": std([line()]), "Notes": [["Packing List", "x"]]})
    r = read_orders(p)
    assert r.kind == "orders" and r.excluded_sheets == [("Notes", INVOICE_REASON)]


def test_export_stamp_read_from_above_the_header(tmp_path):
    assert read_orders(export(tmp_path, "a.xlsx", [fline()], stamp="2026-05-01 10:00:00")).export_time == dt.datetime(2026, 5, 1, 10)
    assert read_orders(export(tmp_path, "b.xlsx", [fline()])).export_time is None


# ---- one order per document, newest export wins ----

def test_lines_roll_up_to_one_order_with_items(tmp_path):
    p = export(tmp_path, "a.xlsx", [fline(code="P-1"), fline(code="P-2", net=50, tax=6, gross=56), fline("SO2")])
    b = build_book([p], master())
    assert [o.document_no for o in b.orders] == ["SO001", "SO2"]
    o = b.orders[0]
    assert (o.net, o.tax, o.gross, len(o.items)) == (150.0, 19.0, 169.0, 2)
    assert o.status_key == "open" and o.entity == CN_ORG and o.source_id == "order:SO001"
    assert summarize(b)["orders"] == 2 and summarize(b)["lines"] == 3


def test_newest_export_wins_and_older_only_documents_stay(tmp_path):
    old = export(tmp_path, "old.xlsx", [fline("SO1", status="Open"), fline("SO2", status="Open")], stamp="2026-03-01 10:00:00")
    new = export(tmp_path, "new.xlsx", [fline("SO1", status="Completed", net=120, gross=135.6), fline("SO3")],
                 stamp="2026-04-01 10:00:00")
    b = build_book([old, new], master())
    got = {o.document_no: o for o in b.orders}
    assert sorted(got) == ["SO1", "SO2", "SO3"]
    assert got["SO1"].status_key == "completed" and got["SO1"].net == 120.0 and got["SO1"].file == "new.xlsx"
    assert got["SO1"].superseded == ["old.xlsx"] and got["SO2"].file == "old.xlsx"
    assert build_book([new, old], master()).orders[0].file == got["SO1"].file  # file order does not matter


def test_export_without_stamp_is_dated_by_its_newest_order(tmp_path):
    old = export(tmp_path, "old.xlsx", [fline("SO1", date=dt.datetime(2026, 1, 5), status="Open")])
    new = export(tmp_path, "new.xlsx", [fline("SO1", date=dt.datetime(2026, 1, 5), status="Completed"),
                                         fline("SO9", date=dt.datetime(2026, 2, 1))])
    b = build_book([old, new], master())
    assert {o.document_no: o.status_key for o in b.orders}["SO1"] == "completed"


def test_sheets_that_repeat_the_total_list_count_each_line_once(tmp_path):
    total = [fline("SO1", code="P-1"), fline("SO1", code="P-2"), fline("SO2")]
    p = book(tmp_path, {"All": [["x"], FULL] + total, "Completed": [["x"], FULL, total[0]],
                        "Open": [["x"], FULL, total[1], total[2]]})
    b = build_book([p], master())
    assert [(o.document_no, len(o.items)) for o in b.orders] == [("SO1", 2), ("SO2", 1)]


def test_lines_only_one_sheet_holds_are_kept(tmp_path):
    p = book(tmp_path, {"A": [["x"], FULL, fline("SO1", code="P-1")], "B": [["x"], FULL, fline("SO1", code="P-2")]})
    assert len(build_book([p], master()).orders[0].items) == 2


def test_byte_identical_copies_are_read_once(tmp_path):
    p = export(tmp_path, "a.xlsx", [fline()], stamp="2026-04-01 10:00:00")
    shutil.copy(p, tmp_path / "0123456789ab_a.xlsx")
    b = build_book([p, str(tmp_path / "0123456789ab_a.xlsx")], master())
    assert len(b.orders) == 1 and len(b.files) == 1 and "identical" in b.skipped_files[0]["error"]


def test_hash_prefixed_copy_with_the_same_stamp_richer_copy_wins(tmp_path):
    rich = export(tmp_path, "0123456789ab_list.xlsx", [fline("SO1", sp="Pat", name="Full widget")], stamp="2026-04-01 10:00:00")
    poor = export(tmp_path, "ba9876543210_list.xlsx", [fline("SO1", sp=None, name=None)], stamp="2026-04-01 10:00:00")
    for order in ([rich, poor], [poor, rich]):
        o = build_book(order, master()).orders[0]
        assert o.salesperson == "Pat" and o.items[0].product_name == "Full widget" and o.file == "list.xlsx"


def test_blank_copy_never_replaces_or_creates_an_order(tmp_path):
    good = export(tmp_path, "good.xlsx", [fline("SO1")], stamp="2026-03-01 10:00:00")
    blank = export(tmp_path, "blank.xlsx", [fline("SO1", cust=None, net=None, tax=None, gross=None, loc_gross=None),
                                             fline("SO2", cust=None, net=None, tax=None, gross=None)],
                   stamp="2026-04-01 10:00:00")
    b = build_book([good, blank], master())
    assert [o.document_no for o in b.orders] == ["SO1"] and b.orders[0].customer == "Acme Test GmbH"
    assert [r.value for r in b.review if r.reason == "blank_order"] == ["SO2"]


def test_zero_value_order_is_an_order_not_a_blank(tmp_path):
    b = build_book([export(tmp_path, "z.xlsx", [fline("SO1", net=0, tax=0, gross=0, loc_gross=0, loc_net=0)])], master())
    assert [o.document_no for o in b.orders] == ["SO1"] and not any(r.reason == "blank_order" for r in b.review)


def test_same_number_from_two_selling_companies_is_two_orders(tmp_path):
    p = export(tmp_path, "a.xlsx", [fline("SO1", org=CN_ORG), fline("SO1", org=US_ORG, cust="Other Test Inc.")])
    b = build_book([p], master())
    assert len(b.orders) == 2 and len({o.source_id for o in b.orders}) == 2


def test_unreadable_file_is_reported_and_the_run_goes_on(tmp_path):
    bad = tmp_path / "bad.xlsx"
    bad.write_bytes(b"not a zip")
    b = build_book([str(bad), export(tmp_path, "a.xlsx", [fline()])], master())
    assert len(b.orders) == 1 and "cannot open" in b.skipped_files[0]["error"]


# ---- currency: proven by the export or not stated, never guessed ----

def test_chinese_entity_local_columns_are_cny_and_foreign_orders_are_marked(tmp_path):
    rows = [fline("SO1", net=100, tax=13, gross=113),  # ordered in the local currency
            fline("SO2", net=1000, tax=130, gross=1130, loc_gross=8960, loc_net=7930)]  # another currency
    o = {x.document_no: x for x in build_book([export(tmp_path, "a.xlsx", rows)], master()).orders}
    assert (o["SO1"].currency, o["SO1"].cny_gross, o["SO1"].cny_net) == ("CNY", 113.0, 100.0)
    assert (o["SO2"].currency, o["SO2"].cny_gross, o["SO2"].gross) == (None, 8960.0, 1130.0)
    body = om.order_body(o["SO2"])
    assert "CNY 8,960.00 incl. tax" in body and "does not name it" in body and "no exchange rate" in body


def test_other_entity_has_no_currency_and_its_functional_columns_are_not_trusted(tmp_path):
    o = build_book([export(tmp_path, "a.xlsx", [fline("SO1", org=US_ORG, loc_gross=99999, loc_net=99999)])], master()).orders[0]
    assert o.currency is None and o.cny_gross is None and o.cny_net is None
    assert "currency not stated in the export" in om.order_body(o) and "CNY" not in om.order_subject(o)
    assert any(r.reason == "currency_not_stated" for r in build_book([export(tmp_path, "b.xlsx", [fline(org=US_ORG)])], master()).review)


def test_document_total_repeated_on_every_line_is_counted_once(tmp_path):
    rows = [fline("SO1", code="P-1", net=100, gross=113, loc_gross=300, loc_net=265),
            fline("SO1", code="P-2", net=100, gross=113, loc_gross=300, loc_net=265),
            fline("SO1", code="P-3", net=100, gross=113, loc_gross=300, loc_net=265)]
    o = build_book([export(tmp_path, "a.xlsx", rows)], master()).orders[0]
    assert (o.cny_gross, o.cny_net) == (300.0, 265.0)


# ---- the Margince adapter, against a fake CRM ----

class OrdersCRM(FakeCRM):
    """FakeCRM plus the open deals the overview note reads."""

    def __init__(self):
        super().__init__()
        self.deals = []

    def paginate(self, path, params=None):
        if path == "/v1/deals":
            yield from [d for d in self.deals if d.get("status") == "open"]
            return
        yield from super().paginate(path, params)


def snapshot(crm):
    return copy.deepcopy((crm.companies, crm.activities, crm.contacts))


def notes(crm, prefix="order:"):
    return {a["source_id"]: a for a in crm.activities.values() if a["source_id"].startswith(prefix)}


def two_exports(tmp_path):
    old = export(tmp_path, "old.xlsx", [fline("SO1", status="Open"), fline("SO2", cust="Beta Test Ltd.")],
                 stamp="2026-03-01 10:00:00")
    new = export(tmp_path, "new.xlsx", [fline("SO1", status="Completed")], stamp="2026-04-01 10:00:00")
    return old, new


def test_commit_writes_an_account_a_note_per_order_and_one_overview(tmp_path):
    crm = OrdersCRM()
    b = build_book(list(two_exports(tmp_path)), master())
    r = om.commit(crm, b)
    assert r["orders_created"] == 2 and r["overviews_created"] == 2 and not r["errors"]
    assert sorted(c["display_name"] for c in crm.companies.values()) == ["Acme Test GmbH", "Beta Test Ltd."]
    assert all(c["lifecycle"] == "customer" and "customer" in c["relationship_types"] for c in crm.companies.values())
    one = notes(crm)["order:SO1"]
    assert one["kind"] == "note" and one["source_system"] == om.SOURCE_SYSTEM and one["links"][0]["entity_type"] == "company"
    assert one["subject"].startswith("ERP order SO1 - Completed - CNY 113.00 incl. tax")
    assert "Read-only" in one["body"] and "ERP export: 2026-04-01 10:00:00" in one["body"] and "Widget" in one["body"]
    assert one["occurred_at"].startswith("2026-03-05T12:00")
    assert len(notes(crm, "orders-summary:")) == 2


def test_second_run_writes_nothing(tmp_path):
    crm = OrdersCRM()
    b = build_book(list(two_exports(tmp_path)), master())
    om.commit(crm, b)
    before = snapshot(crm)
    r = om.commit(crm, build_book(list(two_exports(tmp_path)), master()))
    assert (r["orders_created"], r["orders_updated"], r["overviews_created"], r["overviews_updated"]) == (0, 0, 0, 0)
    assert r["orders_unchanged"] == 2 and r["overviews_unchanged"] == 2
    assert snapshot(crm) == before


def test_preview_counts_and_writes_nothing(tmp_path):
    crm = OrdersCRM()
    b = build_book(list(two_exports(tmp_path)), master())
    before = snapshot(crm)
    p = om.plan(crm, b)
    assert p["orders"]["create"] == 2 and p["overviews"]["create"] == 2 and p["companies_new"] == 2
    assert snapshot(crm) == before
    om.commit(crm, b)
    p = om.plan(crm, b)
    assert p["orders"] == {"create": 0, "update": 0, "same": 2, "newer": 0} and p["companies_present"] == 2


def test_a_later_export_updates_only_what_changed(tmp_path):
    crm = OrdersCRM()
    old, new = two_exports(tmp_path)
    om.commit(crm, build_book([old], master()))
    assert "Open" in notes(crm)["order:SO1"]["subject"]
    r = om.commit(crm, build_book([old, new], master()))
    assert (r["orders_updated"], r["orders_created"], r["orders_unchanged"]) == (1, 0, 1)
    assert "Completed" in notes(crm)["order:SO1"]["subject"]
    assert len(notes(crm)) == 2


def test_an_older_export_never_overwrites_a_newer_one(tmp_path):
    crm = OrdersCRM()
    old, new = two_exports(tmp_path)
    om.commit(crm, build_book([old, new], master()))
    r = om.commit(crm, build_book([old], master()))
    assert r["orders_updated"] == 0 and r["orders_kept_newer"] == 1
    assert "Completed" in notes(crm)["order:SO1"]["subject"]
    assert r["overviews_updated"] == 0  # the old run knows fewer orders than the overview already counts


def test_orders_of_a_company_nobody_has_answered_yet_stay_out(tmp_path):
    crm = OrdersCRM()
    similar = "Cosmic Stellar Company Ltd"  # looks like the master's entry without being it
    p = export(tmp_path, "a.xlsx", [fline("SO1"), fline("SO2", cust=similar)])
    b = build_book([p], master())
    assert any(r.reason == "company" and r.value == similar for r in b.review)
    r = om.commit(crm, b)
    assert r["orders_created"] == 1 and r["orders_waiting_for_a_company_answer"] == 1
    assert list(notes(crm)) == ["order:SO1"] and len(crm.companies) == 1


def test_group_company_orders_are_flagged_like_receivables(tmp_path):
    p = export(tmp_path, "a.xlsx", [fline("SO1"), fline("SO2", cust="Fibro Rundtische GmbH")])
    b = build_book([p], master())
    assert [(r.reason, r.value) for r in b.review if r.reason == "group_company"] == [("group_company", "Fibro Rundtische GmbH")]


def test_overview_sets_orders_against_open_deals_without_converting(tmp_path):
    crm = OrdersCRM()
    b = build_book([export(tmp_path, "a.xlsx", [fline("SO1", status="Completed"), fline("SO2", gross=226, net=200),
                                                 fline("SO3", org=US_ORG)])], master())
    om.commit(crm, b)
    cid = next(iter(crm.companies))
    crm.deals = [{"company_id": cid, "status": "open", "amount_minor": 5000000, "currency": "CNY"},
                 {"company_id": cid, "status": "open", "amount_minor": 120000, "currency": "EUR"},
                 {"company_id": "someone-else", "status": "open", "amount_minor": 1, "currency": "CNY"}]
    r = om.commit(crm, b)
    body = next(iter(notes(crm, "orders-summary:").values()))["body"]
    assert r["overviews_updated"] == 1 and r["orders_unchanged"] == 3
    assert "2026: 3 orders, CNY 300.00 excl. tax; 1 not completed (CNY 200.00); 1 in a currency the export does not name" in body
    assert "Open deals: 1 in CNY: 50,000.00; 1 in EUR: 1,200.00" in body
    assert "not added up or converted" in body and "Orders counted: SO1, SO2" in body


def test_cancelled_orders_are_not_in_the_overview(tmp_path):
    body = om.overview_body(build_book([export(tmp_path, "a.xlsx", [fline("SO1"), fline("SO2", status="Cancelled")])], master()).orders, [])
    assert "1 order," in body and "SO2" not in body


def test_one_bad_note_does_not_stop_the_rest(tmp_path):
    from common.client import ApiError
    crm = OrdersCRM()
    real = crm.log_activity

    def flaky(body):
        if body["source_id"] == "order:SO1":
            import requests
            resp = requests.Response()
            resp.status_code = 422
            resp._content = b'{"code": "validation_error"}'
            resp.request = requests.Request("POST", "http://fake/v1/activities").prepare()
            raise ApiError(resp)
        return real(body)

    crm.log_activity = flaky
    r = om.commit(crm, build_book([export(tmp_path, "a.xlsx", [fline("SO1"), fline("SO2")])], master()))
    assert r["orders_created"] == 1 and len(r["errors"]) == 1 and r["errors"][0].startswith("SO1")


# ---- what must never reach a record or a payload ----

SECRETS = ["6222-0000-1111-2222", "TAXID-91310000TEST", "银行账号-OPAQUE", "INV-TEST-1", "PACKING-CRATE-SECRET"]


def test_no_bank_account_tax_id_or_invoice_value_reaches_any_record_or_payload(tmp_path):
    header = FULL + ["Bank Account", "Customer Tax ID", "纳税人识别号", "Remarks"]
    rows = [fline("SO1") + [SECRETS[0], SECRETS[1], SECRETS[2], "REMARK-NOT-WHITELISTED"]]
    orders_file = export(tmp_path, "orders.xlsx", rows, stamp="2026-04-01 10:00:00", header=header)
    inv = book(tmp_path, {"Invoice": [["Test Rotary US Inc."], [None, None, None, None, "Invoice #:", SECRETS[3]],
                                      ["Item", "Amount"], ["PACKING-CRATE-SECRET", 777777]],
                          "Packing List Crate 1": [["Packing List"], ["PACKING-CRATE-SECRET", 5]]}, name="x_Order_List.xlsx")
    crm = OrdersCRM()
    b = build_book([orders_file, inv], master())
    om.commit(crm, b)
    everything = json.dumps([b.orders, b.review, [vars(f) for f in b.files], summarize(b), crm.companies,
                             crm.activities, om.plan(crm, b)], default=lambda o: getattr(o, "__dict__", str(o)),
                            ensure_ascii=False)
    for secret in SECRETS + ["777777"]:
        assert secret not in everything, secret
    assert "Widget" in everything and "REMARK-NOT-WHITELISTED" not in json.dumps(crm.activities)  # only whitelisted columns travel
    read = read_orders(orders_file)
    assert not any(k in ("Bank Account", "Customer Tax ID", "纳税人识别号") for ln in read.lines for k in ln.extra)
    assert sum("not read (bank account or tax id)" in w for w in read.warnings) == 3


# ---- the real export folder: counts only ----

def test_real_folder_classification_and_book():
    if not os.path.isdir(SAMPLE_DIR):
        pytest.skip("sample folder not present")
    paths = sorted(glob.glob(os.path.join(SAMPLE_DIR, "*.xlsx")))
    kinds = {}
    for p in paths:
        kinds.setdefault(read_orders(p).kind, []).append(os.path.basename(p))
    assert {k: len(v) for k, v in kinds.items()} == {"orders": 5, "invoice_packing": 2, "formula_workbook": 1, "stock_list": 1}
    assert kinds["formula_workbook"] == ["Actual_mit_Formeln.xlsx"]
    b = build_book(paths, master())
    s = summarize(b)
    assert (s["files"], s["files_read"], s["files_excluded"], s["lines"]) == (9, 5, 4, 857)
    assert s["orders"] == 202 and not s["skipped_files"]  # 200 with an amount + 2 zero-value orders
    assert not [o for o in b.orders if o.file == "Actual_mit_Formeln.xlsx"]
    assert {o.status_key for o in b.orders} <= {"open", "to_ship", "in_progress", "completed"}


def test_order_without_a_date_is_asked_about_and_not_written(tmp_path):
    crm = OrdersCRM()
    b = build_book([export(tmp_path, "a.xlsx", [fline("SO1", date=None), fline("SO2")])], master())
    assert [r.value for r in b.review if r.reason == "no_date"] == ["SO1"]
    r = om.commit(crm, b)
    assert r["orders_created"] == 1 and list(notes(crm)) == ["order:SO2"] and not r["errors"]


def test_file_kind_for_the_wiring(tmp_path):
    from excel.orders import file_kind
    assert file_kind(export(tmp_path, "a.xlsx", [fline()])) == ("orders", "")
    assert file_kind(invoice_book(tmp_path)) == ("invoice_packing", INVOICE_REASON)


def test_a_line_that_equals_the_sum_of_the_others_holds_the_order_out(tmp_path):
    """SO-style double count: one line of 220,000 next to four identical lines of 55,000."""
    from excel.orders import _order
    from types import SimpleNamespace as NS
    def line(amount):
        return NS(customer="Acme Co., Ltd.", net_amount=amount, tax_amount=0.0, gross_amount=amount, currency="CNY", date_iso="2026-01-02",
                  salesperson="", status="", product_code="x", product_name="x", quantity=1.0, unit="pc", row=1, sheet="s", extra={})
    info = NS(name="f.xlsx", export_iso=None)
    clean = _order("SO1", [line(55000.0), line(55000.0), line(30000.0)], info)
    dup = _order("SO2", [line(55000.0)] * 4 + [line(220000.0)], info)
    assert clean.issues == []
    assert any("double count" in i for i in dup.issues)


def test_two_exports_with_two_amounts_hold_the_order_out(tmp_path):
    old = export(tmp_path, "old.xlsx", [fline("SO1"), fline("SO2")], stamp="2026-03-01 10:00:00")
    new = export(tmp_path, "new.xlsx", [fline("SO1", net=120, gross=135.6), fline("SO2")], stamp="2026-04-01 10:00:00")
    got = {o.document_no: o for o in build_book([old, new], master()).orders}
    assert any("exports disagree" in i for i in got["SO1"].issues)
    assert got["SO2"].issues == []  # the same amount in both exports is no doubt


def test_order_of_another_sales_entity_is_not_attached_to_a_china_customer_by_name(tmp_path):
    p = export(tmp_path, "a.xlsx", [fline("SO1", cust="Cosmos Stellar", org=US_ORG), fline("SO2", cust="Cosmos Stellar", org=CN_ORG)])
    got = {o.document_no: o for o in build_book([p], master()).orders}
    assert any("another sales entity" in i for i in got["SO1"].issues)
    assert got["SO2"].issues == []
