"""Writes tiny synthetic visit workbooks for the excel reader tests. Every value
here is made up; no real customer data goes in tests."""
import os
import tempfile
import unittest

import openpyxl

WEEKLY_HEADER = {
    "B8": "Sales", "C8": "CW", "D8": "Date", "E8": "Location/City",
    "F8": "Customer \n(full address if new)", "G8": "Potential\n(A, B, C)",
    "H8": "Customer \n(new, exist, dealer)", "I8": "Contact Name", "J8": "Phone No.",
    "K8": "WeChat", "L8": "Email", "M8": "Industries", "N8": "Purpose_Topics Visit",
    "O8": "Competiton Products", "P8": "Next Step / Action", "Q8": "Due Date", "R8": "Status",
}

LUCAS_HEADER = {
    "B14": "Visiting Customer ", "C14": "Date", "D14": "Location/City", "E14": "Customer",
    "F14": "Potential\n(A, B, C)", "G14": "Customer \n(new, exist, dealer)", "H14": "Name",
    "I14": "Phone", "J14": "Office ", "K14": "Email", "L14": "Purpose for Visiting ",
    "M14": "Next Step / Action", "N14": "follow up", "O14": "deadline ",
}


def write_book(path, sheets):
    """sheets: {title: {"cells": {"B8": value, ...}, "merges": ["F9:F12"]}}"""
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, spec in sheets.items():
        ws = wb.create_sheet(title)
        for coord, value in spec.get("cells", {}).items():
            ws[coord] = value
        for rng in spec.get("merges", ()):
            ws.merge_cells(rng)
    wb.save(path)
    return path


class TmpCase(unittest.TestCase):
    def setUp(self):
        tmp = tempfile.TemporaryDirectory()
        self.addCleanup(tmp.cleanup)
        self.tmp = tmp.name
        # The invented customers of these tests are answered "new company", as a person would
        # in the review step; otherwise they only look like a weak match of a real customer.
        import json
        from unittest import mock

        from excel import run
        decisions = os.path.join(self.tmp, "company_decisions.json")
        with open(decisions, "w", encoding="utf-8") as f:
            json.dump({"Acme Test Co.": None, "Beta Test Co.": None}, f)
        env = mock.patch.dict(os.environ, {"COMPANY_DECISIONS_FILE": decisions})
        env.start()
        self.addCleanup(env.stop)
        run._master.cache_clear()
        self.addCleanup(run._master.cache_clear)

    def path(self, name):
        return os.path.join(self.tmp, name)
