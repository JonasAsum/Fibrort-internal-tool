"""The site's front door and import flow, end to end through the Flask client with an in-memory CRM.
Every workbook here is a tiny invented one (webui/selftest.py builds one per family)."""
import functools
import json
import os

import openpyxl
import pytest

from excel import run as excel_run
from excel.companies import Master
from tests.fake_site_crm import SiteCRM
from tests.site_fixture import site  # noqa: F401  (the fixture)
from webui import selftest, server

FAMILY_FILES = {  # the family key each selftest workbook must be sent to, and what its upload row must say
    "visits": ("visits", "01", 2), "accounts": ("accounts", "04", 2), "contacts": ("contacts", "06", 2),
    "todos": ("todos", "07", 3), "orders": ("orders", "08", 2), "receivables": ("receivables", "09", 2),
}


@pytest.fixture
def tiny(monkeypatch):
    """A company master with one invented customer instead of the real customer lists."""
    master = Master.from_records(selftest.MASTER)
    monkeypatch.setattr(excel_run, "_master", functools.lru_cache(None)(lambda importer_dir: master))
    return master


@pytest.fixture
def books(tmp_path):
    return {k: p for k, (p, _) in selftest.make_workbooks(str(tmp_path)).items()}


def upload_all(site, books, *names):
    return site.upload(*[(books[k], os.path.basename(books[k])) for k in (names or books)])


def invoice_list(tmp_path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    wb.create_sheet("Invoice").append(["Item", "Qty", "Unit price", "Amount"])
    wb.create_sheet("Packing List Crate 1").append(["Crate", "Weight"])
    wb.create_sheet("Original Stock").append(["Part", "Qty"])
    path = str(tmp_path / "2026-03-09_-_317_Fibro_Order_List.xlsx")
    wb.save(path)
    return path


@pytest.mark.parametrize("key", list(FAMILY_FILES))
def test_every_family_is_recognised_and_routed_by_its_content(site, tiny, books, key):
    family, folder, found = FAMILY_FILES[key]
    d = site.upload((books[key], os.path.basename(books[key])))
    (row,) = d["batch"]["entry"]
    assert (row["status"], row["family"], row["folder"]) == ("ready", family, folder)
    assert row["records"] > 0 and row["target"] and row["outcome"]
    assert "no visit list recognised" not in json.dumps(d).lower()
    assert d["batch"]["ready_families"] == [family]


def test_a_workbook_is_judged_by_what_is_in_it_not_by_its_name(site, tiny, books, tmp_path):
    renamed = tmp_path / "list.xlsx"
    renamed.write_bytes(open(books["contacts"], "rb").read())
    (row,) = site.upload((str(renamed), "list.xlsx"))["batch"]["entry"]
    assert (row["family"], row["status"]) == ("contacts", "ready")


def test_the_folder_a_browser_reports_does_not_decide_either(site, tiny, books):
    (row,) = site.upload((books["orders"], "06 Contact Lists/orders.xlsx"))["batch"]["entry"]
    assert row["family"] == "orders"


def test_lists_the_review_says_do_not_belong_are_recognised_and_kept_out_not_errors(site, tiny, tmp_path):
    (row,) = site.upload((invoice_list(tmp_path), "08 Sales Order Exports/2026-03-09_-_317_Fibro_Order_List.xlsx"))["batch"]["entry"]
    assert (row["status"], row["kind"], row["family"]) == ("kept_out", "kept_out", None)
    assert row["outcome"].startswith("Recognised") and "kept out of the CRM on purpose" in row["outcome"]
    assert "invoice and packing lists do not belong in the CRM" in row["outcome"]


def test_a_file_nobody_knows_says_so_and_why(site, tiny, tmp_path):
    path = str(tmp_path / "mystery.xlsx")
    wb = openpyxl.Workbook()
    wb.active["A1"] = "Birthday plans"
    wb.save(path)
    d = site.upload((path, "mystery.xlsx"))
    (row,) = d["batch"]["entry"]
    assert row["status"] == "unknown" and "no sheet has the header columns" in row["outcome"]
    assert d["batch"]["ready_families"] == [] and d["batch"]["families"] == []


def test_a_broken_file_is_named_and_the_good_one_next_to_it_still_loads(site, tiny, books, tmp_path):
    broken = tmp_path / "broken.xlsx"
    broken.write_bytes(b"this is not a zip file")
    d = site.upload((str(broken), "broken.xlsx"), (books["todos"], "todos.xlsx"))
    by = {r["display_name"]: r for r in d["batch"]["entry"]}
    assert by["broken.xlsx"]["status"] == "unreadable" and "could not be opened" in by["broken.xlsx"]["outcome"].lower()
    assert by["todos.xlsx"]["status"] == "ready"
    assert d["batch"]["ready_families"] == ["todos"]


def test_identical_copies_are_named_and_read_once(site, tiny, books):
    d = site.upload((books["contacts"], "a/one.xlsx"), (books["contacts"], "b/two.xlsx"))
    first, second = d["batch"]["entry"]
    assert first["status"] == "ready" and second["status"] == "duplicate" and second["duplicate_of"] == "one.xlsx"
    (fam,) = d["batch"]["families"]
    assert fam["files"] == ["one.xlsx"]


def test_a_blank_template_is_recognised_and_says_it_holds_nothing(site, tiny, tmp_path):
    from tests.excel_book import write_book
    from webui.selftest import TODO_HEAD
    path = write_book(str(tmp_path / "to_do_template.xlsx"), {"Projekt-OPL": {"cells": {
        **{f"{chr(65 + i)}3": h for i, h in enumerate(TODO_HEAD)}}}})
    d = site.upload((path, "to_do_template.xlsx"))
    (row,) = d["batch"]["entry"]
    assert (row["family"], row["status"], row["records"]) == ("todos", "empty", 0)
    assert "holds nothing to import" in row["outcome"] and d["batch"]["ready_families"] == []
    assert site.client.post("/api/commit").status_code == 400  # nothing to write is a refusal, not an empty job


def test_non_excel_file_is_refused_and_keeps_nothing(site, tiny, books, tmp_path):
    site.upload((books["todos"], "todos.xlsx"))
    txt = tmp_path / "notes.txt"
    txt.write_text("hi")
    site.upload((str(txt), "notes.txt"), expect=400)
    assert site.batch()["batch"]["entry"] == []


# ---- the import: every family, in a safe order, one failing never blocks the others ----

def test_mixed_upload_commits_in_a_safe_order_and_a_second_commit_writes_nothing(site, tiny, books):
    d = upload_all(site, books)
    assert d["batch"]["ready_families"] == ["accounts", "contacts", "visits", "todos", "orders", "receivables"]
    job = site.commit()
    assert job["status"] == "done" and job["order"] == d["batch"]["ready_families"]
    fams = site.families(job)
    assert list(fams) == job["order"]  # reported in commit order
    assert all(f["status"] == "done" and f["written"] > 0 and f["failed"] == 0 for f in fams.values()), fams
    assert fams["visits"]["rows"]  # the visit report has its own lines

    before, writes = site.crm.snapshot(), len(site.crm.writes)
    again = site.commit()
    assert {k: f["status"] for k, f in site.families(again).items()} == {k: "nothing_new" for k in fams}
    assert all(f["written"] == 0 and f["present"] > 0 for f in site.families(again).values())
    assert len(site.crm.writes) == writes and site.crm.snapshot() == before


def test_each_family_in_the_mixed_upload_wrote_to_the_crm_object_the_review_names(site, tiny, books):
    upload_all(site, books)
    site.commit()
    crm = site.crm
    assert crm.companies and crm.contacts
    kinds = {a["kind"] for a in crm.activities.values()}
    assert {"meeting", "task", "note"} & kinds
    assert any(a["source_id"].startswith("order:") for a in crm.activities.values())  # order history is notes
    assert {f["label"] for f in crm.fields["company"]} >= {"AR open (CNY)", "AR overdue (CNY)", "AR as-of date"}


def test_one_family_failing_never_blocks_the_others(site, tiny, books, monkeypatch):
    from webui import families

    def boom(*a, **k):
        raise RuntimeError("the CRM said no")

    monkeypatch.setattr(families.ADAPTERS["contacts"], "commit", boom)
    upload_all(site, books)
    job = site.commit()
    assert job["status"] == "done"
    fams = site.families(job)
    assert fams["contacts"]["status"] == "failed" and "the CRM said no" in fams["contacts"]["errors"][0]
    assert all(f["status"] == "done" for k, f in fams.items() if k != "contacts")


def test_progress_reports_each_family_while_the_job_runs(site, tiny, books):
    upload_all(site, books, "todos", "orders")
    job = site.commit()
    assert job["order"] == ["todos", "orders"]
    assert [job["families"][k]["label"] for k in job["order"]] == ["To-dos and sales topics (07)", "Sales orders (08)"]


def test_dry_run_writes_nothing_and_says_what_each_family_would_write(site, tiny, books):
    upload_all(site, books)
    before, writes = site.crm.snapshot(), len(site.crm.writes)
    plan = site.dry_run()
    assert plan["wrote"] == 0
    rows = {f["key"]: {r["label"]: r["value"] for r in f["rows"]} for f in plan["families"]}
    assert list(rows) == ["accounts", "contacts", "visits", "todos", "orders", "receivables"]
    assert rows["accounts"]["New accounts"] == 2 and rows["todos"]["New tasks"] == 3
    assert all(f["error"] == "" for f in plan["families"])
    assert len(site.crm.writes) == writes and site.crm.snapshot() == before
    site.commit()
    after = {f["key"]: {r["label"]: r["value"] for r in f["rows"]} for f in site.dry_run()["families"]}
    assert after["todos"]["New tasks"] == 0 and after["todos"]["Tasks already in the CRM"] == 3


def test_dry_run_with_a_crm_that_cannot_be_reached_says_so(site, tiny, books, monkeypatch):
    def down():
        raise ConnectionError("connection refused")

    upload_all(site, books, "todos")
    monkeypatch.setattr(server, "_client", down)
    r = site.client.post("/api/dry_run")
    assert r.status_code == 400 and "could not log in" in r.get_json()["error"]


def test_a_second_import_while_one_runs_is_still_refused(site, tiny, books):
    upload_all(site, books, "todos")
    assert server.COMMIT_LOCK.acquire(blocking=False)
    try:
        assert site.client.post("/api/commit").status_code == 409
    finally:
        server.COMMIT_LOCK.release()


# ---- flagged groups: one click to keep or leave out ----

def test_the_it_list_is_flagged_left_out_by_default_and_one_click_brings_it_in(site, tiny, tmp_path):
    from tests.excel_book import write_book
    from webui.selftest import TODO_HEAD
    rows = {f"{chr(65 + i)}3": h for i, h in enumerate(TODO_HEAD)}
    rows.update({"D4": "Server", "E4": "migrate the server", "F4": "Chenley", "H4": "WIP", "D5": "Mail", "E5": "move the mailboxes",
                 "F5": "Ramon", "H5": "Open"})
    it = write_book(str(tmp_path / "to_do_list__JET-IT.xlsx"), {"Projekt-OPL": {"cells": rows}})
    d = site.upload((it, "to_do_list__JET-IT.xlsx"))
    (fam,) = d["batch"]["families"]
    (flag,) = fam["flags"]
    assert (flag["id"], flag["included"], flag["count"]) == ("it_list", False, 2)  # not known to be sales work: out until ticked in
    assert fam["ready"] == 0
    assert site.client.post("/api/flag", json={"family": "todos", "flag": "it_list", "include": True}).get_json()["batch"]["families"][0]["ready"] == 2

    off = site.client.post("/api/flag", json={"family": "todos", "flag": "it_list", "include": False}).get_json()
    assert off["batch"]["families"][0]["ready"] == 0 and off["batch"]["families"][0]["flags"][0]["included"] is False
    assert off["batch"]["ready_families"] == []
    on = site.client.post("/api/flag", json={"family": "todos", "flag": "it_list", "include": True}).get_json()
    assert on["batch"]["families"][0]["ready"] == 2
    job = site.commit()
    assert site.families(job)["todos"]["written"] == 2
    assert all("needs-check" in a["subject"] for a in site.crm.activities.values())  # tagged where the owner can filter


def test_flag_requests_for_something_unknown_are_refused(site, tiny, books):
    upload_all(site, books, "todos")
    for body in ({"family": "todos", "flag": "nope", "include": True}, {"family": "nope", "flag": "it_list", "include": True},
                 {"family": "todos", "flag": "it_list", "include": "yes"}):
        assert site.client.post("/api/flag", json=body).status_code == 400


# ---- company questions are asked for every family, with one shared answer file ----

def test_a_company_question_of_another_family_is_asked_and_answered(site, tiny, tmp_path):
    path = selftest.make_workbooks(str(tmp_path))["orders"][0]
    wb = openpyxl.load_workbook(path)
    for row in wb.active.iter_rows(min_row=3):
        row[2].value = "Selftest Customer Holdings Co., Ltd."  # close to, not the same as, the known customer
    wb.save(path)
    d = site.upload((path, "orders.xlsx"))
    (fam,) = d["batch"]["families"]
    (q,) = fam["questions"]
    assert q["name"] == "Selftest Customer Holdings Co., Ltd." and q["candidates"][0]["code"] == "KD99999"
    assert fam["waiting"] > 0 and fam["ready"] == 0

    r = site.client.post("/api/company_answer", json={"name": q["name"], "code": "KD99999"})
    assert r.status_code == 200, r.get_json()
    (fam,) = r.get_json()["batch"]["families"]
    assert fam["questions"] == [] and fam["waiting"] == 0 and fam["ready"] > 0
    assert site.client.post("/api/company_answer", json={"name": "Nobody Asked", "code": None}).status_code == 400


# ---- the Checks page ----

def test_checks_report_pass_warn_fail_rows_and_say_how_to_fix_a_failing_one(site, tmp_path, monkeypatch):
    cfg = tmp_path / "cfg.json"
    cfg.write_text(json.dumps({"base_url": "http://127.0.0.1:9", "email": "x@example.test", "password": "pw"}))
    monkeypatch.setattr(server, "CONFIG_PATH", str(cfg))
    monkeypatch.setattr(server, "_client", lambda: (_ for _ in ()).throw(ConnectionError("down")))
    d = site.client.get("/api/checks").get_json()
    env = {r["id"]: r for r in d["environment"]}
    assert env["crm"]["status"] == "fail" and "Start the CRM" in env["crm"]["fix"]
    assert env["login"]["status"] == "warn" and env["login"]["fix"]
    assert env["modules"]["status"] == "pass" and env["modules"]["fix"] == ""
    assert env["folders"]["status"] == "pass"
    assert {r["id"] for r in d["selftest"]} == {f"self-{k}" for k in FAMILY_FILES}
    assert all(r["status"] == "pass" for r in d["selftest"]), d["selftest"]
    assert all(set(r) == {"id", "title", "status", "detail", "fix"} for r in d["environment"] + d["selftest"])


def test_checks_with_a_working_crm_check_the_custom_fields(site, tmp_path, monkeypatch):
    class Up(SiteCRM):
        def get(self, path, **kw):
            return {"user": {"display_name": "Tester"}, "workspace_name": "Test"} if path == "/v1/me" else super().get(path, **kw)

    crm = Up()
    monkeypatch.setattr(server, "_client", lambda: crm)
    monkeypatch.setattr("requests.get", lambda url, timeout=0: type("R", (), {"status_code": 200})())
    env = {r["id"]: r for r in site.client.get("/api/checks").get_json()["environment"]}
    assert (env["crm"]["status"], env["login"]["status"]) == ("pass", "pass")
    assert env["fields"]["status"] == "warn" and "created automatically" in env["fields"]["detail"]
    selftest_ok = site.client.get("/api/checks").get_json()["selftest"]
    assert not site.crm.writes and crm.fields == {"company": [], "contact": [], "lead": []}  # checking never creates a field
    assert all(r["status"] == "pass" for r in selftest_ok)


# ---- what the upload table says will happen ----

def _edited_copy(src, dst, **cells):
    wb = openpyxl.load_workbook(src)
    for coord, value in cells.items():
        wb.active[coord] = value
    wb.save(dst)
    return str(dst)


def test_a_macro_workbook_is_accepted(site, tiny, books, tmp_path):
    macro = tmp_path / "Contacts macro.xlsm"
    macro.write_bytes(open(books["contacts"], "rb").read())
    (row,) = site.upload((str(macro), macro.name))["batch"]["entry"]
    assert (row["family"], row["status"]) == ("contacts", "ready")


def test_a_visit_file_says_how_many_go_in_and_how_many_wait(site, tiny, books):
    (row,) = site.upload((books["visits"], os.path.basename(books["visits"])))["batch"]["entry"]
    assert row["status"] == "ready" and "2 visit(s) found: 2 go in" in row["outcome"] and "wait" not in row["outcome"]


def test_a_visit_waiting_for_a_fix_is_not_counted_as_going_in(site, tiny, books, tmp_path):
    bad = _edited_copy(books["visits"], tmp_path / "Template for Weekly Sales Reporting 2026_CW10_Other Rep.xlsx", F9=None, P9="Call")
    (row,) = site.upload((bad, os.path.basename(bad)))["batch"]["entry"]
    assert "1 go in" in row["outcome"] and "1 wait" in row["outcome"]


def test_an_older_copy_of_a_visit_report_names_the_file_that_holds_its_visits(site, tiny, books, tmp_path):
    older = _edited_copy(books["visits"], tmp_path / "Template for Weekly Sales Reporting 2026_CW10_Older Copy.xlsx", N9="x")
    d = site.upload((books["visits"], os.path.basename(books["visits"])), (older, os.path.basename(older)))
    rows = {r["display_name"]: r for r in d["batch"]["entry"]}
    statuses = sorted(r["status"] for r in rows.values())
    assert statuses == ["ready", "superseded"]
    gone = next(r for r in rows.values() if r["status"] == "superseded")
    kept = next(r for r in rows.values() if r["status"] == "ready")
    assert kept["display_name"] in gone["outcome"] and "blank template" not in gone["outcome"]


def test_a_contact_list_with_only_blank_rows_adds_nothing(site, tiny, books, tmp_path):
    blank = tmp_path / "Contacts blank.xlsx"
    wb = openpyxl.load_workbook(books["contacts"])
    ws = wb.active
    for r in range(2, ws.max_row + 1):
        for c in range(2, ws.max_column + 1):
            ws.cell(r, c).value = None
    wb.save(blank)
    (row,) = site.upload((str(blank), blank.name))["batch"]["entry"]
    assert row["status"] != "ready" or row["records"] == 0
    assert "2 rows" not in row["outcome"]


def test_an_older_order_export_says_its_orders_are_in_the_newer_one(site, tiny, books, tmp_path):
    newer = _edited_copy(books["orders"], tmp_path / "Sales Order List selftest newer.xlsx", H3=5)  # the same two orders, one changed
    d = site.upload((books["orders"], os.path.basename(books["orders"])), (newer, os.path.basename(newer)))
    outcomes = [r["outcome"] for r in d["batch"]["entry"]]
    assert any("also in a newer export" in o for o in outcomes), outcomes


def test_a_second_ar_snapshot_says_which_one_is_used(site, tiny, books, tmp_path):
    other = _edited_copy(books["receivables"], tmp_path / "2026_Financial_Status_20260301.xlsx", Z1="other")
    d = site.upload((books["receivables"], os.path.basename(books["receivables"])), (other, os.path.basename(other)))
    rows = {r["display_name"]: r for r in d["batch"]["entry"]}
    gone = [r for r in rows.values() if r["status"] == "superseded"]
    assert len(gone) == 1 and sorted(r["status"] for r in rows.values()) == ["ready", "superseded"]
    used = next(n for n, r in rows.items() if r["status"] == "ready")
    assert used.removesuffix(".xlsx") in gone[0]["outcome"], gone[0]


def test_a_visit_only_the_excel_can_fix_is_listed_with_its_place_and_marked(site, tiny, books, tmp_path):
    bad = _edited_copy(books["visits"], tmp_path / "Template for Weekly Sales Reporting 2026_CW10_Other Rep.xlsx", F9=None)
    d = site.upload((bad, os.path.basename(bad)))
    assert d["needs_excel_fix"] == 1
    (row,) = [r for r in d["unsure_rows"] if r["needs_excel_fix"]]
    assert (row["sheet"], row["row"]) == ("CW10", 9) and "no customer" in row["why"]
