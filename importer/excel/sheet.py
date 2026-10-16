"""Loads a workbook's sheets into a bounded grid that the layout readers index by
(row, column), 1-based. A merged cell reports its top-left value, and the other
cells of the merge are flagged so a block's continuation rows are not read as
new data. Formulas are read lazily, only when a cached value is missing."""
import io
import re
import zipfile

import openpyxl

MAX_COL = 40
MAX_ROW = 20000


_NAMELESS_STYLE = re.compile(r"<cellStyle(?![^>]*\bname=)[^>]*/>")


def open_workbook(path, **kwargs):
    """openpyxl.load_workbook, except for workbooks Excel opens but openpyxl refuses because a
    <cellStyle> in styles.xml has no name (seen in ERP exports): those are read from a copy in
    memory without the nameless styles. The file on disk is never touched. Any other failure
    is raised as before."""
    try:
        return openpyxl.load_workbook(path, **kwargs)
    except TypeError:
        repaired = _without_nameless_styles(path)
        if repaired is None:
            raise
        return openpyxl.load_workbook(repaired, **kwargs)


def _without_nameless_styles(path):
    try:
        with zipfile.ZipFile(path) as zin:
            styles = zin.read("xl/styles.xml").decode("utf8")
            fixed = _NAMELESS_STYLE.sub("", styles)
            if fixed == styles:
                return None
            buf = io.BytesIO()
            with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
                for item in zin.infolist():
                    data = zin.read(item.filename)
                    zout.writestr(item, fixed.encode("utf8") if item.filename == "xl/styles.xml" else data)
    except (KeyError, zipfile.BadZipFile):
        return None
    buf.seek(0)
    return buf


class _Formulas:
    """Opens the workbook without cached values, once, the first time a formula is asked for."""

    def __init__(self, path):
        self.path = path
        self.book = None

    def cell(self, name, r, c):
        if self.book is None:
            self.book = open_workbook(self.path)
        value = self.book[name].cell(row=r, column=c).value
        return value if isinstance(value, str) and value.startswith("=") else None


class Sheet:
    def __init__(self, name, max_row, max_col, grid, top, formulas=None, ws=None, cut=""):
        self.cut = cut  # why part of the sheet was left out (rows or columns beyond the caps); "" when all was read
        self.name = name
        self.max_row = max_row
        self.max_col = max_col
        self.grid = grid  # (row, col) -> value
        self.top = top  # merged continuation (row, col) -> top-left (row, col)
        self.ws = ws  # openpyxl worksheet (xlsx only), for delegating to visits.reader
        self._formulas = formulas

    def value(self, r, c):
        r, c = self.top.get((r, c), (r, c))
        return self.grid.get((r, c))

    def is_merged_continuation(self, r, c):
        return (r, c) in self.top

    def formula(self, r, c):
        if self._formulas is None:
            return None
        return self._formulas.cell(self.name, r, c)


def _cut_note(rows, cols):
    parts = []
    if rows:
        parts.append(f"rows beyond {MAX_ROW}")
    if cols:
        parts.append(f"columns beyond {MAX_COL}")
    return ("only the first %d rows and %d columns were read: the sheet has data in %s" %
            (MAX_ROW, MAX_COL, " and ".join(parts))) if parts else ""


def load(path):
    """All sheets of an .xlsx or .xls workbook, in workbook order."""
    if path.lower().endswith(".xls"):
        return _load_xls(path)
    wb = open_workbook(path, data_only=True)
    formulas = _Formulas(path)
    out = []
    for ws in wb.worksheets:
        max_row = min(ws.max_row, MAX_ROW)
        max_col = min(ws.max_column, MAX_COL)
        grid = {}
        for row in ws.iter_rows(min_row=1, max_row=max_row, max_col=max_col):
            for cell in row:
                if cell.value is not None:
                    grid[(cell.row, cell.column)] = cell.value
        top = {}
        for rng in ws.merged_cells.ranges:
            for r in range(rng.min_row, min(rng.max_row, max_row) + 1):
                for c in range(rng.min_col, min(rng.max_col, max_col) + 1):
                    if (r, c) != (rng.min_row, rng.min_col):
                        top[(r, c)] = (rng.min_row, rng.min_col)
        beyond = [cell for (r, c), cell in ws._cells.items()
                  if (r > MAX_ROW or c > MAX_COL) and cell.value is not None]
        cut = _cut_note(any(cell.row > MAX_ROW for cell in beyond), any(cell.column > MAX_COL for cell in beyond))
        out.append(Sheet(ws.title, max_row, max_col, grid, top, formulas, ws, cut))
    return out


def _load_xls(path):
    import xlrd

    try:
        book = xlrd.open_workbook(path, formatting_info=True)
    except NotImplementedError:
        book = xlrd.open_workbook(path)
    out = []
    for sh in book.sheets():
        max_row = min(sh.nrows, MAX_ROW)
        max_col = min(sh.ncols, MAX_COL)
        grid = {}
        for r in range(max_row):
            for c in range(max_col):
                ctype = sh.cell_type(r, c)
                if ctype in (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK):
                    continue
                value = sh.cell_value(r, c)
                if ctype == xlrd.XL_CELL_DATE:
                    value = xlrd.xldate.xldate_as_datetime(value, book.datemode)
                grid[(r + 1, c + 1)] = value
        top = {}
        for rlo, rhi, clo, chi in getattr(sh, "merged_cells", []):
            for r in range(rlo, min(rhi, max_row)):
                for c in range(clo, min(chi, max_col)):
                    if (r, c) != (rlo, clo):
                        top[(r + 1, c + 1)] = (rlo + 1, clo + 1)
        empty = (xlrd.XL_CELL_EMPTY, xlrd.XL_CELL_BLANK)
        rows_cut = any(t not in empty for r in range(MAX_ROW, sh.nrows) for t in sh.row_types(r))
        cols_cut = any(t not in empty for r in range(min(sh.nrows, MAX_ROW))
                       for t in sh.row_types(r, start_colx=MAX_COL))
        out.append(Sheet(sh.name, max_row, max_col, grid, top, cut=_cut_note(rows_cut, cols_cut)))
    return out
