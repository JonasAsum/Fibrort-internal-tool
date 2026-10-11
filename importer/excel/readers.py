"""One reader per layout. read_sheet(sheet, path) detects a sheet's layout and
returns (layout, visits, week, review); a sheet with no known layout comes back
as (None, [], None, []). `mock` delegates to visits.reader.parse_sheet unchanged.

Block rule (weekly2026, lucas): a visit starts on a row whose customer cell (or
date, when the customer is empty) has a value and is not a merged continuation;
it runs to the row before the next start. Each next-step cell in the block is a
green note; purpose and communication continuation cells are appended.
"""
import os
import re
from dataclasses import dataclass, field
from datetime import date, datetime, timedelta

from openpyxl.utils import get_column_letter

from visits.reader import WEEKDAY_OFFSETS, WORKBOOK_YEAR, DayPlan, Note, Visit, Week, cw_monday, parse_sheet

from .dates import parse_date, parse_due
from .headers import find_header, normalize
from .layouts import detect, label_cell
from .sheet import Sheet
from .split import SUFFIX_END, is_other_name, is_placeholder, latin_alias, phone_numbers, split_customer, visit_kind

REPS = ["Jason Zhao", "Devin Guo", "Mark Wang", "Sean Wang", "Ethan Lv", "Jason Yi", "Andrew Zhang", "Lucas Guo"]
POTENTIAL = {"A", "B", "C"}
SHEET_CW = re.compile(r"^\s*(?:cw|week|w)?\s*(\d{1,2})\s*$", re.I)
FILE_CW = re.compile(r"CW\s*(\d{1,2})", re.I)
YEAR_IN_NAME = re.compile(r"(20\d{2})")
SHEET_MD = re.compile(r"^\s*(\d{1,2})\.(\d{1,2})")
PLACEHOLDER = re.compile(r"^\d+[.)]?$")  # "2." lines left in an empty next-step cell
ITEM_FIELDS = (("id", "id"), ("typ", "typ"), ("date", "date"), ("from", "from"), ("topic", "general topic"),
               ("details", "details"), ("wv", "wv"), ("wer", "wer"), ("status", "status"))


@dataclass
class ReviewItem:
    family: str
    file: str
    sheet: str
    row: int
    field: str
    value: str
    reason: str
    candidates: list = field(default_factory=list)


@dataclass
class Ctx:
    path: str
    family: str
    year: int
    folder_owner: str
    year_known: bool = True  # False: neither the file name nor sheet CW + file date told the year

    @property
    def file(self) -> str:
        return os.path.basename(self.path)

    def review(self, sheet, row, fld, value, reason, candidates=()):
        return ReviewItem(self.family, self.file, sheet, row, fld,
                          "" if value is None else str(value), reason, list(candidates))


def _text(value) -> str:
    """The cell as text; placeholders ("/", "none", "无" ...) come back empty."""
    if value is None:
        return ""
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    text = str(value).strip()
    return "" if is_placeholder(text) else text


def _raw(sheet: Sheet, r, c):
    """The cell's value, or its formula text when no cached value was saved."""
    value = sheet.value(r, c)
    return value if value is not None else sheet.formula(r, c)


def normalize_owner(raw) -> str:
    """One of REPS for a name or folder name, "" when unknown or ambiguous."""
    t = _text(raw).lower()
    if not t:
        return ""
    for rep in REPS:
        if t == rep.lower():
            return rep
    first = [rep for rep in REPS if t.split()[0] == rep.split()[0].lower()]
    if len(first) == 1 and len(t.split()) == 1:
        return first[0]
    hits = [rep for rep in REPS if rep.lower() in t]
    return hits[0] if len(hits) == 1 else ""


def owner_from_filename(name: str) -> str:
    """A rep named in a file name: the full name, else a first name only one rep has."""
    stem = re.sub(r"[\W_\d]+", " ", os.path.splitext(os.path.basename(name))[0]).lower()
    words = stem.split()
    full = [rep for rep in REPS if rep.lower() in stem]
    if len(full) == 1:
        return full[0]
    firsts = {w for w in words}
    hits = [rep for rep in REPS if rep.split()[0].lower() in firsts]
    first_names = [rep.split()[0] for rep in REPS]
    return hits[0] if len(hits) == 1 and first_names.count(hits[0].split()[0]) == 1 else ""


def _sheet_owner(sheet: Sheet, folder_owner: str) -> str:
    at = label_cell(sheet, "owner", max_row=8)
    if at:
        r, c = at
        for cc in range(c + 1, c + 6):
            owner = normalize_owner(sheet.value(r, cc))
            if owner:
                return owner
    return folder_owner


def _sheet_cw(name: str) -> int | None:
    m = SHEET_CW.match(name)
    return int(m.group(1)) if m else None


def _file_cw(path: str) -> int | None:
    m = FILE_CW.search(os.path.basename(path))
    return int(m.group(1)) if m else None


def _monday(year: int, cw: int | None) -> date | None:
    if not cw:
        return None
    try:
        return cw_monday(year, cw)
    except ValueError:
        return None


def _first_col_right(sheet: Sheet, r: int, c: int) -> int | None:
    for cc in range(c + 1, sheet.max_col + 1):
        if _text(sheet.value(r, cc)):
            return cc
    return None


def _order_in_day(visits: list[Visit]) -> None:
    seen: dict[str, int] = {}
    for v in visits:
        key = v.date_iso or v.date_raw
        v.order_in_day = seen.get(key, 0)
        seen[key] = v.order_in_day + 1


def _is_start(sheet: Sheet, r: int, cols: dict, year: int) -> bool:
    c = cols.get("customer")
    if c is not None and _text(sheet.value(r, c)) and not sheet.is_merged_continuation(r, c):
        return True
    # A customer cell merged over several dated rows (one customer, several visits): each
    # row with its own date is a visit of that customer.
    # Without a customer the date must read as a date, so footnotes in the date column do not start a visit.
    d = cols.get("date")
    if d is None or sheet.is_merged_continuation(r, d):
        return False
    iso, flag = parse_date(_raw(sheet, r, d), None, year)
    return iso is not None or flag in ("ambiguous", "range")


def _read_blocks(sheet, ctx, header_row, cols, owner, cw_hint):
    starts = [r for r in range(header_row + 1, sheet.max_row + 1) if _is_start(sheet, r, cols, ctx.year)]
    visits, review = [], []
    for i, start in enumerate(starts):
        end = starts[i + 1] - 1 if i + 1 < len(starts) else sheet.max_row
        visit, items = _block_visit(sheet, ctx, start, end, cols, owner, cw_hint)
        visits.append(visit)
        review += items
    _order_in_day(visits)
    return visits, review


def _year_items(ctx, sheet_name, row, date_raw, date_iso):
    """The year the file name (or the sheet's CW and the file's date) gives must be the date's ISO year."""
    if not ctx.year_known:
        return [ctx.review(sheet_name, row, "date", date_raw, "no_year")]
    if date.fromisoformat(date_iso).isocalendar()[0] != ctx.year:
        return [ctx.review(sheet_name, row, "date", date_raw, "week_mismatch")]
    return []


def _block_visit(sheet, ctx, start, end, cols, owner, cw_hint):
    review = []

    def cell(fld, r=start):
        c = cols.get(fld)
        return "" if c is None else _text(sheet.value(r, c))

    date_c = cols.get("date")
    date_cell = _raw(sheet, start, date_c) if date_c else None
    date_raw = _text(date_cell)
    ref = _monday(ctx.year, cw_hint)
    date_iso, date_flag = parse_date(date_cell, ref, ctx.year, lookup=sheet.value)
    visit_day = date.fromisoformat(date_iso) if date_iso else (ref or date(ctx.year, 1, 1))

    # "上海某某有限公司(Some Name)" on the first line: the latin part is the English name.
    first, sep, rest = str(cell("customer") or "").partition("\n")
    first, name_en = latin_alias(first)
    # A full legal name on the first line with its translation below: the second line is
    # the other-language name, not an address and not part of the name.
    if rest and SUFFIX_END.search(first.strip()) and is_other_name(rest):
        name_en, rest, sep = name_en or " ".join(rest.split()), "", ""
    name, address, sure = split_customer(first + sep + rest)
    purpose = [cell("purpose")]
    communication = [cell("communication")]
    for r in range(start + 1, end + 1):
        for fld, parts in (("purpose", purpose), ("communication", communication)):
            c = cols.get(fld)
            if c is not None and not sheet.is_merged_continuation(r, c) and _text(sheet.value(r, c)):
                parts.append(_text(sheet.value(r, c)))

    dues, first_due = {}, None
    due_c = cols.get("due")
    for r in range(start, end + 1):
        raw = sheet.value(r, due_c) if due_c else None
        if not _text(raw):
            continue
        iso, phrase, dflag = parse_due(raw, visit_day)
        dues[r] = iso or phrase
        if dflag == "decimal":
            review.append(ctx.review(sheet.name, r, "due", raw, "due_decimal"))
        first_due = first_due or dues[r]

    notes = []
    ns_c = cols.get("next_step")
    for r in range(start, end + 1):
        if ns_c is None or sheet.is_merged_continuation(r, ns_c):
            continue
        text = _text(sheet.value(r, ns_c))
        if text and not PLACEHOLDER.match(text):
            notes.append(Note(text=text, highlight="green", deadline=dues.get(r) or first_due or "",
                              excel_row=r))

    if date_iso is None:
        review.append(ctx.review(sheet.name, start, "date", date_raw, "no_date"))
    elif date_flag in ("ambiguous", "range"):
        review.append(ctx.review(sheet.name, start, "date", date_raw, f"{date_flag}_date"))
    if date_iso and cw_hint and abs(date.fromisoformat(date_iso).isocalendar()[1] - cw_hint) > 1:
        review.append(ctx.review(sheet.name, start, "date", date_raw, "week_mismatch"))
    elif date_iso:
        review += _year_items(ctx, sheet.name, start, date_raw, date_iso)
    if not name:
        review.append(ctx.review(sheet.name, start, "customer", cell("customer"), "no_customer"))
    elif not sure:
        review.append(ctx.review(sheet.name, start, "customer", cell("customer"), "customer_split"))

    row_owner = owner or normalize_owner(cell("sales"))
    if not row_owner:
        review.append(ctx.review(sheet.name, start, "owner", "", "no_owner"))

    extra = {
        "wechat": cell("wechat"), "industry": cell("industry"), "competitor": cell("competitor"),
        "communication": "\n".join(p for p in communication if p), "status": cell("status"),
        "address": address, "sales": cell("sales"), "plan": cell("plan"), "cw_cell": cell("cw"),
        "source_cells": {fld: f"{get_column_letter(c)}{start}" for fld, c in cols.items()},
    }
    if "follow_up" in cols:
        extra["follow_up"] = cell("follow_up")
    extra["name_en"] = name_en
    extra["kind"] = visit_kind(cell("customer"))
    phones = phone_numbers(sheet.value(start, cols["phone"])) if "phone" in cols else []
    if len(phones) > 1:
        extra["phone_more"] = "; ".join(phones[1:])
    potential = cell("potential").upper()
    visit = Visit(
        workbook=ctx.file, sheet=sheet.name, owner=row_owner,
        year=int(date_iso[:4]) if date_iso else ctx.year,
        cw=date.fromisoformat(date_iso).isocalendar()[1] if date_iso else (cw_hint or 0),
        excel_row_start=start, excel_row_end=end, date_raw=date_raw, date_iso=date_iso,
        city=cell("city"), customer_raw=name,
        potential=potential if potential in POTENTIAL else "",
        customer_type_raw=cell("customer_type"), contact_raw=cell("contact"),
        phone_raw=phones[0] if phones else "",
        office_phone=cell("office"), email=cell("email"),
        purpose="\n".join(p for p in purpose if p), notes=notes, extra=extra,
    )
    return visit, review


def read_weekly(sheet, ctx, owner, found):
    row, cols = found
    visits, review = _read_blocks(sheet, ctx, row, cols, owner, _sheet_cw(sheet.name))
    return visits, None, review


def read_lucas(sheet, ctx, owner, found):
    row, cols = found
    cw = _file_cw(ctx.path)  # the Timing cell is stale in some files; the file name is not
    visits, review = _read_blocks(sheet, ctx, row, cols, owner, cw)
    first = next((v for v in visits if v.date_iso), None)
    week = None
    if first:
        day = date.fromisoformat(first.date_iso)
        monday = day - timedelta(days=day.weekday())
        iso_year, iso_week, _ = day.isocalendar()
        days = []
        for r in range(1, row):
            label = _text(sheet.value(r, 3))
            offset = WEEKDAY_OFFSETS.get(label[:3].lower()) if label.isalpha() else None
            if offset is not None:
                days.append(DayPlan(weekday=label, date_iso=(monday + timedelta(days=offset)).isoformat(),
                                    plan=_text(sheet.value(r, 4))))
        week = Week(workbook=ctx.file, sheet=sheet.name, owner=owner, year=iso_year, cw=iso_week,
                    monday_iso=monday.isoformat(), days=days, summary=[])
    return visits, week, review


ITEM_HEADER_MAX_ROW = 40
EVENT = re.compile(r"\b(?:exhibition|expo|fair|show|trip)\b|展会|展览", re.I)
VISIT_TOPIC = re.compile(r"(?:\bvisit(?:ed|ing)?\b(?:\s+to\b)?|\bmeeting\s+with\b|拜访)\s*(.+)", re.I)
TOPIC_CUT = re.compile(r"\s+and\s|\s*\(|\s+to\s|\s+for\s|,(?!\s*(?:Co|Ltd))|[，；;]", re.I)
ROLE_WORDS = {"manager", "director", "chairman", "engineer", "supervisor", "technical", "technology",
              "technician", "production", "sales", "general", "chief", "senior", "deputy", "assistant",
              "president", "head", "vice", "workshop", "quality", "purchasing", "procurement", "project",
              "owner", "lead", "ceo", "cto", "coo"}
COMPANY_TAIL = re.compile(r"\b(?:Co|Ltd|Technology|Machine Tools|GmbH|Inc|Corp|Group|Tools)\b", re.I)


def _topic_company(topic: str) -> str:
    """The company after "visit to", "visited", "meeting with" or "拜访" in a General topic."""
    m = VISIT_TOPIC.search(topic or "")
    if not m:
        return ""
    return " ".join(TOPIC_CUT.split(m.group(1), maxsplit=1)[0].split()).strip(" .,;，、")


def _role_company(role: str) -> str:
    """"director of Destaco" -> "Destaco"; "manager Nupa Technology Co., Ltd" -> "Nupa Technology Co., Ltd"."""
    text = " ".join(role.split())
    parts = re.split(r"\s+of\s+", text, flags=re.I)
    if len(parts) > 1:
        return parts[-1].strip(" .,;")
    words = text.split()
    while len(words) > 1 and words[0].lower().strip(".,") in ROLE_WORDS:
        words.pop(0)
    rest = " ".join(words).strip(" ,;")
    return rest if COMPANY_TAIL.search(rest) else ""


def _co_key(text: str) -> str:
    return re.sub(r"[\W_]+", "", text.casefold())


def _same_company(a: str, b: str) -> bool:
    ka, kb = _co_key(a), _co_key(b)
    return len(ka) >= 3 and len(kb) >= 3 and (ka in kb or kb in ka)


def _person(raw) -> str:
    """A rep from a "from"/"Wer" cell, which may read "Devin.Guo"."""
    return normalize_owner(re.sub(r"[._\-]+", " ", _text(raw)))


def _item_header(sheet: Sheet):
    """(header row, {field: col}, last header col) for the item table."""
    for r in range(1, min(ITEM_HEADER_MAX_ROW, sheet.max_row) + 1):
        labels = {}
        for c in range(1, sheet.max_col + 1):
            t = normalize(sheet.value(r, c))
            if t:
                labels.setdefault(t, c)
        if "id" in labels and "general topic" in labels:
            cols = {}
            for fld, prefix in ITEM_FIELDS:
                for t, c in labels.items():
                    if t == prefix or t.startswith(prefix):
                        cols.setdefault(fld, c)
                        break
            return r, cols, max(labels.values())
    return None


def _sheet_md(name: str, year: int) -> str | None:
    m = SHEET_MD.match(name)
    if not m:
        return None
    try:
        return date(year, int(m.group(1)), int(m.group(2))).isoformat()
    except ValueError:
        return None


def _detail_visit(sheet, ctx, rows, cols, last, sheet_customer, event, participants, owner):
    """One item: its head row plus the continuation rows below it."""
    review = []
    head = rows[0]

    def field(fld, rs=rows):
        col = cols.get(fld)
        if col is None:
            return ""
        return "\n".join(t for t in (_text(sheet.value(r, col)) for r in rs) if t)

    topic = field("topic", [head])
    details = field("details")
    wv_raw = sheet.value(head, cols["wv"]) if "wv" in cols else None
    date_cell = _raw(sheet, head, cols["date"]) if "date" in cols else None
    date_iso, date_flag = parse_date(date_cell, None, ctx.year, lookup=sheet.value)
    if date_flag in ("ambiguous", "range"):
        review.append(ctx.review(sheet.name, head, "date", _text(date_cell), f"{date_flag}_date"))
    if date_iso is None:
        date_iso = _sheet_md(sheet.name, ctx.year)
    visit_day = date.fromisoformat(date_iso) if date_iso else date(ctx.year, 1, 1)

    customer = _topic_company(topic) or sheet_customer
    contact = next((p["name"] for p in participants
                    if p["company"] and customer and _same_company(p["company"], customer)), "")
    owner = _person(field("from", [head])) or _person(field("wer", [head])) or owner
    wer = field("wer", [head])

    notes = []
    if details:
        notes.append(Note(text=details, highlight=None, deadline="", excel_row=head))
    if _text(wv_raw):
        iso, phrase, _ = parse_due(wv_raw, visit_day)
        notes.append(Note(text=f"{topic or details} (Wer: {wer})", highlight="green",
                          deadline=iso or phrase, excel_row=head))

    extra_notes = [_text(sheet.value(r, cc)) for r in rows for cc in range(last + 1, sheet.max_col + 1)]
    if not customer:
        review.append(ctx.review(sheet.name, head, "customer", "", "no_customer"))
    if date_iso is None:
        review.append(ctx.review(sheet.name, head, "date", _text(date_cell), "no_date"))
    if not owner:
        review.append(ctx.review(sheet.name, head, "owner", "", "no_owner"))
    if date_iso:
        review += _year_items(ctx, sheet.name, head, _text(date_cell), date_iso)
    if participants and not contact:
        review.append(ctx.review(sheet.name, head, "participants",
                                 "; ".join(p["name"] for p in participants), "participants"))

    visit = Visit(
        workbook=ctx.file, sheet=sheet.name, owner=owner,
        year=int(date_iso[:4]) if date_iso else ctx.year,
        cw=date.fromisoformat(date_iso).isocalendar()[1] if date_iso else 0,
        excel_row_start=head, excel_row_end=rows[-1], date_raw=_text(date_cell), date_iso=date_iso,
        city="", customer_raw=customer, potential="", customer_type_raw="", contact_raw=contact,
        phone_raw="", office_phone="", email="", purpose=topic, notes=notes,
        extra={
            "kind": visit_kind(customer), "event": event,
            "participants": [dict(p) for p in participants], "details": details,
            "status_text": field("status"), "typ": field("typ", [head]),
            "from_person": field("from", [head]), "wv": _text(wv_raw), "wer": wer,
            "extra_notes": " | ".join(t for t in extra_notes if t),
            "source_cells": {fld: f"{get_column_letter(col)}{head}" for fld, col in cols.items()},
        },
    )
    return visit, review


def read_detail(sheet, ctx, owner):
    """A trip report: one Visit per item row. The "Customer name:" cell can be an event;
    the company comes from the item's topic, else from that cell."""
    review = []
    label = label_cell(sheet, "customer name")
    if label is None:
        return [], None, review
    lr, lc = label
    c = _first_col_right(sheet, lr, lc)
    sheet_customer = _text(sheet.value(lr, c)) if c else ""
    event = sheet_customer if EVENT.search(sheet_customer) else ""

    participants = []
    p = label_cell(sheet, "participants")
    if p:
        pr, pc = p
        for r in range(pr, min(pr + 20, sheet.max_row) + 1):
            name_c = _first_col_right(sheet, r, pc)
            if name_c is None:
                break
            role_c = _first_col_right(sheet, r, name_c)
            role = _text(sheet.value(r, role_c)) if role_c else ""
            participants.append({"name": _text(sheet.value(r, name_c)), "role": role,
                                 "company": _role_company(role)})

    head = _item_header(sheet)
    if head is None:
        return [], None, review
    hr, cols, last = head
    items = []
    for r in range(hr + 1, sheet.max_row + 1):
        starts = bool(_text(sheet.value(r, cols["topic"]))) or (
            "date" in cols and bool(_text(_raw(sheet, r, cols["date"]))))
        if starts:
            items.append([r])
        elif items and any(_text(sheet.value(r, cc)) for cc in range(1, sheet.max_col + 1)):
            items[-1].append(r)

    visits = []
    for rows in items:
        visit, found = _detail_visit(sheet, ctx, rows, cols, last, sheet_customer, event,
                                     participants, owner)
        visits.append(visit)
        review += found
    _order_in_day(visits)
    return visits, None, review


def read_mock(sheet, ctx):
    visits, week = parse_sheet(sheet.ws, ctx.file, ctx.year)
    return visits, week, []


def _year_for(path: str, cw: int | None = None) -> tuple[int, bool]:
    """(year, known). The year in the file name; else the year whose CW `cw` falls nearest the
    file's modification date; else unknown (the file date's year, or today's, as a placeholder)."""
    name = os.path.basename(path)
    if m := WORKBOOK_YEAR.search(name):
        return int(m.group(1)), True
    if m := YEAR_IN_NAME.search(name):
        return int(m.group(1)), True
    try:
        modified = datetime.fromtimestamp(os.path.getmtime(path)).date()
    except OSError:
        return date.today().year, False
    if cw:
        near = [(abs((_monday(y, cw) - modified).days), y) for y in range(modified.year - 1, modified.year + 2)
                if _monday(y, cw)]
        if near:
            return min(near)[1], True
    return modified.year, False


def read_sheet(sheet: Sheet, path: str):
    """(layout, visits, week, review) for one sheet of the workbook at path."""
    layout = detect(sheet, path)
    if layout is None:
        return None, [], None, []
    folder_owner = (normalize_owner(os.path.basename(os.path.dirname(os.path.abspath(path))))
                    or owner_from_filename(path))
    cw = _file_cw(path) if layout == "lucas" else _sheet_cw(sheet.name) or _file_cw(path)
    year, known = _year_for(path, cw)
    ctx = Ctx(path=path, family=layout, year=year, folder_owner=folder_owner, year_known=known)
    if layout == "mock":
        visits, week, review = read_mock(sheet, ctx)
    elif layout == "detail":
        visits, week, review = read_detail(sheet, ctx, folder_owner)
    else:
        owner = _sheet_owner(sheet, folder_owner)
        found = find_header(sheet)
        if layout == "lucas":
            visits, week, review = read_lucas(sheet, ctx, owner, found)
        else:
            visits, week, review = read_weekly(sheet, ctx, owner, found)
    return layout, visits, week, review
