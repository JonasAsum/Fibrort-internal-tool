"""Built-in self-test: writes one tiny invented workbook per supported family into a temporary
folder, reads it through the real reader and compares the counts with what was put in. Writes
nothing to the CRM, to the upload folder or to the company answers. Every name here is made up.
"""
import datetime as dt
import os
import tempfile

import openpyxl

from excel import accounts, contacts, orders, receivables, todos
from excel.companies import Master
from excel.run import read_paths as read_visits
from webui.entry import classify

CUSTOMER_CN, CUSTOMER_EN = "自检测试有限公司", "Selftest Customer Co., Ltd."
MASTER = [{"source": "master", "code": "KD99999", "cn": CUSTOMER_CN, "en": CUSTOMER_EN, "short": ""}]
DAY = dt.datetime(2026, 3, 5)

WEEKLY = {"B8": "Sales", "C8": "CW", "D8": "Date", "E8": "Location/City", "F8": "Customer \n(full address if new)",
          "G8": "Potential\n(A, B, C)", "H8": "Customer \n(new, exist, dealer)", "I8": "Contact Name",
          "J8": "Phone No.", "K8": "WeChat", "L8": "Email", "M8": "Industries", "N8": "Purpose_Topics Visit",
          "O8": "Competiton Products", "P8": "Next Step / Action", "Q8": "Due Date", "R8": "Status"}
ORDER_HEAD = ["Document No.", "Document Date", "Customer", "Salesperson", "Order Status", "Product Code",
              "Product Name", "Sales Quantity", "Sales Unit", "Amount Excl. Tax", "Tax Amount", "Amount Incl. Tax"]
AR_HEAD = ["Customer name(EN)", "Customer name", "Sales invoice no.", "Invoice Date", "Due Date", "Payment Term in Days",
           "Invoice Value Local Currency（RMB）", "Due days", "Days since Invoice", "Remark", "币种", "Original amt", "Currency"]
TODO_HEAD = ["No.", "Categorie", "Date", "Topics", "To Do / Next Steps", "Responsible", "Due Date", "Status", "Prio",
             "Remarks", "Documents or Websites (Hyperlink)"]


def _book(path, sheets):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    for title, rows in sheets.items():
        ws = wb.create_sheet(title)
        if isinstance(rows, dict):
            for coord, value in rows.items():
                ws[coord] = value
        else:
            for r in rows:
                ws.append(r)
    wb.save(path)
    return path


def make_workbooks(folder) -> dict:
    """{family: (path, expected counts)} - the workbooks are invented and tiny."""
    p = lambda name: os.path.join(folder, name)  # noqa: E731
    weekly = dict(WEEKLY, E4="Owner：", F4="Jason Zhao", B9="Jason Zhao", D9=DAY, F9=CUSTOMER_EN, P9="Call",
                  B10="Jason Zhao", D10=DAY + dt.timedelta(days=1), F10=CUSTOMER_EN, P10="Visit")
    return {
        "visits": (_book(p("Template for Weekly Sales Reporting 2026_CW10_Self Test.xlsx"), {"CW10": weekly}),
                   {"found": 2}),
        "accounts": (_book(p("Customer List_North Region selftest.xlsx"), {"North": [
            ["序号", "客户编码", "客户名称", "客户简称", "地区名称"],
            [1, "CR990001", CUSTOMER_CN, "自检", "North"], [2, "CR990002", "自检二号机电有限公司", "自检二", "North"]]}),
                     {"accounts": 2}),
        "contacts": (_book(p("Contacts selftest.xlsx"), {"Contacts": [
            ["序号", "公司名称", "联系人", "电话", "邮箱"],
            [1, CUSTOMER_CN, "Selftest Anna", "+8613900000001", "anna@selftest.example"],
            [2, CUSTOMER_CN, "Selftest Bo", "+8613900000002", "bo@selftest.example"]]}),
                     {"contacts": 2}),
        "todos": (_book(p("to-do selftest.xlsx"), {"Projekt-OPL": [["TOPIC - Selftest"], ["Project Responsible:"], TODO_HEAD,
            [1, None, DAY, "Weekly report", "send the weekly report", "Ethan", DAY, "WIP", 1, None, None],
            [2, None, DAY, "Customer list", "update the customer list", "Ethan", DAY, "Done", 2, None, None],
            [3, None, DAY, "Quotation", "send the quotation", "Ethan", DAY, "Open", 3, None, None]]}),
                  {"tasks": 3}),
        "orders": (_book(p("Sales Order List selftest.xlsx"), {"Orders": [["Sales Order List"], ORDER_HEAD,
            ["SELF001", DAY, CUSTOMER_CN, "Pat", "Open", "P-1", "Widget", 2, "pc", 100, 13, 113],
            ["SELF002", DAY, CUSTOMER_CN, "Pat", "Completed", "P-2", "Gadget", 1, "pc", 50, 6.5, 56.5]]}),
                   {"orders": 2}),
        "receivables": (_ar_book(p("2026_Financial_Status_20260410.xlsx")), {"customers": 1, "invoices": 2}),
    }


def _ar_book(path):
    wb = openpyxl.Workbook()
    wb.remove(wb.active)
    ov = wb.create_sheet("Overview")
    ov["G2"], ov["H2"] = "Date:", dt.datetime(2026, 4, 10)
    ws = wb.create_sheet("AR")
    ws["G2"], ws["G4"] = "Total Account Receivables", "Total Amount Overdue"
    ws["A3"], ws["G3"] = "Account Receivables", 1500
    ws["A5"], ws["B5"], ws["G5"] = dt.datetime(2026, 4, 13), "Today !", 1000
    for c, label in enumerate(AR_HEAD, 1):
        ws.cell(7, c, label)
    for i, (no, due, amount) in enumerate((("SELF-IN-1", dt.datetime(2026, 1, 1), 1000),
                                           ("SELF-IN-2", dt.datetime(2026, 9, 1), 500)), 8):
        for c, v in enumerate([CUSTOMER_EN, CUSTOMER_CN, no, dt.datetime(2025, 12, 1), due, 30, amount, None, None,
                               None, "人民币", amount, "CNY"], 1):
            ws.cell(i, c, v)
    wb.save(path)
    return path


def run() -> list[dict]:
    """One result per family: {"family", "ok", "expected", "got", "detected_as", "error"}. Never raises."""
    master = Master.from_records(MASTER)
    results = []
    with tempfile.TemporaryDirectory(prefix="importer-selftest-") as folder:
        books = make_workbooks(folder)
        readers = {
            "visits": lambda path: {"found": len(read_visits([path]).parsed.visits)},
            "accounts": lambda path: {"accounts": len(accounts.read_paths([path], master).accepted_accounts())},
            "contacts": lambda path: {"contacts": len(contacts.read_paths([path], master).accepted_contacts())},
            "todos": lambda path: {"tasks": len(todos.read_paths([path], master).todos)},
            "orders": lambda path: {"orders": len(orders.build_book([path], master).accepted())},
            "receivables": lambda path: _ar_counts(receivables.build_book([path], master)),
        }
        for family, (path, expected) in books.items():
            row = {"family": family, "expected": expected, "got": None, "ok": False, "error": "", "detected_as": None}
            try:
                row["detected_as"] = classify(path)["family"]  # the front door must send it to this reader
                row["got"] = readers[family](path)
                row["ok"] = row["got"] == expected and row["detected_as"] == family
            except Exception as exc:  # a reader that raises is a failed check, not a failed page
                row["error"] = f"{type(exc).__name__}: {exc}"
            results.append(row)
    return results


def _ar_counts(book) -> dict:
    return {"customers": len(book.accepted()), "invoices": len(book.sheet.invoices) if book.sheet else 0}
